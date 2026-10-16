#!/usr/bin/env python3
"""
Teamplaner – konfigurierbarer Wochenplaner (Version 1.0, Oktober 2026)
=======================================================================
Scheduler, Regelprüfung und Excel-Export für Teams, deren Regeln sich als
«so viele Personen pro Tag / Halbtag für Aufgabe X» beschreiben lassen –
zurzeit das KSV-Team (`team_ksv.json`). Das KSC-Team hat wegen seiner vielen
Spezialregeln einen eigenen Scheduler (`arbeitskalender.py`).

Alles Team-Spezifische steht in der Konfiguration und ist in der App unter
«Team & Regeln» änderbar – ohne Code:

- Personen: Pensum, Anwesenheit Mo–Fr (ganz / VM / NM / frei), alle 2 Wochen
  frei, Zeilenfarbe im Excel, Bemerkung
- Wer darf was: pro Aufgabe «ja», «Reserve» (nur wenn sonst niemand frei ist)
  oder «–»
- Feste Einsätze: Aufgabe und/oder Zusatztext (z. B. Arbeitszeit «12-16»),
  jede / gerade / ungerade KW
- Aufgaben: Code, Bedeutung, Farbe, Dauer (ganzer Tag / Halbtag / nur VM /
  nur NM / nur manuell) und Anzahl Personen pro Wochentag
- Auffüllen: womit freie Halbtage gefüllt werden (KSV: KSV2)

Planungsreihenfolge:
  0. Abwesenheiten / Sonderwünsche (Krank, Ferien, Frei, Termin, HO, Aufgabe)
  1. Feste Einsätze
  2. Aufgaben in der Reihenfolge der Liste (= Wichtigkeit):
     - ganzer Tag: dieselbe Person VM + NM; ist niemand ganztags frei,
       übernehmen zwei Personen je einen Halbtag (Hinweis)
     - Halbtag: pro Halbtag; bevorzugt Personen, deren anderer Halbtag schon
       belegt ist – so bleiben ganze Tage für Ganztags-Aufgaben frei
     - Verteilung: wer die Aufgabe diese Woche am wenigsten hat, dann wer sie
       gestern nicht hatte, dann wer insgesamt am wenigsten Aufgaben hat;
       «Reserve» nur, wenn niemand mit «ja» frei ist
     - mehrere Personen am selben Halbtag → nummeriert («contact 1», «contact 2»)
  3. Rest: Auffüll-Aufgabe (KSV2)
  4. Zusatztexte aus festen Einsätzen anhängen; Nicht-Arbeitstage = «Frei»

Das Modul importiert KEIN Streamlit und hat keine Nebenwirkungen: derselbe
Montag + dieselbe Konfiguration ergeben immer denselben Plan.
"""

from __future__ import annotations

import argparse
import copy
import json
import random
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

VERSION = "1.0"

# ============================================================
# KONSTANTEN
# ============================================================

DAYS = ["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag"]
DAY_KEYS = ["Mo", "Di", "Mi", "Do", "Fr"]
SLOTS = ["Vormittag", "Nachmittag"]
SLOT_KEYS = ["VM", "NM"]
HALBTAG_WERTE = ["ganz", "VM", "NM", "frei"]
PARITAET = ["gerade", "ungerade"]

# Dauer einer Aufgabe: welche Halbtage sie belegt und ob der Generator sie plant
DAUER = ["ganzer Tag", "Halbtag", "nur VM", "nur NM", "nur manuell"]
AUTO_DAUER = DAUER[:4]
DAUER_SLOTS = {"ganzer Tag": (0, 1), "Halbtag": (0, 1), "nur VM": (0,), "nur NM": (1,), "nur manuell": ()}
DAUER_HILFE = {
    "ganzer Tag": "dieselbe Person am Vormittag und am Nachmittag",
    "Halbtag": "pro Halbtag – Vormittag und Nachmittag einzeln besetzt",
    "nur VM": "nur am Vormittag",
    "nur NM": "nur am Nachmittag",
    "nur manuell": "plant der Generator nie – über feste Einsätze, Schritt 1 oder «Plan bearbeiten»",
}

# Wer darf was: Stufen pro Person und Aufgabe
STUFEN = ["ja", "Reserve", "–"]

# Abwesenheiten (fest eingebaut): Code → Farbe. Gross-/Kleinschreibung egal.
ABWESEND = {"KRANK": "BFBFBF", "Ferien": "00B0F0", "Frei": "00B0F0"}

# Farben mit Namen (für die Auswahl in der App) – Excel-Standardfarben
FARBEN: Dict[str, str] = {
    "Hellgrün": "92D050", "Grün": "00B050", "Rot": "FF0000", "Orange": "FFC000",
    "Gelb": "FFFF00", "Violett": "7030A0", "Blau": "00B0F0", "Dunkelblau": "0070C0",
    "Türkis": "00FFFF", "Pink": "FF66CC", "Grau": "BFBFBF", "Hellgrau": "D9D9D9",
    "Rosa": "FFC7CE", "Mint": "C6EFCE", "Hellgelb": "FFEB9C", "Hellblau": "BDD7EE",
    "Lavendel": "D9D2E9", "Pfirsich": "FCE4D6", "Blassgrün": "E2EFDA", "Weiss": "FFFFFF",
}
# Zeilenfarben für neue Personen (Reihenfolge wie in der KSV-Vorlage)
ZEILENFARBEN = ["FFC7CE", "C6EFCE", "FFEB9C", "BDD7EE", "D9D2E9", "FCE4D6", "E2EFDA"]

DEFAULT_REGELN = {
    "auffuellen": "",        # Code für freie Halbtage ("" = leer lassen)
    "frei_anzeigen": True,   # Nicht-Arbeitstage laut Anwesenheit als «Frei» eintragen
}

DEFAULT_CONFIG_PATH = Path(__file__).resolve().parent / "team_ksv.json"


# ============================================================
# HILFEN
# ============================================================

def norm_hex(value: str, fallback: str = "D9D9D9") -> str:
    """'#92d050' → '92D050'; ungültig → fallback."""
    h = str(value or "").strip().lstrip("#").upper()
    return h if re.fullmatch(r"[0-9A-F]{6}", h) else fallback


def is_dark(hex_color: str) -> bool:
    h = norm_hex(hex_color)
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255 < 0.55


def farbname(hex_color: str) -> str:
    """Name zur Farbe aus FARBEN, sonst '#RRGGBB'."""
    h = norm_hex(hex_color)
    for name, val in FARBEN.items():
        if val == h:
            return name
    return f"#{h}"


def split_name(name: str) -> Tuple[str, str]:
    """'Medine A.' → ('Medine', 'A.'), 'Jana' → ('Jana', '')."""
    parts = (name or "").split()
    if len(parts) >= 2 and parts[-1].endswith("."):
        return " ".join(parts[:-1]), parts[-1]
    return (name or "").strip(), ""


def halbtag_slots(value: str) -> List[int]:
    return {"ganz": [0, 1], "VM": [0], "NM": [1], "frei": []}.get(value, [0, 1])


def kw_paritaet(week_number: int) -> str:
    return "gerade" if week_number % 2 == 0 else "ungerade"


def get_next_monday(ref: Optional[datetime] = None) -> datetime:
    ref = ref or datetime.now()
    days = 0 - ref.weekday()
    if days <= 0:
        days += 7
    return ref + timedelta(days=days)


def week_index(monday) -> int:
    if isinstance(monday, datetime):
        monday = monday.date()
    return (monday - date(2024, 1, 1)).days // 7


# ============================================================
# KONFIGURATION
# ============================================================

def load_config(path: Optional[Path | str] = None) -> dict:
    p = Path(path) if path else DEFAULT_CONFIG_PATH
    if not p.exists():
        raise FileNotFoundError(f"Konfiguration {p} nicht gefunden.")
    with open(p, "r", encoding="utf-8") as fh:
        return normalize_config(json.load(fh))


def save_config(cfg: dict, path: Optional[Path | str] = None) -> Path:
    p = Path(path) if path else DEFAULT_CONFIG_PATH
    with open(p, "w", encoding="utf-8") as fh:
        fh.write(config_to_json(cfg))
    return p


def config_to_json(cfg: dict) -> str:
    return json.dumps(cfg, ensure_ascii=False, indent=2) + "\n"


def normalize_config(cfg: dict) -> dict:
    """Füllt Standardwerte auf, damit alter/handgeschriebener JSON-Inhalt läuft."""
    cfg = copy.deepcopy(cfg or {})
    cfg.setdefault("version", VERSION)
    regeln = cfg.setdefault("regeln", {})
    for k, v in DEFAULT_REGELN.items():
        regeln.setdefault(k, v)
    for t in cfg.setdefault("aufgaben", []):
        t["code"] = str(t.get("code", "")).strip()
        t.setdefault("bedeutung", "")
        t["farbe"] = norm_hex(t.get("farbe"))
        if t.get("dauer") not in DAUER:
            t["dauer"] = "nur manuell"
        anz = t.get("anzahl") or {}
        t["anzahl"] = {dk: max(0, int(anz.get(dk, 0) or 0)) for dk in DAY_KEYS}
    cfg.setdefault("backup_team", [])
    for b in cfg["backup_team"]:
        b.setdefault("name", b.get("short", ""))
        b.setdefault("standard", "")
    for i, m in enumerate(cfg.setdefault("mitarbeiter", [])):
        m.setdefault("name", m.get("short", ""))
        m.setdefault("pct", 100)
        m.setdefault("aktiv", True)
        m.setdefault("kann", [])
        m.setdefault("reserve", [])
        m.setdefault("anwesend", {})
        m.setdefault("frei_alle_2_wochen", None)
        m.setdefault("fix", [])
        m["farbe"] = norm_hex(m.get("farbe"), ZEILENFARBEN[i % len(ZEILENFARBEN)])
        m.setdefault("bemerkung", "")
        for fx in m["fix"]:
            fx.setdefault("task", "")
            fx.setdefault("zusatz", "")
    return cfg


def task_list(cfg: dict) -> List[dict]:
    return cfg.get("aufgaben", [])


def task_codes(cfg: dict) -> List[str]:
    return [t["code"] for t in task_list(cfg) if t.get("code")]


def auto_tasks(cfg: dict) -> List[dict]:
    """Aufgaben, die der Generator plant – in der Reihenfolge der Liste (= Wichtigkeit)."""
    return [t for t in task_list(cfg) if t.get("code") and t.get("dauer") in AUTO_DAUER]


def task_by_code(cfg: dict, code: str) -> Optional[dict]:
    return next((t for t in task_list(cfg) if t.get("code") == code), None)


def _match(task: str, codes: Iterable[str]) -> str:
    """Längster Code, mit dem die Zelle beginnt (Gross-/Kleinschreibung egal)."""
    t = (task or "").strip().lower()
    if not t:
        return ""
    for code in sorted(codes, key=len, reverse=True):
        if t.startswith(code.lower()):
            return code
    return ""


def is_absent_code(task: str) -> bool:
    return bool(_match(task, ABWESEND))


def base_code(task: str, cfg: dict) -> str:
    """Welcher bekannte Code steckt in der Zelle? ('contact 2' → 'contact', 'KSV2 12-16' → 'KSV2').
    Abwesenheiten liefern ihren Code (KRANK/Ferien/Frei), Unbekanntes ''."""
    return _match(task, list(ABWESEND) + task_codes(cfg))


def get_task_color(task: str, cfg: dict) -> Optional[str]:
    code = base_code(task, cfg)
    if not code:
        return None
    if code in ABWESEND:
        return ABWESEND[code]
    t = task_by_code(cfg, code)
    return t["farbe"] if t else None


def legend_items(cfg: dict) -> List[Tuple[str, str, str]]:
    """(Code, Bedeutung, Farbe) für Legende in App und Excel."""
    items = [(t["code"], t.get("bedeutung") or t["code"], t["farbe"]) for t in task_list(cfg) if t.get("code")]
    items += [("Frei", "frei (Teilzeit / Kompensation)", ABWESEND["Frei"]),
              ("Ferien", "Ferien", ABWESEND["Ferien"]), ("KRANK", "krank", ABWESEND["KRANK"])]
    return items


def editor_options(cfg: dict) -> List[str]:
    """Auswahl im Plan-Editor: alle Codes, bei mehreren Personen pro Tag nummeriert."""
    opts = [""]
    for t in task_list(cfg):
        if not t.get("code"):
            continue
        opts.append(t["code"])
        top = max(t["anzahl"].values()) if t.get("dauer") in AUTO_DAUER else 0
        opts += [f"{t['code']} {i}" for i in range(1, top + 1)] if top > 1 else []
    return opts + ["Frei", "Ferien", "KRANK"]


# ── Prüfung der Konfiguration ─────────────────────────────────────────

def validate_config(cfg: dict) -> List[str]:
    return config_errors(cfg) + config_hints(cfg)


def config_errors(cfg: dict) -> List[str]:
    """Fehler, mit denen die Konfiguration nicht übernommen werden darf."""
    probs: List[str] = []
    codes = [t.get("code", "") for t in task_list(cfg)]
    if any(not c for c in codes):
        probs.append("Aufgabe ohne Code")
    low = [c.lower() for c in codes if c]
    dups = {c for c in codes if c and low.count(c.lower()) > 1}
    if dups:
        probs.append(f"Aufgabe doppelt: {', '.join(sorted(dups))}")
    reserved = [c for c in codes if c and c.lower() in {a.lower() for a in ABWESEND}]
    if reserved:
        probs.append(f"«{reserved[0]}» ist für Abwesenheiten reserviert – bitte anderen Code wählen")
    for t in task_list(cfg):
        if t.get("dauer") not in DAUER:
            probs.append(f"Aufgabe {t.get('code')}: Dauer «{t.get('dauer')}» ungültig")
    fill = cfg.get("regeln", {}).get("auffuellen", "")
    if fill and fill not in codes:
        probs.append(f"Auffüllen mit «{fill}»: diese Aufgabe gibt es nicht")
    shorts = [m.get("short", "") for m in cfg.get("mitarbeiter", [])]
    if any(not s for s in shorts):
        probs.append("Person ohne Kurzname")
    dup_s = {s for s in shorts if s and shorts.count(s) > 1}
    if dup_s:
        probs.append(f"Kurzname doppelt: {', '.join(sorted(dup_s))}")
    backup_shorts = [b.get("short", "") for b in cfg.get("backup_team", [])]
    overlap = set(shorts) & set(backup_shorts)
    if overlap:
        probs.append(f"Kurzname im Team UND im BackUp-Team: {', '.join(sorted(overlap))}")
    names = [(m.get("name") or m.get("short", "")).strip() for m in cfg.get("mitarbeiter", [])] + \
            [(b.get("name") or b.get("short", "")).strip() for b in cfg.get("backup_team", [])]
    dup_n = {n for n in names if n and names.count(n) > 1}
    if dup_n:
        probs.append(f"Name doppelt: {', '.join(sorted(dup_n))}")
    for m in cfg.get("mitarbeiter", []):
        who = m.get("name") or m.get("short")
        for r in list(m.get("kann", [])) + list(m.get("reserve", [])):
            if r not in codes:
                probs.append(f"{who}: unbekannte Aufgabe «{r}» unter «Wer darf was»")
        for dk, val in m.get("anwesend", {}).items():
            if dk not in DAY_KEYS or val not in HALBTAG_WERTE:
                probs.append(f"{who}: Anwesenheit {dk}={val} ungültig")
        f2 = m.get("frei_alle_2_wochen")
        if f2:
            if f2.get("kw") not in PARITAET:
                probs.append(f"{who}: frei_alle_2_wochen.kw muss gerade/ungerade sein")
            for dk, val in f2.get("tage", {}).items():
                if dk not in DAY_KEYS or val not in HALBTAG_WERTE:
                    probs.append(f"{who}: frei_alle_2_wochen {dk}={val} ungültig")
        for fx in m.get("fix", []):
            if fx.get("tag") not in DAY_KEYS or fx.get("halbtag") not in HALBTAG_WERTE:
                probs.append(f"{who}: fester Einsatz {fx} ungültig")
            if not fx.get("task") and not fx.get("zusatz"):
                probs.append(f"{who}: fester Einsatz ohne Aufgabe und ohne Zusatz")
            if fx.get("kw") not in (None, "", "gerade", "ungerade"):
                probs.append(f"{who}: fester Einsatz – Wochen muss leer/gerade/ungerade sein")
    return probs


def config_hints(cfg: dict) -> List[str]:
    """Hinweise: der Generator läuft, kann aber z. B. eine Aufgabe nie besetzen."""
    probs: List[str] = []
    act = [m for m in cfg.get("mitarbeiter", []) if m.get("aktiv", True)]
    for t in auto_tasks(cfg):
        if sum(t["anzahl"].values()) and not any(t["code"] in m.get("kann", []) + m.get("reserve", []) for m in act):
            probs.append(f"{t['code']}: niemand darf diese Aufgabe – unter «Wer darf was» eintragen")
    for d, dk in enumerate(DAY_KEYS):
        full = sum(t["anzahl"][dk] for t in auto_tasks(cfg) if t["dauer"] == "ganzer Tag")
        here = sum(1 for m in act if m.get("anwesend", {}).get(dk, "ganz") == "ganz")
        if full > here:
            probs.append(f"{DAYS[d]}: {full} Personen ganztags nötig, laut Anwesenheit nur {here} ganztags da")
    if not cfg.get("regeln", {}).get("auffuellen"):
        probs.append("Kein «Auffüllen mit» gewählt – freie Halbtage bleiben leer")
    return probs


# ── Team bearbeiten (für die App; ohne Streamlit testbar) ─────────────

def new_short(name: str, cfg: dict) -> str:
    """Interner Schlüssel für eine neue Person: Vorname, bei Kollision der ganze Name."""
    name = name.strip()
    taken = {m["short"] for m in cfg.get("mitarbeiter", [])} | {b["short"] for b in cfg.get("backup_team", [])}
    first = name.split()[0] if name.split() else name
    for cand in (first, name):
        if cand and cand not in taken:
            return cand
    i = 2
    while f"{name} {i}" in taken:
        i += 1
    return f"{name} {i}"


def rotation_refs(cfg: dict, short: str) -> List[str]:
    """Dieser Planer kennt keine Rotationslisten – Kompatibilität mit der App."""
    return []


def remove_member(cfg: dict, short: str) -> None:
    cfg["mitarbeiter"] = [m for m in cfg["mitarbeiter"] if m["short"] != short]


def stufe(m: dict, code: str) -> str:
    """'ja' / 'Reserve' / '–' für Person m und Aufgabe code."""
    if code in m.get("kann", []):
        return "ja"
    if code in m.get("reserve", []):
        return "Reserve"
    return "–"


def set_stufe(m: dict, code: str, value: str) -> None:
    m["kann"] = [c for c in m.get("kann", []) if c != code] + ([code] if value == "ja" else [])
    m["reserve"] = [c for c in m.get("reserve", []) if c != code] + ([code] if value == "Reserve" else [])


def rename_task(cfg: dict, old: str, new: str) -> None:
    """Code umbenennen – auch in Wer-darf-was, festen Einsätzen und «Auffüllen»."""
    for t in task_list(cfg):
        if t["code"] == old:
            t["code"] = new
    for m in cfg.get("mitarbeiter", []):
        m["kann"] = [new if c == old else c for c in m.get("kann", [])]
        m["reserve"] = [new if c == old else c for c in m.get("reserve", [])]
        for fx in m.get("fix", []):
            if fx.get("task") == old:
                fx["task"] = new
    if cfg["regeln"].get("auffuellen") == old:
        cfg["regeln"]["auffuellen"] = new


def remove_task(cfg: dict, code: str) -> None:
    """Aufgabe löschen – samt Häkchen unter «Wer darf was» und «Auffüllen»."""
    cfg["aufgaben"] = [t for t in task_list(cfg) if t["code"] != code]
    for m in cfg.get("mitarbeiter", []):
        m["kann"] = [c for c in m.get("kann", []) if c != code]
        m["reserve"] = [c for c in m.get("reserve", []) if c != code]
    if cfg["regeln"].get("auffuellen") == code:
        cfg["regeln"]["auffuellen"] = ""


def recurring_rows(cfg: dict) -> List[dict]:
    """Feste Einsätze und «alle 2 Wochen frei» als eine Liste:
    {short, tag, halbtag, task, zusatz, kw}  (task "frei" = alle 2 Wochen frei)."""
    rows = []
    for m in cfg["mitarbeiter"]:
        for fx in m.get("fix", []):
            rows.append({"short": m["short"], "tag": fx.get("tag", "Mo"), "halbtag": fx.get("halbtag", "ganz"),
                         "task": fx.get("task", ""), "zusatz": fx.get("zusatz", ""), "kw": fx.get("kw") or None})
        f2 = m.get("frei_alle_2_wochen") or {}
        for dk, val in f2.get("tage", {}).items():
            rows.append({"short": m["short"], "tag": dk, "halbtag": val, "task": "frei", "zusatz": "",
                         "kw": f2.get("kw")})
    return rows


def apply_recurring_rows(cfg: dict, rows: List[dict]) -> List[str]:
    """Schreibt die Liste aus recurring_rows() zurück. Liefert Fehler (dann bleibt cfg unverändert)."""
    names = {m["short"]: m.get("name") or m["short"] for m in cfg["mitarbeiter"]}
    per = {s: {"fix": [], "frei": {}, "kw": None} for s in names}
    errors = []
    for r in rows:
        p = per.get(r["short"])
        if p is None or not (r.get("task") or r.get("zusatz")):
            continue
        if r.get("task") == "frei":
            if r.get("kw") not in PARITAET:
                errors.append(f"{names[r['short']]}: «frei» jede Woche bitte unter Personen → Anwesenheit "
                              f"({r['tag']}) eintragen; hier nur «gerade KW» oder «ungerade KW»")
            elif p["kw"] and p["kw"] != r["kw"]:
                errors.append(f"{names[r['short']]}: «frei» geht nur in geraden ODER in ungeraden KW, nicht gemischt")
            else:
                p["kw"] = r["kw"]
                p["frei"][r["tag"]] = r["halbtag"]
        else:
            fx = {"tag": r["tag"], "halbtag": r["halbtag"], "task": r.get("task") or "",
                  "zusatz": (r.get("zusatz") or "").strip()}
            if r.get("kw") in PARITAET:
                fx["kw"] = r["kw"]
            p["fix"].append(fx)
    if errors:
        return errors
    for m in cfg["mitarbeiter"]:
        p = per[m["short"]]
        m["fix"] = p["fix"]
        m["frei_alle_2_wochen"] = {"kw": p["kw"], "tage": p["frei"]} if p["frei"] else None
    return []


REGEL_LABELS = {"auffuellen": "Auffüllen mit", "frei_anzeigen": "Nicht-Arbeitstage als «Frei»"}


def describe_changes(old: dict, new: dict) -> List[str]:
    """Lesbare Liste der Unterschiede zwischen zwei Konfigurationen (für «ungespeichert»)."""
    old, new = normalize_config(old), normalize_config(new)
    out: List[str] = []
    om = {m["short"]: m for m in old["mitarbeiter"]}
    nm = {m["short"]: m for m in new["mitarbeiter"]}

    def name(s):
        return (nm.get(s) or om.get(s) or {}).get("name") or s

    def fx_str(fx):
        kw = f" ({fx['kw']} KW)" if fx.get("kw") else ""
        txt = " ".join(x for x in (fx.get("task"), fx.get("zusatz")) if x)
        return f"{fx.get('tag')} {fx.get('halbtag', 'ganz')} {txt}{kw}"

    out += [f"Neue Person: {name(s)}" for s in nm if s not in om]
    out += [f"Entfernt: {name(s)}" for s in om if s not in nm]
    for s, m in nm.items():
        o = om.get(s)
        if o is None:
            continue
        n = m.get("name") or s
        if o.get("name") != m.get("name"):
            out.append(f"{o.get('name')} heisst neu {m.get('name')}")
        if o.get("pct") != m.get("pct"):
            out.append(f"{n}: Pensum {o.get('pct')} → {m.get('pct')} %")
        if bool(o.get("aktiv", True)) != bool(m.get("aktiv", True)):
            out.append(f"{n}: {'wieder aktiv' if m.get('aktiv', True) else 'inaktiv (wird nicht eingeplant)'}")
        for dk in DAY_KEYS:
            a, b = o["anwesend"].get(dk, "ganz"), m["anwesend"].get(dk, "ganz")
            if a != b:
                out.append(f"{n}: Anwesenheit {dk} {a} → {b}")
        for code in sorted(set(o["kann"] + o["reserve"] + m["kann"] + m["reserve"])):
            a, b = stufe(o, code), stufe(m, code)
            if a != b:
                out.append(f"{n}: {code} {a} → {b}")
        ofx, nfx = sorted(map(fx_str, o["fix"])), sorted(map(fx_str, m["fix"]))
        out += [f"{n}: fester Einsatz {x}" for x in nfx if x not in ofx]
        out += [f"{n}: fester Einsatz entfernt: {x}" for x in ofx if x not in nfx]
        if (o.get("frei_alle_2_wochen") or None) != (m.get("frei_alle_2_wochen") or None):
            f2 = m.get("frei_alle_2_wochen")
            out.append(f"{n}: alle 2 Wochen frei → " +
                       (f"{', '.join(f'{d} {v}' for d, v in f2['tage'].items())} ({f2['kw']} KW)" if f2 else "keine"))
        if o.get("farbe") != m.get("farbe"):
            out.append(f"{n}: Zeilenfarbe {farbname(o.get('farbe'))} → {farbname(m.get('farbe'))}")
        if (o.get("bemerkung") or "") != (m.get("bemerkung") or ""):
            out.append(f"{n}: Bemerkung geändert")
    if [s for s in om if s in nm] != [s for s in nm if s in om]:
        out.append("Reihenfolge im Plan geändert")
    ot = {t["code"]: t for t in old["aufgaben"]}
    nt = {t["code"]: t for t in new["aufgaben"]}
    out += [f"Neue Aufgabe: {c}" for c in nt if c not in ot]
    out += [f"Aufgabe entfernt: {c}" for c in ot if c not in nt]
    for c, t in nt.items():
        o = ot.get(c)
        if o is None:
            continue
        if o["dauer"] != t["dauer"]:
            out.append(f"{c}: Dauer {o['dauer']} → {t['dauer']}")
        if o["anzahl"] != t["anzahl"] and t["dauer"] in AUTO_DAUER:
            anz = " ".join(f"{d} {t['anzahl'][d]}" for d in DAY_KEYS)
            out.append(f"{c}: Anzahl pro Tag {anz}")
        if o["farbe"] != t["farbe"]:
            out.append(f"{c}: Farbe {farbname(o['farbe'])} → {farbname(t['farbe'])}")
        if (o.get("bedeutung") or "") != (t.get("bedeutung") or ""):
            out.append(f"{c}: Bedeutung «{t.get('bedeutung')}»")
    if [c for c in ot if c in nt] != [c for c in nt if c in ot]:
        out.append("Reihenfolge der Aufgaben geändert")
    for k, lbl in REGEL_LABELS.items():
        if old["regeln"].get(k) != new["regeln"].get(k):
            out.append(f"Regel «{lbl}»: {old['regeln'].get(k)} → {new['regeln'].get(k)}")
    if old["backup_team"] != new["backup_team"]:
        out.append("BackUp-Team geändert")
    if not out and config_to_json(old) != config_to_json(new):
        out.append("Weitere Einstellungen geändert")
    return out


# ============================================================
# PERSONEN
# ============================================================

@dataclass
class Employee:
    short: str
    name: str
    pct: int
    kann: set = field(default_factory=set)
    reserve: set = field(default_factory=set)
    available: Dict[Tuple[int, int], bool] = field(default_factory=dict)
    fix: List[Tuple[int, int, str, str]] = field(default_factory=list)   # (Tag, Halbtag, Aufgabe, Zusatz)
    farbe: str = "FFFFFF"
    bemerkung: str = ""

    def is_available(self, day: int, slot: int) -> bool:
        return self.available.get((day, slot), False)

    def level(self, code: str) -> str:
        return "ja" if code in self.kann else "Reserve" if code in self.reserve else "–"


def create_employees(week_number: int, cfg: Optional[dict] = None) -> Dict[str, Employee]:
    """Alle aktiven Personen für diese KW (Anwesenheit inkl. 2-Wochen-Rhythmus)."""
    cfg = normalize_config(cfg if cfg is not None else load_config())
    par = kw_paritaet(week_number)
    employees: Dict[str, Employee] = {}
    for m in cfg["mitarbeiter"]:
        if not m.get("aktiv", True):
            continue
        e = Employee(short=m["short"], name=m.get("name") or m["short"], pct=int(m.get("pct", 100)),
                     kann=set(m.get("kann", [])), reserve=set(m.get("reserve", [])),
                     farbe=m.get("farbe") or "FFFFFF", bemerkung=m.get("bemerkung", ""))
        for d, dk in enumerate(DAY_KEYS):
            slots = halbtag_slots(m.get("anwesend", {}).get(dk, "ganz"))
            for s in range(2):
                e.available[(d, s)] = s in slots
        f2 = m.get("frei_alle_2_wochen")
        if f2 and f2.get("kw") == par:
            for dk, val in f2.get("tage", {}).items():
                if dk in DAY_KEYS:
                    for s in halbtag_slots(val):
                        e.available[(DAY_KEYS.index(dk), s)] = False
        for fx in m.get("fix", []):
            if fx.get("kw") in PARITAET and fx["kw"] != par:
                continue
            if fx.get("tag") not in DAY_KEYS:
                continue
            d = DAY_KEYS.index(fx["tag"])
            for s in halbtag_slots(fx.get("halbtag", "ganz")):
                e.fix.append((d, s, fx.get("task", ""), fx.get("zusatz", "")))
        employees[e.short] = e
    return employees


def two_week_notes(week_number: int, cfg: dict) -> List[str]:
    """Für die UI: wer hat in dieser KW seinen 2-Wochen-Rhythmus."""
    par = kw_paritaet(week_number)
    notes = []
    for m in normalize_config(cfg)["mitarbeiter"]:
        if not m.get("aktiv", True):
            continue
        n = m.get("name") or m["short"]
        f2 = m.get("frei_alle_2_wochen")
        if f2 and f2.get("kw") == par:
            tage = ", ".join(f"{dk} {val}" if val != "ganz" else dk for dk, val in f2.get("tage", {}).items())
            notes.append(f"{n}: {tage} frei")
        for fx in m.get("fix", []):
            if fx.get("kw") == par:
                txt = " ".join(x for x in (fx.get("task"), fx.get("zusatz")) if x)
                notes.append(f"{n}: {fx.get('tag')} {fx.get('halbtag', 'ganz')} {txt}")
    return notes


# ============================================================
# SCHEDULE
# ============================================================

class Schedule:
    def __init__(self, employees: Dict[str, Employee], week_number: int, week_start_date, cfg: dict):
        self.employees = employees
        self.week_number = week_number
        self.week_start_date = week_start_date
        self.cfg = cfg
        self.schedule: Dict[str, Dict[Tuple[int, int], str]] = {n: {} for n in employees}
        self.notes: Dict[Tuple[str, int, int], str] = {}
        self.backup_tasks: Dict[Tuple[str, int, int], str] = {}
        self.conflicts: List[str] = []

    def assign(self, name, day, slot, task) -> bool:
        if name not in self.employees or (day, slot) in self.schedule[name]:
            return False
        if not self.employees[name].is_available(day, slot):
            return False
        self.schedule[name][(day, slot)] = task
        return True

    def force_assign(self, name, day, slot, task) -> None:
        if name in self.employees:
            self.schedule[name][(day, slot)] = task

    def is_free(self, name, day, slot) -> bool:
        return (name in self.employees and (day, slot) not in self.schedule[name]
                and self.employees[name].is_available(day, slot))

    def is_free_day(self, name, day) -> bool:
        return self.is_free(name, day, 0) and self.is_free(name, day, 1)

    def get_task(self, name, day, slot) -> str:
        return self.schedule.get(name, {}).get((day, slot), "")

    def get_backup_task(self, name, day, slot) -> str:
        if (name, day, slot) in self.backup_tasks:
            return self.backup_tasks[(name, day, slot)]
        for b in self.cfg.get("backup_team", []):
            if b.get("short") == name:
                return b.get("standard", "")
        return ""

    def is_absent(self, name, day, slot) -> bool:
        return is_absent_code(self.get_task(name, day, slot)) or (name, day, slot) in self.notes

    def is_present(self, name, day, slot) -> bool:
        return (name in self.employees and self.employees[name].is_available(day, slot)
                and not self.is_absent(name, day, slot))

    def present_anytime(self, name) -> bool:
        return any(self.is_present(name, d, s) for d in range(5) for s in range(2))

    def holders(self, code: str, day: int, slot: int) -> List[str]:
        return [n for n in self.employees if base_code(self.get_task(n, day, slot), self.cfg) == code]

    def count_task(self, name, pred) -> int:
        return sum(1 for t in self.schedule.get(name, {}).values() if pred(t))


# ============================================================
# SCHEDULER
# ============================================================

def build_schedule(week_number: int, week_start_date, overrides: Optional[dict] = None,
                   cfg: Optional[dict] = None, seed: Optional[int] = None) -> Schedule:
    """Baut den Wochenplan.

    overrides: {(short, day, slot): task} – "KRANK"/"Ferien"/"Frei", "*Notiz" (Termin,
               Halbtag bleibt leer) oder ein beliebiger Code (HO, Einführung …).
               Einträge für BackUp-Namen landen 1:1 in sched.backup_tasks.
    seed:      Zufalls-Seed; Standard = ISO-Jahr*100 + ISO-KW (deterministisch).
    """
    cfg = normalize_config(cfg if cfg is not None else load_config())
    employees = create_employees(week_number, cfg)
    sched = Schedule(employees, week_number, week_start_date, cfg)
    order = [m["short"] for m in cfg["mitarbeiter"] if m["short"] in employees]
    backup_names = {b["short"] for b in cfg["backup_team"]}
    iso_year, iso_week, _ = week_start_date.isocalendar()
    rng = random.Random(seed if seed is not None else iso_year * 100 + iso_week)

    # 0. Abwesenheiten / Sonderwünsche / BackUp-Einträge
    for (name, day, slot), task in (overrides or {}).items():
        task = task or ""
        if name in employees:
            if task.startswith("*"):
                sched.notes[(name, day, slot)] = task[1:] or "Termin"
                sched.force_assign(name, day, slot, "")
            elif is_absent_code(task):
                sched.force_assign(name, day, slot, task)
            else:
                sched.assign(name, day, slot, task)
        elif name in backup_names:
            sched.backup_tasks[(name, day, slot)] = task.lstrip("*")

    # 1. Feste Einsätze
    for name in order:
        for d, s, task, _zusatz in employees[name].fix:
            if task:
                sched.assign(name, d, s, task)

    # 2. Aufgaben in Listen-Reihenfolge
    belastung: Dict[str, int] = {n: 0 for n in order}   # Halbtage mit geplanten Aufgaben diese Woche

    def best(cands: List[str], key) -> Optional[str]:
        scored = sorted((key(n), rng.random(), n) for n in cands)
        return scored[0][2] if scored else None

    for t in auto_tasks(cfg):
        code, dauer = t["code"], t["dauer"]
        used: Dict[str, float] = {}
        gestern: set = set()
        for day in range(5):
            need = t["anzahl"][DAY_KEYS[day]]
            heute: set = set()

            def take(name, slots):
                for s in slots:
                    sched.assign(name, day, s, code)
                used[name] = used.get(name, 0) + len(slots) / (2 if dauer == "ganzer Tag" else 1)
                belastung[name] += len(slots)
                heute.add(name)

            def pool(stufe_):
                return [n for n in order if employees[n].level(code) == stufe_]

            def note_reserve(name, wo):
                sched.conflicts.append(f"ℹ️ {DAYS[day]}{wo}: {code} – Reserve eingesetzt ({employees[name].name})")

            def chancen(name, slots):
                """An wie vielen Tagen ab heute könnte die Person diese Aufgabe noch übernehmen?
                Wer seltener da ist (Teilzeit), kommt an seinen Tagen zuerst dran."""
                return sum(1 for d in range(day, 5) if t["anzahl"][DAY_KEYS[d]]
                           and all(sched.is_free(name, d, s) for s in slots))

            if need > 0 and dauer == "ganzer Tag":
                for _ in range(max(0, min(need - len(sched.holders(code, day, s)) for s in (0, 1)))):
                    chosen = None
                    for st in ("ja", "Reserve"):
                        cands = [n for n in pool(st) if sched.is_free_day(n, day)]
                        chosen = best(cands, lambda n: (used.get(n, 0), n in gestern, chancen(n, (0, 1)),
                                                        belastung[n]))
                        if chosen:
                            if st == "Reserve":
                                note_reserve(chosen, "")
                            break
                    if not chosen:
                        break
                    take(chosen, (0, 1))
            if need > 0:
                # Halbtags-Aufgaben – und Ganztags-Aufgaben, die niemand ganztags übernehmen kann
                for s in DAUER_SLOTS[dauer]:
                    while len(sched.holders(code, day, s)) < need:
                        chosen = None
                        for st in ("ja", "Reserve"):
                            cands = [n for n in pool(st) if sched.is_free(n, day, s)]
                            chosen = best(cands, lambda n: (used.get(n, 0), sched.is_free(n, day, 1 - s),
                                                            n in gestern, belastung[n]))
                            if chosen:
                                if st == "Reserve":
                                    note_reserve(chosen, f" {SLOT_KEYS[s]}")
                                break
                        if not chosen:
                            break
                        take(chosen, (s,))
                # mehrere Personen am selben Halbtag → nummerieren
                if need > 1:
                    for s in DAUER_SLOTS[dauer]:
                        hs = sorted(sched.holders(code, day, s),
                                    key=lambda n: (base_code(sched.get_task(n, day, 1 - s), cfg) != code,
                                                   order.index(n) if n in order else 99))
                        for i, n in enumerate(hs, 1):
                            if sched.get_task(n, day, s).strip().lower() == code.lower():
                                sched.schedule[n][(day, s)] = f"{code} {i}"
            gestern = heute

    # 3. Rest auffüllen
    fill = cfg["regeln"].get("auffuellen") or ""
    if fill:
        for name in order:
            for d in range(5):
                for s in range(2):
                    sched.assign(name, d, s, fill)

    # 4. Zusatztexte (z. B. Arbeitszeit) anhängen, Nicht-Arbeitstage = «Frei»
    for name in order:
        for d, s, _task, zusatz in employees[name].fix:
            cur = sched.get_task(name, d, s)
            if zusatz and cur and not is_absent_code(cur) and not cur.endswith(zusatz):
                sched.schedule[name][(d, s)] = f"{cur} {zusatz}"
    if cfg["regeln"].get("frei_anzeigen"):
        for name in order:
            for d in range(5):
                for s in range(2):
                    if not employees[name].is_available(d, s) and not sched.get_task(name, d, s):
                        sched.schedule[name][(d, s)] = "Frei"
    return sched


# ============================================================
# AUSWERTUNG & REGELPRÜFUNG
# ============================================================

def day_overview(sched: Schedule) -> List[Tuple[str, List[str]]]:
    """Pro geplanter Aufgabe und Tag, wer sie hat: [(Code, [Mo … Fr])]."""
    cfg = sched.cfg
    out = []
    for t in auto_tasks(cfg):
        code = t["code"]
        cells = []
        for d in range(5):
            need = t["anzahl"][DAY_KEYS[d]]
            per = {s: [sched.employees[n].name for n in sched.holders(code, d, s)] for s in range(2)}
            slots = DAUER_SLOTS[t["dauer"]]
            if not need and not any(per.values()):
                cells.append("–")
                continue
            if len(slots) == 2 and per[0] == per[1]:
                txt = ", ".join(per[0]) or "⚠️ OFFEN"
            else:
                txt = " · ".join(f"{SLOT_KEYS[s]}: {', '.join(per[s]) or '⚠️ OFFEN'}" for s in slots)
            if any(len(per[s]) < need for s in slots):
                txt += " ⚠️" if "OFFEN" not in txt else ""
            cells.append(txt)
        out.append((code, cells))
    return out


def person_summary(sched: Schedule) -> List[dict]:
    """Verteilung pro Person: Halbtage je geplanter Aufgabe (für die Übersicht)."""
    cfg = sched.cfg
    codes = [t["code"] for t in auto_tasks(cfg)]
    fill = cfg["regeln"].get("auffuellen") or ""
    if fill and fill not in codes:
        codes.append(fill)
    rows = []
    for m in cfg["mitarbeiter"]:
        n = m["short"]
        if n not in sched.employees:
            continue
        cnt = {c: 0 for c in codes}
        da = 0
        for d in range(5):
            for s in range(2):
                b = base_code(sched.get_task(n, d, s), cfg)
                if b in cnt:
                    cnt[b] += 1
                if sched.is_present(n, d, s):
                    da += 1
        rows.append({"name": sched.employees[n].name, "da": da, **cnt})
    return rows


def coverage_table(sched: Schedule) -> List[dict]:
    """Pro Halbtag: geplante Aufgaben mit Ist/Soll."""
    out = []
    for d in range(5):
        for s in range(2):
            items = []
            for t in auto_tasks(sched.cfg):
                if s not in DAUER_SLOTS[t["dauer"]]:
                    continue
                need = t["anzahl"][DAY_KEYS[d]]
                have = len(sched.holders(t["code"], d, s))
                if need or have:
                    items.append({"code": t["code"], "have": have, "need": need, "ok": have >= need})
            out.append({"day": DAYS[d], "slot": SLOTS[s], "items": items})
    return out


def validate_schedule(sched: Schedule) -> List[str]:
    """Prüft den Plan (auch nach manuellen Änderungen).
    ❌ = Besetzung fehlt / nicht erlaubt · ⚠️ = Soll-Regel · ℹ️ = Info."""
    cfg = sched.cfg
    issues: List[str] = list(sched.conflicts)
    emps = sched.employees
    autos = {t["code"]: t for t in auto_tasks(cfg)}
    fill = cfg["regeln"].get("auffuellen")
    for d in range(5):
        for s in range(2):
            for code, t in autos.items():
                if s not in DAUER_SLOTS[t["dauer"]]:
                    continue
                need = t["anzahl"][DAY_KEYS[d]]
                have = len(sched.holders(code, d, s))
                if have < need:
                    issues.append(f"❌ {DAYS[d]} {SLOTS[s]}: {code} {have}/{need} besetzt")
    for name, emp in emps.items():
        for d in range(5):
            vm, nm = sched.get_task(name, d, 0), sched.get_task(name, d, 1)
            bvm, bnm = base_code(vm, cfg), base_code(nm, cfg)
            for code, t in autos.items():
                if t["dauer"] != "ganzer Tag" or (bvm == code) == (bnm == code):
                    continue
                other = 1 if bvm == code else 0
                if sched.is_present(name, d, other):
                    issues.append(f"⚠️ {emp.name} {DAYS[d]}: {code} nur am {SLOT_KEYS[1 - other]} – "
                                  f"ist eine Ganztags-Aufgabe (aufgeteilt)")
            for s, task in ((0, vm), (1, nm)):
                where = f"{emp.name} {DAYS[d]} {SLOTS[s]}"
                b = base_code(task, cfg)
                if b in autos and emp.level(b) == "–":
                    issues.append(f"❌ {where}: {b} nicht erlaubt (Wer darf was)")
                if task and not is_absent_code(task) and not emp.is_available(d, s):
                    issues.append(f"⚠️ {where}: {task} eingetragen, laut Anwesenheit aber nicht da")
                if fill and not task and emp.is_available(d, s) and (name, d, s) not in sched.notes:
                    issues.append(f"⚠️ {where}: ohne Aufgabe")
    return issues


def check_rule_impact(absences: List[Tuple[str, int, int]], employees: Dict[str, Employee],
                      cfg: dict) -> List[str]:
    """Warnungen VOR dem Planen: betroffene feste Einsätze und knappe Besetzung."""
    cfg = normalize_config(cfg)
    warnings: List[str] = []
    for name, day, slot in absences:
        emp = employees.get(name)
        if not emp:
            continue
        for d, s, task, _z in emp.fix:
            if d == day and s == slot and task:
                warnings.append(f"{emp.name} {DAYS[day]} {SLOTS[slot]}: fester Einsatz {task} fällt aus")
    autos = {t["code"] for t in auto_tasks(cfg)}
    for day in range(5):
        for slot in range(2):
            absent = {n for n, d, s in absences if d == day and s == slot}
            # wer an diesem Halbtag einen festen Einsatz ausserhalb der geplanten Aufgaben hat, fällt weg
            busy = {e.short for e in employees.values()
                    if any(d == day and s == slot and task and base_code(task, cfg) not in autos
                           for d, s, task, _z in e.fix)}
            available = sum(1 for e in employees.values()
                            if e.is_available(day, slot) and e.short not in absent | busy)
            needed = sum(t["anzahl"][DAY_KEYS[day]] for t in auto_tasks(cfg) if slot in DAUER_SLOTS[t["dauer"]])
            if available < needed:
                warnings.append(f"KRITISCH {DAYS[day]} {SLOTS[slot]}: nur {available} Personen da – "
                                f"{needed} Aufgaben zu besetzen")
            elif available == needed:
                warnings.append(f"Knapp: {DAYS[day]} {SLOTS[slot]} – {available} Personen für {needed} Aufgaben, "
                                f"niemand fürs Auffüllen")
    return warnings


# ============================================================
# EXCEL – Block-Format wie im KSV-Jahresplan (direkt einfügbar)
# ============================================================

def write_excel(sched: Schedule, output) -> None:
    """Schreibt den Plan nach `output` (Dateipfad oder BytesIO).
    Spalten A–L entsprechen genau einem Wochenblock im Jahresplan (Woche | Name |
    Vorname | Mo VM … Fr NM) und lassen sich dort direkt einfügen."""
    cfg = sched.cfg
    wb = Workbook()
    ws = wb.active
    ws.title = f"KW{sched.week_number}"
    start = sched.week_start_date
    end = start + timedelta(days=4)
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base_font = dict(name="Calibri", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def task_cell(row, col, task):
        c = ws.cell(row=row, column=col, value=task or None)
        c.border = border
        c.alignment = center
        color = get_task_color(task, cfg)
        c.font = Font(**base_font, color="FFFFFF" if color and is_dark(color) else "000000")
        if color:
            c.fill = PatternFill("solid", fgColor=color)
        return c

    ws.cell(row=1, column=1, value=f"Woche {sched.week_number} | {start.strftime('%d.%m.%Y')} – "
                                   f"{end.strftime('%d.%m.%Y')}").font = Font(**base_font, bold=True)
    for col, val in ((1, "Name"), (2, "Vorname")):
        c = ws.cell(row=2, column=col, value=val)
        c.font = Font(**base_font, bold=True)
        c.border = border
        ws.cell(row=3, column=col).border = border
    for d in range(5):
        col = 3 + 2 * d
        c = ws.cell(row=2, column=col, value=f"{DAYS[d]}\n{(start + timedelta(days=d)).strftime('%d.%m.%Y')}")
        c.font = Font(**base_font, bold=True)
        c.alignment = center
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        for s in range(2):
            ws.cell(row=2, column=col + s).border = border
            h = ws.cell(row=3, column=col + s, value=SLOTS[s])
            h.font = Font(**base_font, bold=True)
            h.alignment = center
            h.border = border
    bem = ws.cell(row=3, column=13, value="Bemerkungen")
    bem.font = Font(**base_font, bold=True)
    ws.row_dimensions[2].height = 32

    row = 4
    for m in cfg["mitarbeiter"]:
        short = m["short"]
        if short not in sched.employees:
            continue
        emp = sched.employees[short]
        vorname, initial = split_name(emp.name)
        for col, val in ((1, initial), (2, vorname)):
            c = ws.cell(row=row, column=col, value=val or None)
            c.font = Font(**base_font)
            c.fill = PatternFill("solid", fgColor=emp.farbe)
            c.border = border
        for d in range(5):
            for s in range(2):
                task_cell(row, 3 + 2 * d + s, sched.get_task(short, d, s))
        notes = [f"{DAY_KEYS[d]} {SLOT_KEYS[s]}: {sched.notes[(short, d, s)]}"
                 for d in range(5) for s in range(2) if (short, d, s) in sched.notes]
        if notes:
            c = ws.cell(row=row, column=13, value="\n".join(notes))
            c.font = Font(**base_font)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    if cfg["backup_team"]:
        c = ws.cell(row=row, column=1, value="BackUp Team")
        c.font = Font(**base_font, bold=True)
        row += 1
        for b in cfg["backup_team"]:
            c = ws.cell(row=row, column=2, value=b.get("name") or b["short"])
            c.font = Font(**base_font)
            c.border = border
            ws.cell(row=row, column=1).border = border
            for d in range(5):
                for s in range(2):
                    task_cell(row, 3 + 2 * d + s, sched.get_backup_task(b["short"], d, s))
            row += 1

    # Legende unterhalb (ausserhalb des Wochenblocks)
    row += 2
    ws.cell(row=row, column=1, value="Legende").font = Font(**base_font, bold=True)
    row += 1
    items = legend_items(cfg)
    per_col = (len(items) + 2) // 3
    for i, (code, expl, color) in enumerate(items):
        r, c0 = row + i % per_col, 3 + 4 * (i // per_col)
        c = ws.cell(row=r, column=c0, value=code)
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(**base_font, color="FFFFFF" if is_dark(color) else "000000")
        c.alignment = center
        c.border = border
        ws.cell(row=r, column=c0 + 1, value=expl).font = Font(name="Calibri", size=10)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    for i in range(10):
        ws.column_dimensions[get_column_letter(3 + i)].width = 14
    ws.column_dimensions["M"].width = 28
    ws.freeze_panes = "C4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(output)


# ============================================================
# KOMMANDOZEILE
# ============================================================

def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Teamplaner (KSV) – Plan für eine KW erzeugen")
    ap.add_argument("--montag", help="Montag der Zielwoche, z.B. 2026-10-19 (Standard: nächste Woche)")
    ap.add_argument("--config", help="Pfad zur Konfiguration (Standard: team_ksv.json)")
    ap.add_argument("--out", help="Excel-Datei (Standard: Arbeitsplan_KSV_KWxx_YYYYMMDD.xlsx)")
    args = ap.parse_args(argv)
    monday = datetime.strptime(args.montag, "%Y-%m-%d") if args.montag else get_next_monday()
    kw = monday.isocalendar()[1]
    cfg = load_config(args.config)
    sched = build_schedule(kw, monday, cfg=cfg)
    issues = validate_schedule(sched)
    out = args.out or f"Arbeitsplan_{cfg.get('team') or 'Team'}_KW{kw}_{monday.strftime('%Y%m%d')}.xlsx"
    write_excel(sched, out)
    print(f"KW {kw}  {monday.strftime('%d.%m.%Y')} – {(monday + timedelta(days=4)).strftime('%d.%m.%Y')}")
    print(f"Datei: {out}")
    for code, cells in day_overview(sched):
        print(f"{code:10s}", " | ".join(f"{DAY_KEYS[d]} {c}" for d, c in enumerate(cells)))
    if issues:
        print(f"\n{len(issues)} Hinweis(e):")
        for i in issues:
            print("  ", i)
    else:
        print("\n✅ Alle Regeln eingehalten.")
    return 0 if not any(i.startswith("❌") for i in issues) else 1


if __name__ == "__main__":
    raise SystemExit(main())
