#!/usr/bin/env python3
"""
Arbeitsplan KSC · KSV – Streamlit-App (Version 8.2, Oktober 2026)
================================================================
Eine App für zwei Teams mit denselben Funktionen. Oben links wird das Team
gewählt (KSC / KSV); jedes Team hat eigene Stammdaten, Regeln und
Abwesenheiten:

- Schritt 1 «Planung»: Abwesenheiten und Sonderwünsche erfassen (auch über
  mehrere Wochen).
- Schritt 2 «Ergebnis»: Wochenplan mit Regelprüfung, Tagesübersicht,
  Plan-Editor und Excel-Download.
- «📆 Jahresplan»: Ferien und freie Tage von–bis eintragen, Feiertage
  festlegen, Übersicht übers ganze Jahr (wer fehlt wann, geht der Plan auf?).
  Alles fliesst automatisch in jeden Wochenplan ein.
- «Team & Regeln»: Personen, Wer darf was, feste Einsätze und Regeln – ohne
  Code, gespeichert in team.json (KSC) bzw. team_ksv.json (KSV).

Scheduler: KSC → arbeitskalender.py (eingebaute Spezialregeln),
           KSV → teamplaner.py (Aufgaben und Besetzung frei konfigurierbar).

Direktlink: …/?team=KSV öffnet das KSV-Team. Mit ARBEITSPLAN_TEAM=KSV
(Umgebungsvariable bzw. Secret auf oberster Ebene) ist die App fest auf ein
Team eingestellt und der Umschalter verschwindet – so lässt sich dasselbe
Repo als eigene App pro Team deployen.

Speichern: lokal; optional direkt ins GitHub-Repo (Secrets [github]
token/repo/branch) – damit überleben Änderungen Neustarts von Streamlit Cloud.

Start lokal:
    pip install -r requirements.txt
    streamlit run streamlit_app.py
"""

from __future__ import annotations

import base64
import copy
import hashlib
import inspect
import json
import os
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import ModuleType

import pandas as pd
import streamlit as st

BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import arbeitskalender as ak  # noqa: E402
import feiertage as ft  # noqa: E402
import teamplaner as tp  # noqa: E402
from arbeitskalender import (  # noqa: E402
    DAYS, DAY_KEYS, HALBTAG_WERTE, NUR_MANUELL, ROLLEN, SCHALTER_CODES,
    SCHALTER_RESERVE_MODI, SLOTS, SLOT_KEYS, TPA_TEL_MODI,
)


# ══════════════════════════════════════════════════════════════════════
#  Teams
# ══════════════════════════════════════════════════════════════════════

@dataclass(frozen=True)
class Team:
    key: str
    untertitel: str
    engine: ModuleType
    config_path: Path
    absence_path: Path
    farbe: str            # Akzentfarbe im Kopf
    farbe_soft: str
    gruende: tuple        # Arten von Einträgen in Schritt 1 (Schlüssel aus GRUENDE)
    codes: dict           # Art → Code im Plan


# Arten von Einträgen in Schritt 1 – der Schlüssel wird so in abwesenheiten*.json gespeichert
GRUENDE = {
    "Krank": ("🤒 Krank", "Halbtag blockiert, zählt als abwesend"),
    "Ferien": ("🏖️ Ferien", "Halbtag blockiert, zählt als abwesend"),
    "Frei": ("🌙 Frei", "freier Halbtag (z. B. Kompensation), zählt als abwesend"),
    "Termin": ("🕑 Termin", "Halbtag bleibt leer, die Notiz steht im Plan"),
    "Home Office": ("🏠 Home Office", "«HO» im Plan, keine andere Aufgabe"),
    "Custom": ("📌 Aufgabe setzen", "eine bestimmte Aufgabe fest eintragen, z. B. Schulung oder ERP"),
}
ABWESEND_ARTEN = {"Krank", "Ferien", "Frei"}

TEAMS = {
    "KSC": Team(
        key="KSC", untertitel="Wochenplanung KSC-Team", engine=ak,
        config_path=Path(os.environ.get("KSC_TEAM_CONFIG", BASE_DIR / "team.json")),
        absence_path=Path(os.environ.get("KSC_ABSENCES", BASE_DIR / "abwesenheiten.json")),
        farbe="#2563eb", farbe_soft="#dbeafe",
        gruende=("Krank", "Ferien", "Termin", "Home Office", "Custom"),
        codes={"Krank": "KRANK", "Ferien": "FERIEN", "Termin": "TERMIN", "Home Office": "HO", "Custom": "CUSTOM"},
    ),
    "KSV": Team(
        key="KSV", untertitel="Wochenplanung KSV-Team", engine=tp,
        config_path=Path(os.environ.get("KSV_TEAM_CONFIG", BASE_DIR / "team_ksv.json")),
        absence_path=Path(os.environ.get("KSV_ABSENCES", BASE_DIR / "abwesenheiten_ksv.json")),
        farbe="#059669", farbe_soft="#d1fae5",
        gruende=("Krank", "Ferien", "Frei", "Termin", "Home Office", "Custom"),
        codes={"Krank": "KRANK", "Ferien": "Ferien", "Frei": "Frei", "Termin": "TERMIN", "Home Office": "HO",
               "Custom": "CUSTOM"},
    ),
}

# Titel in Browser-Tab und Kopfzeile («Arbeitsplan KSC» / «Arbeitsplan KSV») – hier ändern
APP_TITEL = "Arbeitsplan"

# Auswahl im KSC-Plan-Editor (leer = Zelle frei)
TASK_OPTIONS = [
    "", "TEL", "ABKL",
    "ERF7", "ERF7/Q", "ERF7/HUB", "ERF7/SCH",
    "ERF8", "ERF8/Q", "ERF8/HUB",
    "ERF9", "ERF9/Q", "ERF9/TEL",
    "ERF4/SCH", "ERF5",
    "PO", "PO/SCAN", "PO/ABKL", "PO/TEL", "scanning", "SCAN",
    "HO", "HO/Q", "TL/Q",
    "KGS", "KTG", "KSC Spez.", "TAGES PA", "TAGES PA/TEL",
    "RX Abo", "VBZ/Q", "ERP", "KSV",
    "KRANK", "FERIEN",
]

# Team festlegen: fest (ARBEITSPLAN_TEAM) › Link ?team=KSV › zuletzt gewählt › KSC
FIXED_TEAM = os.environ.get("ARBEITSPLAN_TEAM", "").strip().upper()
FIXED_TEAM = FIXED_TEAM if FIXED_TEAM in TEAMS else ""
# Neuere Streamlit-Versionen halten den Umschalter selbst mit der Adresse (?team=KSV) synchron;
# sonst wird die Adresse von Hand nachgeführt (erscheint dann erst bei der nächsten Aktion).
TEAM_BIND = "bind" in inspect.signature(st.segmented_control).parameters
if FIXED_TEAM:
    st.session_state.team = FIXED_TEAM
elif st.session_state.get("team") not in TEAMS:   # erster Aufruf oder Umschalter abgewählt
    _qp = str(st.query_params.get("team", "")).strip().upper()
    st.session_state.team = _qp if _qp in TEAMS else (st.session_state.get("team_last") or "KSC")
st.session_state.team_last = st.session_state.team
TEAM: str = st.session_state.team
T: Team = TEAMS[TEAM]
E: ModuleType = T.engine
IS_KSC = TEAM == "KSC"


def wk(name: str) -> str:
    """Widget-Schlüssel pro Team – so behält jedes Team seine Eingaben beim Umschalten."""
    return f"{TEAM}·{name}"


# ══════════════════════════════════════════════════════════════════════
#  Streamlit-Setup & CSS
# ══════════════════════════════════════════════════════════════════════

st.set_page_config(page_title=f"{APP_TITEL} {TEAM}", page_icon="📅", layout="wide",
                   initial_sidebar_state="collapsed")

CUSTOM_CSS = """
<style>
:root {
  --bg: #f4f6fb; --surface: #ffffff; --surface-2: #f8fafc;
  --border: #e3e8f0; --border-2: #cbd5e1;
  --text: #0f172a; --text-dim: #475569; --text-mute: #94a3b8; --heading: #0b1020;
  --primary: #0f172a; --primary-2: #1e293b;
  --accent: #2563eb; --accent-soft: #dbeafe;
  --ok: #16a34a; --ok-soft: #dcfce7; --warn: #d97706; --warn-soft: #fef3c7;
  --err: #dc2626; --err-soft: #fee2e2; --info-soft: #eff6ff; --phc: #00008B;
}
html, body, [data-testid="stAppViewContainer"] { background: var(--bg); color: var(--text); }
[data-testid="stHeader"] { background: transparent; }
[data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] { display: none; }
.block-container { padding-top: 1.2rem !important; padding-bottom: 2rem !important; max-width: 1500px; }

.st-key-topbar { background: var(--surface); border: 1px solid var(--border); border-top: 4px solid var(--team);
  border-radius: 12px; padding: 10px 14px 6px; margin-bottom: 4px; box-shadow: 0 1px 2px rgba(15,23,42,.04); }
.st-key-stepbar { border-top: 1px solid var(--border); margin-top: 4px; padding-top: 8px; }
.brand { display: flex; align-items: center; gap: 12px; }
.brand-icon { width: 40px; height: 40px; border-radius: 10px; background: var(--team); color: white;
  display: grid; place-items: center; font-size: 19px; flex-shrink: 0; }
.brand-title { font-size: 16px; font-weight: 750; color: var(--heading); line-height: 1.15; white-space: nowrap; }
.brand-title .team { color: var(--team); }
.brand-sub { font-size: 10.5px; color: var(--text-mute); letter-spacing: .6px; text-transform: uppercase; white-space: nowrap; }
.team-fixed { display: inline-block; padding: 6px 14px; border-radius: 999px; background: var(--team-soft); color: var(--team);
  font-weight: 800; font-size: 13px; letter-spacing: .5px; }

.st-key-week_offset [data-baseweb="select"] > div, .st-key-week_offset [role="group"] {
  background: var(--surface-2); border-radius: 10px; cursor: pointer; }
.st-key-week_offset input, .st-key-week_offset [data-baseweb="select"] { font-weight: 700; color: var(--heading);
  font-size: 13.5px; text-overflow: ellipsis; }

div.stButton button[kind="primary"], div.stFormSubmitButton button[kind="primaryFormSubmit"] { background: var(--primary) !important; border: 1px solid var(--primary) !important;
  color: white !important; font-weight: 600 !important; border-radius: 9px !important; }
div.stButton button[kind="primary"]:hover, div.stFormSubmitButton button[kind="primaryFormSubmit"]:hover { background: var(--primary-2) !important; border-color: var(--primary-2) !important; }
div.stButton button[kind="secondary"] { background: var(--surface) !important; border: 1px solid var(--border) !important;
  color: var(--text) !important; border-radius: 8px !important; font-weight: 500 !important; }
div.stButton button[kind="secondary"]:hover { background: var(--surface-2) !important; border-color: var(--border-2) !important; }
div.stButton > button, div.stButton > button > div, div.stButton > button p { white-space: nowrap !important; }
div.stButton button:disabled { opacity: .45; cursor: not-allowed; }
div.stButton > button { padding-left: 6px !important; padding-right: 6px !important; }
div.stDownloadButton button { background: var(--ok) !important; border: 1px solid var(--ok) !important;
  color: white !important; border-radius: 9px !important; font-weight: 600 !important; }
div.stDownloadButton button p { color: white !important; }
div.stDownloadButton button:hover { filter: brightness(1.08); }

.stats-strip { background: var(--surface); border: 1px solid var(--border); border-radius: 12px; padding: 10px 16px;
  margin: 8px 0 14px; display: flex; gap: 22px; align-items: center; flex-wrap: wrap; box-shadow: 0 1px 2px rgba(15,23,42,.04); }
.stat-item .lbl { font-size: 10px; color: var(--text-mute); text-transform: uppercase; letter-spacing: .8px; font-weight: 600; margin-bottom: 2px; }
.stat-item .val { font-size: 15px; font-weight: 700; color: var(--heading); font-feature-settings: "tnum"; }
.stat-item .val small { font-size: 11.5px; font-weight: 500; color: var(--text-dim); margin-left: 4px; }
.stat-divider { height: 28px; width: 1px; background: var(--border); }
.info-pill { margin-left: auto; display: flex; align-items: center; gap: 8px; background: var(--info-soft);
  border: 1px solid #bfdbfe; color: #1e3a8a; padding: 6px 14px; border-radius: 999px; font-size: 12.5px; }
.info-pill .ico { width: 16px; height: 16px; border-radius: 50%; background: #2563eb; color: white;
  display: grid; place-items: center; font-size: 10px; font-weight: 700; flex-shrink: 0; }

.howto { display: flex; flex-wrap: wrap; align-items: center; gap: 8px; margin: 0 0 12px; font-size: 13px; color: var(--text-dim); }
.howto-step { background: var(--surface); border: 1px solid var(--border); border-radius: 999px; padding: 6px 12px 6px 6px; }
.howto-step b { display: inline-grid; place-items: center; width: 20px; height: 20px; border-radius: 50%; background: var(--team);
  color: white; font-size: 11px; margin-right: 7px; }
.howto-arrow { color: var(--text-mute); }

.card { background: var(--surface); border: 1px solid var(--border); border-radius: 12px; padding: 18px 22px;
  margin-bottom: 16px; box-shadow: 0 1px 2px rgba(15,23,42,.04); }
.card h2 { font-size: 15.5px; font-weight: 700; color: var(--heading); margin: 0 0 4px; }
.card .sub { font-size: 12.5px; color: var(--text-mute); margin-bottom: 14px; }
.card-head { display: flex; align-items: center; justify-content: space-between; margin-bottom: 12px; }
.field-label { font-size: 10.5px; color: var(--text-mute); text-transform: uppercase; letter-spacing: .8px; font-weight: 700; margin: 8px 0 4px 2px; }
.field-label span { font-weight: 500; text-transform: none; letter-spacing: 0; }
.preview-line { background: var(--surface-2); border: 1px dashed var(--border-2); border-radius: 9px; padding: 8px 12px;
  font-size: 13px; color: var(--text-dim); margin: 6px 0 2px; }

.empty-state { padding: 34px 18px; text-align: center; color: var(--text-mute); font-size: 13px; }
.empty-state .icon { width: 44px; height: 44px; margin: 0 auto 10px; border-radius: 12px; background: var(--surface-2);
  border: 1px solid var(--border); display: grid; place-items: center; color: var(--text-mute); font-size: 20px; }
.empty-state .ttl { color: var(--text); font-weight: 600; margin-bottom: 4px; font-size: 13px; }

.entry-row { background: var(--surface); border: 1px solid var(--border); border-radius: 9px; padding: 9px 12px;
  display: flex; align-items: center; gap: 10px; font-size: 13px; flex-wrap: wrap; }
.entry-dot { width: 8px; height: 8px; border-radius: 50%; flex-shrink: 0; }
.entry-dot.Krank { background: var(--err); } .entry-dot.Ferien { background: var(--warn); } .entry-dot.Frei { background: #0ea5e9; }
.entry-dot.Termin { background: #06b6d4; } .entry-dot.HO { background: var(--accent); } .entry-dot.Custom { background: var(--text-mute); }
.entry-name { font-weight: 600; color: var(--text); }
.entry-meta { color: var(--text-dim); font-size: 12.5px; margin-left: 4px; }
.entry-note { color: #b45309; font-style: italic; margin-left: 6px; font-size: 12px; }
.group-label { font-size: 10.5px; color: var(--text-mute); text-transform: uppercase; letter-spacing: .8px; font-weight: 700; margin: 10px 0 6px 2px; }

.banner { display: flex; align-items: flex-start; gap: 10px; padding: 12px 16px; border-radius: 10px; margin-bottom: 14px; border: 1px solid; font-size: 13.5px; }
.banner .check { width: 22px; height: 22px; border-radius: 50%; display: grid; place-items: center; flex-shrink: 0; font-weight: 700; font-size: 13px; }
.banner-ok { background: var(--ok-soft); border-color: rgba(22,163,74,.3); color: #14532d; } .banner-ok .check { background: var(--ok); color: white; }
.banner-warn { background: var(--warn-soft); border-color: rgba(217,119,6,.3); color: #78350f; } .banner-warn .check { background: var(--warn); color: white; }
.banner-error { background: var(--err-soft); border-color: rgba(220,38,38,.3); color: #7f1d1d; } .banner-error .check { background: var(--err); color: white; }
.banner-info { background: var(--info-soft); border-color: #bfdbfe; color: #1e3a8a; } .banner-info .check { background: #2563eb; color: white; }
.banner h4 { margin: 0 0 4px; font-size: 13px; font-weight: 700; }
.banner ul { margin: 4px 0 0; padding-left: 18px; font-size: 12.5px; }

.plan-wrap { overflow-x: auto; border-radius: 10px; border: 1px solid var(--border); background: var(--surface); }
table.plan { width: 100%; border-collapse: separate; border-spacing: 0; font-size: 11.5px;
  font-family: -apple-system, "Segoe UI", "Inter", sans-serif; min-width: 1100px; background: var(--surface); }
table.plan th, table.plan td { padding: 6px 6px; border-right: 1px solid var(--border); border-bottom: 1px solid var(--border); vertical-align: middle; }
table.plan th { background: var(--surface-2); color: var(--text-dim); font-weight: 700; font-size: 10.5px; letter-spacing: .6px; text-transform: uppercase; text-align: center; }
table.plan th.team-head { text-align: left; padding: 10px 14px; vertical-align: top; background: var(--surface); }
table.plan th.team-head .legend { margin-top: 4px; color: var(--text-mute); font-size: 10.5px; font-weight: 500; text-transform: none; letter-spacing: 0; }
table.plan th.day-head { background: var(--surface); padding: 8px 8px; }
table.plan th.day-head .day-name { display: block; color: var(--text-dim); font-weight: 700; font-size: 11px; letter-spacing: .8px; text-transform: uppercase; }
table.plan th.day-head .day-date { display: block; color: var(--text-mute); font-weight: 600; font-size: 10.5px; letter-spacing: 0; text-transform: none; margin-top: 1px; }
.tv-pill { display: inline-flex; align-items: center; gap: 6px; margin-top: 6px; padding: 3px 9px; background: var(--surface-2);
  border: 1px solid var(--border); border-radius: 999px; font-size: 11px; font-weight: 600; color: var(--text); text-transform: none; letter-spacing: 0; }
.tv-tag { font-size: 9.5px; font-weight: 800; padding: 1px 5px; border-radius: 4px; letter-spacing: .4px; }
.tv-pill.open .tv-name { color: var(--warn); }
table.plan th.slot-head { background: var(--surface-2); font-size: 9.5px; padding: 5px 4px; color: var(--text-mute); font-weight: 600; }
table.plan td.name-cell { background: var(--surface); padding: 6px 12px; text-align: left; font-size: 12px; color: var(--text); font-weight: 500; white-space: nowrap; position: sticky; left: 0; }
table.plan td.pct-cell { background: var(--surface); text-align: right; padding-right: 14px; color: var(--text-mute); font-size: 11px; font-weight: 600; }
table.plan tbody tr:nth-child(even) td.name-cell, table.plan tbody tr:nth-child(even) td.pct-cell, table.plan tbody tr:nth-child(even) td.task-cell { background: var(--surface-2); }
td.task-cell { text-align: center; padding: 5px 4px; background: var(--surface); }
table.plan td.sep-cell { background: #e2e8f0 !important; color: var(--text-dim); font-weight: 700; font-size: 10.5px; letter-spacing: .6px; text-transform: uppercase; padding: 6px 12px; text-align: left; }
.swatch { display: inline-block; width: 10px; height: 10px; border-radius: 3px; margin-right: 8px; vertical-align: -1px; border: 1px solid rgba(15,23,42,.12); }
.pill { display: inline-block; padding: 4px 10px; border-radius: 999px; font-size: 11px; font-weight: 600; letter-spacing: .2px; min-width: 64px; text-align: center; }
.pill-phc { background: var(--phc) !important; color: white !important; font-weight: 700; }
.pill-termin { background: #fef3c7; color: #78350f; }
.pill-note { display:block; font-size: 9.5px; color: var(--text-mute); margin-top: 2px; }

table.roles { width: 100%; border-collapse: collapse; font-size: 12px; margin-top: 14px; }
table.roles th, table.roles td { border: 1px solid var(--border); padding: 6px 10px; text-align: center; }
table.roles th { background: var(--surface-2); color: var(--text-dim); font-size: 10.5px; text-transform: uppercase; letter-spacing: .6px; }
table.roles td.lbl { text-align: left; font-weight: 600; color: var(--text); background: var(--surface-2); white-space: nowrap; }
table.roles td.open { color: var(--warn); font-weight: 700; }
.section-title { font-size: 12px; font-weight: 700; color: var(--text-dim); text-transform: uppercase; letter-spacing: .7px; margin: 18px 0 0; }

.wa-strip { display: grid; grid-template-columns: repeat(3, 1fr); gap: 12px; margin-top: 14px; }
.wa-card { background: var(--surface); border: 1px solid var(--border); border-radius: 12px; padding: 12px 16px; display: flex; align-items: center; gap: 12px; }
.wa-icon { width: 32px; height: 32px; border-radius: 8px; background: var(--surface-2); border: 1px solid var(--border); display: grid; place-items: center; font-size: 14px; }
.wa-meta .lbl { font-size: 10px; color: var(--text-mute); text-transform: uppercase; letter-spacing: .8px; font-weight: 700; }
.wa-meta .val { font-size: 14px; font-weight: 700; color: var(--heading); margin-top: 2px; }
.legend-wrap { display: flex; flex-wrap: wrap; gap: 8px 14px; }
.legend-item { display: inline-flex; align-items: center; gap: 6px; font-size: 12px; color: var(--text-dim); }
.element-container { margin-bottom: 0.4rem; }

.page-head h2 { font-size: 18px; font-weight: 700; color: var(--heading); margin: 0 0 2px; padding: 0; }
.page-head .sub { font-size: 12.5px; color: var(--text-mute); margin-bottom: 10px; }
.st-key-savebar_ok { background: var(--ok-soft); border-color: rgba(22,163,74,.3) !important; color: #14532d; }
.st-key-savebar_dirty { background: var(--warn-soft); border-color: rgba(217,119,6,.35) !important; color: #78350f; }
.st-key-savebar_ok p, .st-key-savebar_dirty p { margin-bottom: 0; font-size: 13.5px; }
.team-hint { background: var(--surface); border: 1px solid var(--border); border-radius: 10px; padding: 9px 14px;
  font-size: 13px; color: var(--text-dim); }
.explain { background: var(--surface); border: 1px solid var(--border); border-left: 4px solid var(--team); border-radius: 10px;
  padding: 10px 14px; font-size: 13px; color: var(--text-dim); margin-bottom: 10px; }
.explain ol { margin: 4px 0 0; padding-left: 20px; }

.entry-dot.Feiertag { background: #a855f7; }
.entry-tag { margin-left: auto; font-size: 13px; cursor: help; }
.fe-row { background: #faf5ff; border-color: #e9d5ff; margin-bottom: 4px; }
.fe-pill { display: inline-block; margin-top: 6px; padding: 3px 9px; background: #f3e8ff; border: 1px solid #e9d5ff;
  border-radius: 999px; font-size: 11px; font-weight: 600; color: #6b21a8; text-transform: none; letter-spacing: 0; }

.jp-wrap { overflow-x: auto; border: 1px solid var(--border); border-radius: 10px; background: var(--surface); }
.jp-wrap.rtl { direction: rtl; }          /* schmaler Bildschirm, zweite Jahreshälfte: rechts beginnen */
.jp-wrap.rtl table.jp { direction: ltr; }
table.jp { border-collapse: separate; border-spacing: 0; font-size: 11px; width: 100%; }
table.jp th, table.jp td { border-bottom: 1px solid var(--border); text-align: center; padding: 2px 0; min-width: 17px;
  height: 24px; white-space: nowrap; }
table.jp th { color: var(--text-mute); font-weight: 600; font-size: 9.5px; background: var(--surface-2); }
table.jp th.jp-month { color: var(--text-dim); font-weight: 700; font-size: 10px; text-transform: uppercase;
  letter-spacing: .6px; text-align: left; padding-left: 4px; }
table.jp .ms { border-left: 1px solid var(--border-2); }
table.jp .jp-name { text-align: left; padding: 2px 10px; position: sticky; left: 0; background: var(--surface);
  font-weight: 600; color: var(--text); z-index: 1; min-width: 108px; border-right: 1px solid var(--border); }
table.jp thead .jp-name { background: var(--surface-2); color: var(--text-mute); font-size: 10px; }
table.jp tr.jp-fe td { background: #fdfaff; }
table.jp tr.jp-status td { border-top: 2px solid var(--border); border-bottom: none; }
table.jp .past { opacity: .35; }
table.jp th.sel, table.jp td.sel { background: var(--team-soft) !important; }
table.jp th.sel { color: var(--team); font-weight: 800; }
.jp-cell { display: inline-block; min-width: 16px; padding: 2px 1px; border-radius: 4px; font-weight: 700; font-size: 9.5px;
  line-height: 1.3; letter-spacing: -.2px; }
.jp-dot { display: inline-block; width: 10px; height: 10px; border-radius: 50%; }
.jp-legend { display: flex; flex-wrap: wrap; gap: 6px 16px; margin-top: 10px; font-size: 12px; color: var(--text-dim); align-items: center; }
.jp-legend .jp-cell { min-width: 22px; text-align: center; }
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)
st.markdown(f"<style>:root {{ --team: {T.farbe}; --team-soft: {T.farbe_soft}; }}</style>", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════
#  Hilfsfunktionen
# ══════════════════════════════════════════════════════════════════════

def html_escape(s) -> str:
    if s is None:
        return ""
    return (str(s).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def _fg_for(bg_hex: str) -> str:
    h = bg_hex.lstrip("#")
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return "#1F1F1F"
    lum = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "#1F1F1F" if lum > 0.6 else "#FFFFFF"


def task_color(task: str) -> str | None:
    """Farbe wie im Excel des aktuellen Teams."""
    return ak.get_task_color(task) if IS_KSC else tp.get_task_color(task, ts["cfg"])


def pill_for_task(task: str) -> tuple[str, str]:
    if not task:
        return "transparent", "transparent"
    if task.startswith("*"):
        return "#fef3c7", "#78350f"
    color = task_color(task)
    if color:
        return f"#{color}", _fg_for(color)
    return "#f1f5f9", "#334155"


def _b(v) -> bool:
    return bool(v) and not (isinstance(v, float) and pd.isna(v))


def _s(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def farb_label(hex_color: str) -> str:
    return tp.farbname(hex_color)


def farb_hex(label: str, fallback: str = "D9D9D9") -> str:
    label = _s(label)
    return tp.FARBEN.get(label) or tp.norm_hex(label, fallback)


def farb_options(*hexes: str) -> list[str]:
    opts = list(tp.FARBEN)
    for h in hexes:
        if farb_label(h) not in opts:
            opts.append(farb_label(h))
    return opts


def team_task_options(cfg: dict) -> list[str]:
    """Codes für Auswahllisten (Plan-Editor, «Aufgabe setzen»)."""
    return TASK_OPTIONS if IS_KSC else tp.editor_options(cfg)


def disp(short: str) -> str:
    """Anzeigename zum internen Kurznamen."""
    cfg = ts["cfg"]
    for p in cfg["mitarbeiter"] + cfg["backup_team"]:
        if p["short"] == short:
            return p.get("name") or short
    return short


def person_label(short: str) -> str:
    cfg = ts["cfg"]
    for m in cfg["mitarbeiter"]:
        if m["short"] == short:
            return m.get("name") or short
    for b in cfg["backup_team"]:
        if b["short"] == short:
            return f"{b.get('name') or short} (BackUp)"
    return short


def active_members() -> list[dict]:
    return [m for m in ts["cfg"]["mitarbeiter"] if m.get("aktiv", True)]


def active_shorts() -> list[str]:
    return [m["short"] for m in active_members()]


def backup_shorts() -> list[str]:
    return [b["short"] for b in ts["cfg"]["backup_team"]]


def employee_options() -> dict[str, str]:
    """Anzeige-Label → Kurzname (aktives Team + BackUp-Team)."""
    return {person_label(s): s for s in active_shorts() + backup_shorts()}


def get_week_info(offset: int = 0) -> dict:
    monday = ak.get_next_monday() + timedelta(weeks=offset)
    iso_year, kw, _ = monday.isocalendar()
    friday = monday + timedelta(days=4)
    return {
        "monday": monday, "friday": friday, "kw": kw, "iso_year": iso_year,
        "widx": ak.week_index(monday), "is_even": kw % 2 == 0,
        "monday_str": monday.strftime("%d.%m.%Y"), "friday_str": friday.strftime("%d.%m.%Y"),
        "two_week_notes": E.two_week_notes(kw, ts["cfg"]),
        "feiertage": ft.pro_tag(ft.woche(ts["feiertage"], monday)),   # {Tag: "Weihnachten"}
    }


def entry_widx_start(entry: dict) -> int:
    jahr = int(entry.get("jahr") or date.today().isocalendar()[0])
    kw = int(entry.get("kw_start") or 1)
    try:
        return ak.week_index(date.fromisocalendar(jahr, kw, 1))
    except ValueError:
        return ak.week_index(date.fromisocalendar(jahr, 1, 1))


def entry_active(entry: dict, widx: int) -> bool:
    start = entry_widx_start(entry)
    weeks = max(1, int(entry.get("weeks", 1)))
    return start <= widx < start + weeks


def kw_range_label(von: date, bis: date) -> str:
    """'KW 44', 'KW 52–53' bzw. über den Jahreswechsel 'KW 53 – KW 1'."""
    (y0, k0, _), (y1, k1, _) = von.isocalendar(), bis.isocalendar()
    if (y0, k0) == (y1, k1):
        return f"KW {k0}"
    return f"KW {k0}–{k1}" if y0 == y1 else f"KW {k0} – KW {k1}"


def kw_label(jahr: int, kw: int, weeks: int) -> str:
    start = date.fromisocalendar(jahr, kw, 1)
    return kw_range_label(start, start + timedelta(weeks=max(1, weeks) - 1))


def entry_kw_label(entry: dict) -> str:
    return kw_label(int(entry.get("jahr") or date.today().year), int(entry.get("kw_start") or 0),
                    max(1, int(entry.get("weeks", 1))))


def days_label(days: list[int]) -> str:
    days = sorted(days)
    if len(days) >= 3 and days == list(range(days[0], days[-1] + 1)):
        return f"{DAY_KEYS[days[0]]}–{DAY_KEYS[days[-1]]}"
    return ", ".join(DAY_KEYS[d] for d in days)


def slot_label(slots: list[int]) -> str:
    return "ganzer Tag" if len(slots) == 2 else ("Vormittag" if slots[0] == 0 else "Nachmittag")


# Einträge gibt es in zwei Formen (beide in abwesenheiten*.json):
#  - Schritt 1: Wochentage ("days") ab einer KW ("jahr", "kw_start") für "weeks" Wochen
#  - Jahresplan: Zeitraum "von"–"bis" (Wochenenden und Feiertage zählen nicht);
#    "days" ist leer, jahr/kw_start/weeks beschreiben nur die betroffenen Wochen
WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def _as_date(x) -> date:
    return x.date() if isinstance(x, datetime) else x


def datum_label(d: date, mit_jahr: bool = False) -> str:
    return f"{WOCHENTAGE[d.weekday()]} {d.day}.{d.month}." + (str(d.year) if mit_jahr else "")


def range_label(von: date, bis: date) -> str:
    if von == bis:
        return datum_label(von, True)
    return f"{datum_label(von, von.year != bis.year)} – {datum_label(bis, True)}"


def entry_range(entry: dict) -> tuple[date, date] | None:
    """(von, bis) bei Einträgen aus dem Jahresplan, sonst None."""
    if not entry.get("von"):
        return None
    try:
        von = date.fromisoformat(str(entry["von"])[:10])
        bis = date.fromisoformat(str(entry.get("bis") or entry["von"])[:10])
    except ValueError:
        return None
    return von, max(von, bis)


def entry_days(entry: dict, monday) -> list[int]:
    """Welche Tage (0 = Mo … 4 = Fr) der Woche ab `monday` betrifft der Eintrag?"""
    monday = _as_date(monday)
    rng = entry_range(entry)
    if rng:
        return [d for d in range(5) if rng[0] <= monday + timedelta(days=d) <= rng[1]]
    return sorted(entry.get("days", [])) if entry_active(entry, ak.week_index(monday)) else []


def entry_span(entry: dict) -> tuple[date, date]:
    """Erster und letzter betroffener Tag – zum Sortieren, Aufräumen und für den Jahresplan."""
    rng = entry_range(entry)
    if rng:
        return rng
    days = sorted(entry.get("days", [])) or [0, 4]
    start = date(2024, 1, 1) + timedelta(weeks=entry_widx_start(entry))
    weeks = max(1, int(entry.get("weeks", 1)))
    return start + timedelta(days=days[0]), start + timedelta(weeks=weeks - 1, days=days[-1])


def entry_when_label(entry: dict) -> str:
    slots = f" ({slot_label(entry['slots'])})"
    rng = entry_range(entry)
    if rng:
        return f"{range_label(*rng)}{slots} · {kw_range_label(*rng)}"
    return f"{days_label(entry['days'])}{slots} · {entry_kw_label(entry)}"


# ══════════════════════════════════════════════════════════════════════
#  Persistenz (lokal + optional GitHub)
# ══════════════════════════════════════════════════════════════════════

def github_conf() -> dict | None:
    try:
        g = st.secrets["github"]
        token, repo = g["token"], g["repo"]
    except Exception:
        return None
    if not token or not repo:
        return None
    try:
        branch = g.get("branch", "main") or "main"
    except Exception:
        branch = "main"
    return {"token": token, "repo": repo, "branch": branch}


def github_put(rel_path: str, content: str, message: str) -> tuple[bool, str]:
    conf = github_conf()
    if not conf:
        return False, "GitHub nicht konfiguriert (Secrets [github] token/repo fehlen)"
    try:
        import requests
    except ImportError:
        return False, "Paket 'requests' fehlt"
    url = f"https://api.github.com/repos/{conf['repo']}/contents/{rel_path}"
    headers = {"Authorization": f"Bearer {conf['token']}", "Accept": "application/vnd.github+json"}
    try:
        r = requests.get(url, headers=headers, params={"ref": conf["branch"]}, timeout=20)
        sha = r.json().get("sha") if r.status_code == 200 else None
        body = {"message": message, "branch": conf["branch"],
                "content": base64.b64encode(content.encode("utf-8")).decode("ascii")}
        if sha:
            body["sha"] = sha
        r = requests.put(url, headers=headers, json=body, timeout=20)
        if r.status_code in (200, 201):
            return True, f"{rel_path} nach GitHub ({conf['repo']}@{conf['branch']}) gespeichert"
        return False, f"GitHub-Fehler {r.status_code}: {r.text[:200]}"
    except Exception as exc:  # Netzwerk etc.
        return False, f"GitHub nicht erreichbar: {exc}"


def write_local(path: Path, content: str) -> tuple[bool, str]:
    try:
        path.write_text(content, encoding="utf-8")
        return True, f"{path.name} lokal gespeichert"
    except Exception as exc:
        return False, f"{path.name} konnte nicht geschrieben werden: {exc}"


def entries_json(entries: list[dict], feiertage: dict) -> str:
    return json.dumps({"_hinweis": "Abwesenheiten, Sonderwünsche und Feiertage der App (automatisch gepflegt)",
                       "entries": entries, "feiertage": ft.normalize(feiertage)}, ensure_ascii=False, indent=2) + "\n"


def absences_json() -> str:
    return entries_json(ts["entries"], ts["feiertage"])


def persist_entries() -> None:
    """Einträge und Feiertage sofort lokal sichern – ins GitHub-Repo erst mit «Speichern»."""
    write_local(T.absence_path, absences_json())


def save_everything(push_github: bool) -> list[tuple[bool, str]]:
    msgs = []
    cfg_json = E.config_to_json(ts["cfg"])
    msgs.append(write_local(T.config_path, cfg_json))
    msgs.append(write_local(T.absence_path, absences_json()))
    if push_github:
        msgs.append(github_put(T.config_path.name, cfg_json, f"{T.config_path.name} über die App aktualisiert ({TEAM})"))
        msgs.append(github_put(T.absence_path.name, absences_json(),
                               f"Abwesenheiten und Feiertage {TEAM} über die App aktualisiert"))
    ts["cfg_saved_json"] = cfg_json
    ts["entries_saved_json"] = absences_json()
    return msgs


def load_absences(path: Path) -> tuple[list[dict], dict]:
    """Einträge und Feiertage aus abwesenheiten*.json – fehlt die Datei: keine Einträge, Standard-Feiertage."""
    data = {}
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    return list(data.get("entries", [])), ft.normalize(data.get("feiertage"))


def replace_cfg(new_cfg: dict) -> None:
    ts["cfg"] = E.normalize_config(new_cfg)
    ts["cfg_version"] += 1
    ts["result"] = None


# ══════════════════════════════════════════════════════════════════════
#  Session-State (pro Team)
# ══════════════════════════════════════════════════════════════════════

st.session_state.setdefault("teams", {})
if TEAM not in st.session_state.teams:
    try:
        _cfg = E.load_config(T.config_path)
    except FileNotFoundError as exc:
        st.error(str(exc))
        st.stop()
    _entries, _feiertage = load_absences(T.absence_path)
    st.session_state.teams[TEAM] = {
        "cfg": _cfg, "cfg_saved_json": E.config_to_json(_cfg), "cfg_version": 0,
        "entries": _entries, "feiertage": _feiertage, "entries_saved_json": entries_json(_entries, _feiertage),
        "fe_version": 0, "result": None, "sched": None,
    }
ts: dict = st.session_state.teams[TEAM]
ts.setdefault("feiertage", ft.default())   # Sitzung aus einer älteren Version
ts.setdefault("fe_version", 0)

st.session_state.setdefault("week_offset", 0)
if st.session_state.week_offset not in range(-1, 53):
    st.session_state.week_offset = min(52, max(-1, int(st.session_state.week_offset)))
st.session_state.setdefault("step", 1)
st.session_state.setdefault("save_msgs", [])
if not FIXED_TEAM and not TEAM_BIND and st.query_params.get("team") != TEAM:
    st.query_params["team"] = TEAM

cfg_dirty = E.config_to_json(ts["cfg"]) != ts["cfg_saved_json"]
entries_dirty = absences_json() != ts["entries_saved_json"]


# ══════════════════════════════════════════════════════════════════════
#  Plan-Generierung
# ══════════════════════════════════════════════════════════════════════

def build_rows(sched) -> tuple[list[dict], list[dict]]:
    cfg = ts["cfg"]
    rows, backup_rows = [], []
    phc = getattr(sched, "phc_liste", {})
    for m in cfg["mitarbeiter"]:
        short = m["short"]
        if short not in sched.employees:
            continue
        emp = sched.employees[short]
        cells = [{"task": sched.get_task(short, d, s),
                  "is_phc": s == 1 and phc.get(d) == short,
                  "note": sched.notes.get((short, d, s), "")}
                 for d in range(5) for s in range(2)]
        rows.append({"full_name": emp.name, "short": short, "pct": emp.pct, "cells": cells,
                     "farbe": None if IS_KSC else m.get("farbe")})
    for b in cfg["backup_team"]:
        cells = [{"task": sched.get_backup_task(b["short"], d, s), "is_phc": False, "note": ""}
                 for d in range(5) for s in range(2)]
        backup_rows.append({"full_name": b.get("name") or b["short"], "short": b["short"], "pct": None,
                            "cells": cells, "farbe": None})
    return rows, backup_rows


def excel_bytes(sched) -> bytes:
    buf = BytesIO()
    E.write_excel(sched, buf)
    return buf.getvalue()


def _team_details(sched) -> dict:
    """Teamspezifische Übersichten unter dem Plan."""
    if IS_KSC:
        return {"coverage": ak.coverage_table(sched)}
    return {"coverage": tp.coverage_table(sched), "overview": tp.day_overview(sched),
            "summary": tp.person_summary(sched)}


def refresh_result_from_sched(sched) -> None:
    res = ts["result"]
    res["issues"] = E.validate_schedule(sched)
    res["rows"], res["backup_rows"] = build_rows(sched)
    res["excel_bytes"] = excel_bytes(sched)
    res.update(_team_details(sched))


def week_inputs(monday) -> dict:
    """Einträge (Schritt 1 und Jahresplan) und Feiertage einer Woche → Eingaben für build_schedule.
    Feiertage gehen vor: an diesen Halbtagen zählt kein Eintrag."""
    cfg = ts["cfg"]
    employees = E.create_employees(_as_date(monday).isocalendar()[1], cfg)
    backups = set(backup_shorts())
    fe = ft.woche(ts["feiertage"], monday)
    overrides, notes, arten, details = {}, {}, {}, {}

    for entry in ts["entries"]:
        days = entry_days(entry, monday)
        name = entry.get("name", "")
        is_backup = name in backups
        if not days or (name not in employees and not is_backup):
            continue
        art = entry.get("type", "Custom")
        code = T.codes.get(art, "CUSTOM")
        note = (entry.get("note") or "").strip()
        task = (entry.get("task") or "").strip()
        if code == "TERMIN":
            label = f"*{note}" if note else "*Termin"
        elif code == "CUSTOM":
            label = task or "HO"
        else:
            label = code
        hit = False
        for day in days:
            for slot in entry["slots"]:
                if (day, slot) in fe:
                    continue
                if is_backup:
                    overrides[(name, day, slot)] = label
                    continue
                if employees[name].is_available(day, slot):
                    overrides[(name, day, slot)] = label
                    arten[(name, day, slot)] = art
                    hit = True
                    if note and code != "TERMIN" and note != task:
                        notes[(name, day, slot)] = note
        if hit:
            what = GRUENDE.get(art, (art, ""))[0] + (f" «{task}»" if art == "Custom" and task else "")
            half = "" if len(entry["slots"]) == 2 else f" {SLOT_KEYS[entry['slots'][0]]}"
            details.setdefault(name, []).append(f"{what} {days_label(days)}{half}")

    halbtage: dict[str, dict[str, int]] = {}   # Person → Art → Halbtage (für den Jahresplan)
    for (name, _d, _s), art in arten.items():
        halbtage.setdefault(name, {})
        halbtage[name][art] = halbtage[name].get(art, 0) + 1
    absences = sorted(k for k, art in arten.items()
                      if art in ABWESEND_ARTEN or str(overrides.get(k, "")).startswith("*"))
    return {"employees": employees, "overrides": overrides, "notes": notes, "absences": absences,
            "feiertage": fe, "halbtage": halbtage, "details": details}


def plan_fp(inp: dict) -> str:
    """Fingerabdruck der Eingaben einer Woche – ändert er sich, ist ein erstellter Plan veraltet."""
    return hashlib.md5(repr((sorted(inp["overrides"].items()), sorted(inp["notes"].items()),
                             sorted(inp["feiertage"].items()))).encode("utf-8")).hexdigest()


def build_week(monday) -> tuple:
    """Plan einer Woche mit allen Einträgen und Feiertagen: (sched, Warnungen, Eingaben)."""
    cfg = ts["cfg"]
    inp = week_inputs(monday)
    sched = E.build_schedule(_as_date(monday).isocalendar()[1], monday, inp["overrides"], cfg=cfg,
                             feiertage=inp["feiertage"])
    for key, n in inp["notes"].items():
        sched.notes[key] = n
    warnings = E.check_rule_impact(inp["absences"], inp["employees"], cfg) if inp["absences"] else []
    return sched, warnings, inp


def run_generation() -> None:
    cfg = ts["cfg"]
    info = get_week_info(st.session_state.week_offset)
    kw, monday = info["kw"], info["monday"]
    sched, warnings, inp = build_week(monday)

    rows, backup_rows = build_rows(sched)
    res = {
        "team": TEAM,
        "filename": f"Arbeitsplan_{'' if IS_KSC else 'KSV_'}KW{kw}_{monday.strftime('%Y%m%d')}.xlsx",
        "excel_bytes": excel_bytes(sched),
        "kw": kw, "iso_year": info["iso_year"], "monday": monday,
        "monday_str": info["monday_str"], "friday_str": info["friday_str"],
        "warnings": warnings,
        "issues": E.validate_schedule(sched),
        "rows": rows, "backup_rows": backup_rows,
        "config_problems": E.validate_config(cfg),
        "feiertage": ft.pro_tag(inp["feiertage"]), "fp": plan_fp(inp),
    }
    if IS_KSC:
        res.update({
            "tagesverantwortung": [sched.tagesverantwortung.get(d, "?") for d in range(5)],
            "phc_liste": [sched.phc_liste.get(d, "?") for d in range(5)],
            "tagespa": [sched.tagespa.get(d, "?") for d in range(5)],
            "po_liste": [sched.po_liste.get(d, "?") for d in range(5)],
            "po_zusatz": [", ".join(sched.po_zusatz.get(d, [])) or "–" for d in range(5)],
            "schalter": [sched.schalter.get(d, "–") for d in range(5)],
            "erf9": [sched.erf9.get(d, "–") for d in range(5)],
            "wochenaufgaben": {"direkt": sched.wochenaufgabe_direkt, "onb": sched.wochenaufgabe_onb,
                               "btm": sched.wochenaufgabe_btm},
        })
    res.update(_team_details(sched))
    ts["result"] = res
    ts["sched"] = sched
    st.session_state.step = 2


def generate_now() -> None:
    with st.spinner("Plan wird generiert …"):
        try:
            run_generation()
        except Exception as exc:
            st.error(f"Fehler beim Generieren: {exc}")
            return
    st.rerun()


# ══════════════════════════════════════════════════════════════════════
#  Kopfzeile
# ══════════════════════════════════════════════════════════════════════

info = get_week_info(st.session_state.week_offset)

# Plan neu erzeugen, wenn er nicht zur gewählten KW passt (Wochen-/Teamwechsel) oder sich die Einträge bzw.
# Feiertage dieser Woche seither geändert haben – sonst bleibt er stehen (auch mit Änderungen von Hand)
if st.session_state.step == 2 and ts["result"] is not None:
    _andere_kw = (ts["result"]["iso_year"], ts["result"]["kw"]) != (info["iso_year"], info["kw"])
    if _andere_kw or ts["result"].get("fp") != plan_fp(week_inputs(info["monday"])):
        try:
            run_generation()
            if not _andere_kw:
                st.toast(f"Plan KW {info['kw']} neu erstellt – Einträge oder Feiertage haben sich geändert.", icon="🔄")
        except Exception as exc:  # z. B. defekte Konfiguration – dann zurück zu Schritt 1
            ts["result"] = None
            st.error(f"Fehler beim Generieren: {exc}")
if st.session_state.step == 2 and ts["result"] is None:   # z. B. Teamwechsel zu einem Team ohne Plan
    st.session_state.step = 1


def _switch_team() -> None:
    new = st.session_state.get("team_pick")
    if not new or new == st.session_state.team:
        st.session_state.team_pick = st.session_state.team
        return
    st.session_state.team = new
    st.query_params["team"] = new   # Adresse sofort anpassen (Lesezeichen / Neuladen)


def _goto(step: int) -> None:
    st.session_state.step = step


WOCHEN_AUSWAHL = list(range(-1, 53))   # laufende Woche bis 52 Wochen im Voraus (0 = nächste Woche)


def week_label(offset: int) -> str:
    monday = ak.get_next_monday() + timedelta(weeks=offset)
    friday = monday + timedelta(days=4)
    tage = (f"{monday:%d.}–{friday:%d.%m.}" if monday.month == friday.month
            else f"{monday:%d.%m.}–{friday:%d.%m.}")
    wann = {-1: " · diese Woche", 0: " · nächste Woche"}.get(offset, "")
    return f"KW {monday.isocalendar()[1]} · {tage}{wann}"


def _week(delta: int) -> None:
    st.session_state.week_offset = min(WOCHEN_AUSWAHL[-1], max(WOCHEN_AUSWAHL[0], st.session_state.week_offset + delta))


with st.container(key="topbar"):
    # Zeile 1 – worum geht es: Team und Woche, dazu die Hauptaktion
    r1 = st.columns([2.2, 1.45, 4.55, 1.8], vertical_alignment="center")
    with r1[0]:
        st.markdown(
            f'<div class="brand"><div class="brand-icon">📅</div><div>'
            f'<div class="brand-title">{APP_TITEL} <span class="team">{TEAM}</span></div>'
            f'<div class="brand-sub">{html_escape(T.untertitel)}</div></div></div>',
            unsafe_allow_html=True,
        )
    with r1[1]:
        team_help = ("Team wechseln. Jedes Team hat eigene Personen, Regeln und Abwesenheiten. "
                     "Direktlink für ein Team: Adresse mit ?team=KSV bzw. ?team=KSC.")
        if FIXED_TEAM:
            st.markdown(f'<span class="team-fixed">Team {TEAM}</span>', unsafe_allow_html=True)
        elif TEAM_BIND:
            st.segmented_control("Team", list(TEAMS), key="team", bind="query-params",
                                 label_visibility="collapsed", help=team_help)
        else:
            st.session_state.team_pick = TEAM
            st.segmented_control("Team", list(TEAMS), key="team_pick", on_change=_switch_team,
                                 label_visibility="collapsed", help=team_help)
    with r1[2]:
        w1, w2, w3 = st.columns([0.8, 5.0, 0.8], vertical_alignment="center")
        w1.button("‹", key="wk_prev", help="Vorwoche", width="stretch", on_click=_week, args=(-1,),
                  disabled=st.session_state.week_offset <= WOCHEN_AUSWAHL[0])
        w2.selectbox("Woche", WOCHEN_AUSWAHL, key="week_offset", format_func=week_label,
                     label_visibility="collapsed",
                     help="Woche wählen: die laufende Woche oder bis zu einem Jahr im Voraus. "
                          "Danach «Plan generieren».")
        w3.button("›", key="wk_next", help="Nächste Woche", width="stretch", on_click=_week, args=(1,),
                  disabled=st.session_state.week_offset >= WOCHEN_AUSWAHL[-1])
    with r1[3]:
        if st.session_state.step == 2 and ts["result"]:
            st.download_button("⬇  Excel herunterladen", data=ts["result"]["excel_bytes"],
                               file_name=ts["result"]["filename"],
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               width="stretch", key="cta_download")
        elif st.button("→  Plan generieren", key="cta_generate", type="primary", width="stretch",
                       help=f"Plan für KW {info['kw']} erstellen – Abwesenheiten aus Schritt 1 werden berücksichtigt"):
            generate_now()
    # Zeile 2 – die Schritte wie Reiter
    with st.container(key="stepbar"):
        s1, s2, s4, s3, _ = st.columns([1.6, 1.6, 1.7, 2.1, 3.0], vertical_alignment="center")
        step = st.session_state.step
        s1.button("● 1  Planung" if step == 1 else "1  Planung", key="step1", width="stretch",
                  type="primary" if step == 1 else "secondary", on_click=_goto, args=(1,),
                  help="Abwesenheiten und Sonderwünsche für die Woche erfassen")
        s2.button("● 2  Ergebnis" if step == 2 else "2  Ergebnis", key="step2", width="stretch",
                  type="primary" if step == 2 else "secondary", disabled=ts["result"] is None,
                  on_click=_goto, args=(2,),
                  help="Wochenplan ansehen, anpassen, Excel herunterladen" if ts["result"] is not None
                  else "Zuerst «Plan generieren»")
        s4.button("📆  Jahresplan", key="step4", width="stretch", type="primary" if step == 4 else "secondary",
                  on_click=_goto, args=(4,),
                  help="Ferien und Feiertage fürs ganze Jahr – fliessen automatisch in jeden Wochenplan")
        s3.button("⚙  Team & Regeln" + ("  •" if cfg_dirty else ""), key="step3", width="stretch",
                  type="primary" if step == 3 else "secondary", on_click=_goto, args=(3,),
                  help="Personen, wer darf was, feste Einsätze und Regeln" +
                       (" – ungespeicherte Änderungen" if cfg_dirty else ""))


# ══════════════════════════════════════════════════════════════════════
#  Info-Leiste
# ══════════════════════════════════════════════════════════════════════

n_active = len(active_shorts())
n_total = len(ts["cfg"]["mitarbeiter"])
notes_str = " · ".join([f"🎉 {DAY_KEYS[d]} {t}" for d, t in info["feiertage"].items()] + info["two_week_notes"]) \
    or "keine Besonderheiten im 2-Wochen-Rhythmus"
storage = []
if cfg_dirty:
    storage.append("Team & Regeln: ungespeichert")
if entries_dirty:
    storage.append("Einträge: ungespeichert")
storage_html = (f'<div class="info-pill" style="margin-left:12px; background:var(--warn-soft); '
                f'border-color:rgba(217,119,6,.3); color:#78350f;"><span class="ico" style="background:var(--warn);">!</span>'
                f'<span>{html_escape(" · ".join(storage))}</span></div>') if storage else ""
n_week = sum(1 for e in ts["entries"] if entry_days(e, info["monday"]))
st.markdown(
    f"""
    <div class="stats-strip">
      <div class="stat-item"><div class="lbl">Team</div><div class="val">{TEAM} <small>{n_active} von {n_total} aktiv</small></div></div>
      <div class="stat-divider"></div>
      <div class="stat-item"><div class="lbl">Woche</div><div class="val">KW {info['kw']} <small>{"gerade" if info["is_even"] else "ungerade"} · {info['monday_str']} – {info['friday_str']}</small></div></div>
      <div class="stat-divider"></div>
      <div class="stat-item"><div class="lbl">Einträge KW {info['kw']}</div><div class="val">{n_week} <small>Abwesenheiten / Wünsche</small></div></div>
      <div class="info-pill" title="Feiertage und 2-Wochen-Rhythmus in dieser KW"><span class="ico">i</span><span>{html_escape(notes_str)}</span></div>
      {storage_html}
    </div>
    """,
    unsafe_allow_html=True,
)


# ══════════════════════════════════════════════════════════════════════
#  Team & Regeln – gemeinsame Bausteine
# ══════════════════════════════════════════════════════════════════════
#  Jede Änderung wird sofort in die Sitzung übernommen (Callback → commit)
#  und gilt für den nächsten Plan; «Speichern» in der Leiste oben macht sie
#  dauerhaft. Intern bleibt der Kurzname ("short") der Schlüssel jeder
#  Person – angezeigt wird überall nur der Name.

TEAM_SECTIONS = {
    "KSC": ["👤 Personen", "🎯 Wer darf was", "📌 Feste Einsätze", "🔄 Rotationen", "⚖️ Regeln", "🗂️ BackUp-Team"],
    "KSV": ["👤 Personen", "🎯 Wer darf was", "📋 Aufgaben & Besetzung", "📌 Feste Einsätze", "🗂️ BackUp-Team"],
}
WOCHEN = {"jede Woche": None, "gerade KW": "gerade", "ungerade KW": "ungerade"}
WOCHEN_LABEL = {v: k for k, v in WOCHEN.items()}
FREI = "frei (nicht da)"
NUR_ZUSATZ = "– (nur Zusatz)"
# Rollen-Matrix KSC: Rolle in team.json → Spaltentitel (Schalter hat eine eigene Auswahl-Spalte)
ROLLEN_SPALTEN = {
    "HUB": "HUB", "ERF8": "ERF8", "ERF9": "ERF9 Mi–Fr", "TPA": "TAGES PA", "PHC": "18-Uhr-Liste",
    "PO": "Postöffnung", "Scanning": "Scanning", "RX Abo": "RX Abo",
    "Direkt": "Direkt", "ONB": "ONB", "BTM": "BTM", "VBZ": "VBZ",
}
SCHALTER_WAHL = ["–", "ja", "Reserve"]
HALBTAG_HILFE = "ganz = ganzer Tag · VM = nur Vormittag · NM = nur Nachmittag · frei = nicht da"

HILFE = {
    "KSC": [
        ("eine neue Person aufnehmen",
         "👤 Personen → «➕ Neue Person». Rollen lassen sich von einer ähnlichen Person übernehmen – danach in 🎯 Wer darf was prüfen."),
        ("jemanden vorübergehend nicht einplanen (krank, Mutterschutz)",
         "👤 Personen → Häkchen «Aktiv» entfernen. Rollen und feste Einsätze bleiben erhalten."),
        ("Pensum, Arbeitstage oder Funktion ändern", "👤 Personen → direkt in der Tabelle."),
        ("jemandem eine Aufgabe erlauben oder wegnehmen", "🎯 Wer darf was → Häkchen setzen oder entfernen."),
        ("einen festen Einsatz ändern (z. B. Mo ERF9 an eine andere Person)",
         "📌 Feste Einsätze → bei der alten Zeile «entfernen» anhaken, oben die neue Zeile hinzufügen."),
        ("«jede 2. Woche frei» eintragen", "📌 Feste Einsätze → Einsatz «frei», Wochen «gerade KW» oder «ungerade KW»."),
        ("korrigieren, wer diese Woche TV, 18-Uhr-Liste, Labor oder eine Wochenaufgabe hat",
         "🔄 Rotationen → bei «diese Woche» die richtige Person wählen; die Folgewochen ergeben sich daraus."),
        ("die Anzahl Personen für TEL, ABKL, PO, HUB, RX Abo … ändern", "⚖️ Regeln"),
        ("eine einmalige Abwesenheit oder einen Sonderwunsch erfassen",
         "Nicht hier – in Schritt 1 «Planung» (eine Woche) oder im 📆 Jahresplan (Ferien von–bis, Feiertage)."),
    ],
    "KSV": [
        ("eine neue Person aufnehmen",
         "👤 Personen → «➕ Neue Person». «Wer darf was» lässt sich von einer ähnlichen Person übernehmen."),
        ("jemanden vorübergehend nicht einplanen (krank, Mutterschutz)",
         "👤 Personen → Häkchen «Aktiv» entfernen. Alles andere bleibt erhalten."),
        ("Teilzeit / freie Halbtage ändern", "👤 Personen → Spalten Mo–Fr (ganz / VM / NM / frei)."),
        ("festlegen, wer contact, KSV1, Redcare … machen darf",
         "🎯 Wer darf was → «ja», «Reserve» (nur wenn sonst niemand frei ist) oder «–»."),
        ("eine neue Aufgabe einführen (so wie Redcare im August)",
         "📋 Aufgaben & Besetzung → «➕ Neue Aufgabe»: Code, Farbe, ganzer Tag oder Halbtag, Personen pro Tag."),
        ("die Anzahl Personen pro Tag ändern (z. B. Montag 2× contact)", "📋 Aufgaben & Besetzung → Spalten Mo–Fr."),
        ("festlegen, was zuerst besetzt wird, wenn zu wenig Leute da sind",
         "📋 Aufgaben & Besetzung → Spalte «Nr.» (1 = wird zuerst besetzt)."),
        ("einen festen Einsatz eintragen (z. B. Quality Di VM)", "📌 Feste Einsätze → oben hinzufügen."),
        ("Arbeitszeiten im Plan anzeigen (z. B. «12-16» im Wochenwechsel)",
         "📌 Feste Einsätze → Einsatz «– (nur Zusatz)», Zusatz «12-16», Wochen «ungerade KW» (und eine zweite Zeile für gerade KW)."),
        ("eine einmalige Abwesenheit oder einen Sonderwunsch erfassen",
         "Nicht hier – in Schritt 1 «Planung» (eine Woche) oder im 📆 Jahresplan (Ferien von–bis, Feiertage)."),
    ],
}

# Regeln, die im KSC-Programm fest eingebaut sind (Änderung = Code-Anpassung)
FESTE_REGELN = [
    "TEL und ABKL: pro Person höchstens ein Halbtag pro Tag (nur-Telefon-Personen ausgenommen).",
    "Tagesverantwortung: dieselbe Teamleitung ganztags, an dem Tag ohne andere Aufgaben (TL/Q).",
    "TAGES PA, Schalter und ERF9 immer ganztags; ERF9 Mi–Fr wechselnd (Mo/Di über feste Einsätze).",
    "ERF8: eine Person pro Nachmittag.",
    "Teamleitungen bekommen nie TAGES PA oder Postöffnung; freie Halbtage werden ERF7/Q.",
    "18-Uhr-Liste: die Person scannt an diesem Tag nicht; nie nur-Telefon- oder Aushilfe-Personen.",
    "Aushilfe: täglich PO/SCAN am Vormittag und scanning am Nachmittag.",
    "Wochenaufgaben Direkt / ONB / BTM: drei verschiedene Personen; wer die ganze Woche fehlt, wird übersprungen.",
    "PO/TEL am Nachmittag wird nie automatisch vergeben (nur Notfall, manuell im Plan).",
]
FESTE_REGELN_KSV = [
    "Abwesenheiten heissen immer KRANK, Ferien oder Frei (Gross-/Kleinschreibung egal) – an diesen Halbtagen wird niemand eingeplant.",
    "Verteilung: zuerst wer die Aufgabe diese Woche am wenigsten hat, dann wer sie gestern nicht hatte, dann wer seltener da ist "
    "(Teilzeit kommt an ihren Tagen zuerst dran), dann wer insgesamt am wenigsten Aufgaben hat.",
    "«Reserve» wird nur eingesetzt, wenn niemand mit «ja» frei ist – im Ergebnis steht dann ein Hinweis.",
    "Ganztags-Aufgabe ohne ganztags freie Person: zwei Personen übernehmen je einen Halbtag (Hinweis «aufgeteilt»).",
    "Mehrere Personen am selben Halbtag werden nummeriert: «contact 1», «contact 2».",
    "Derselbe Montag ergibt mit denselben Einstellungen immer denselben Plan.",
]


def _member(cfg: dict, short: str) -> dict:
    return next(m for m in cfg["mitarbeiter"] if m["short"] == short)


def _cfg_copy() -> dict:
    return copy.deepcopy(ts["cfg"])


def _editor_height(n_rows: int) -> int:
    """Höhe, bei der eine Tabelle ganz ohne eigene Scrollleiste sichtbar ist."""
    return 35 * (n_rows + 1) + 3


def try_commit(new_cfg: dict, msg: str = "", remount: bool = False) -> list[str]:
    """Übernimmt eine geänderte Konfiguration sofort in die Sitzung.
    Liefert Fehler – dann bleibt alles beim Alten."""
    new_cfg = E.normalize_config(new_cfg)
    errors = E.config_errors(new_cfg)
    if errors:
        return errors
    ts["cfg"] = new_cfg
    ts["result"] = None
    if remount:  # Tabellen frisch aufbauen: st.data_editor verwirft seinen Bearbeitungsstand,
        ts["cfg_version"] += 1  # sobald sich die Daten ändern – sonst gehen Folge-Edits verloren
    if msg:
        st.toast(msg, icon="✅")
    return []


def commit_cb(new_cfg: dict, msg: str = "", remount: bool = False) -> None:
    """Für Widget-Callbacks: Fehler oben anzeigen und Widgets auf den gültigen Stand zurücksetzen."""
    errors = try_commit(new_cfg, msg, remount)
    if errors:
        _fail(errors)


def _fail(errors: list[str]) -> None:
    st.session_state[wk("team_error")] = errors
    ts["cfg_version"] += 1


# ── Speicherleiste ────────────────────────────────────────────────────

def _save_cb() -> None:
    st.session_state.save_msgs = save_everything(push_github=bool(github_conf()))


def _discard_cb() -> None:
    replace_cfg(json.loads(ts["cfg_saved_json"]))
    st.toast("Änderungen verworfen – wieder der zuletzt gespeicherte Stand.", icon="↩️")


def render_save_bar() -> None:
    changes = E.describe_changes(json.loads(ts["cfg_saved_json"]), ts["cfg"])
    if entries_dirty:
        changes.append("Einträge (Schritt 1 / Jahresplan) oder Feiertage")
    # Gleiche Höhe in beiden Zuständen, damit die Seite beim ersten Ändern nicht springt
    with st.container(border=True, key="savebar_dirty" if changes else "savebar_ok"):
        c1, c2, c3, c4 = st.columns([5.2, 1.6, 1.2, 1.2], vertical_alignment="center")
        n = len(changes)
        c1.markdown(f"**{n} ungespeicherte Änderung{'en' if n != 1 else ''}** – gilt schon für den nächsten Plan, "
                    "geht aber verloren, wenn die App neu startet." if changes else
                    f"✓ **Alles gespeichert** – der Generator arbeitet mit diesem Stand ({T.config_path.name}).")
        with c2.popover("Was wurde geändert?", width="stretch", disabled=not changes):
            st.markdown("\n".join(f"- {c}" for c in changes))
        c3.button("💾  Speichern", key=wk("savebar_save"), type="primary", width="stretch", on_click=_save_cb,
                  disabled=not changes)
        c4.button("↶  Verwerfen", key=wk("savebar_discard"), width="stretch", on_click=_discard_cb,
                  disabled=not cfg_dirty, help="Team & Regeln auf den zuletzt gespeicherten Stand zurücksetzen")
    for ok, msg in st.session_state.pop("save_msgs", []):
        (st.success if ok else st.warning)(msg)


# ── 👤 Personen ───────────────────────────────────────────────────────

def _on_people_edit(key: str, shorts: list[str]) -> None:
    cfg = _cfg_copy()
    moves, errors = [], []
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        short = shorts[int(ridx)]
        m = _member(cfg, short)
        for col, val in changes.items():
            if col == "Name":
                if _s(val):
                    m["name"] = _s(val)
                else:
                    errors.append("Name darf nicht leer sein")
            elif col == "%" and _s(val):
                m["pct"] = int(float(val))
            elif col == "Aktiv":
                m["aktiv"] = _b(val)
            elif col == "Funktion":
                ak.set_funktion(cfg, short, _s(val) or "Team")
            elif col == "Farbe" and _s(val):
                m["farbe"] = farb_hex(val, m.get("farbe") or "FFFFFF")
            elif col in DAY_KEYS:
                anw = dict(m.get("anwesend", {}))
                anw[col] = _s(val) or "ganz"
                m["anwesend"] = {dk: anw[dk] for dk in DAY_KEYS if anw.get(dk, "ganz") != "ganz"}
            elif col == "Bemerkung":
                m["bemerkung"] = _s(val)
            elif col == "Nr." and _s(val):
                moves.append((short, int(float(val))))
    if errors:
        _fail(errors)
        return
    for short, pos in moves:
        lst = cfg["mitarbeiter"]
        m = _member(cfg, short)
        lst.remove(m)
        lst.insert(max(0, min(len(lst), pos - 1)), m)
    commit_cb(cfg, msg="Reihenfolge geändert" if moves else "", remount=True)


@st.dialog("Neue Person", width="large")
def dialog_add_person() -> None:
    cfg = ts["cfg"]
    shorts = [m["short"] for m in cfg["mitarbeiter"]]
    name = st.text_input("Name, wie er im Plan erscheint", placeholder="z. B. Sandra K.",
                         help="Vorname und Initial des Nachnamens – im KSV-Excel stehen sie in «Vorname» und «Name».")
    c1, c2 = st.columns(2)
    pct = c1.number_input("Pensum %", min_value=10, max_value=100, value=100, step=10)
    if IS_KSC:
        fkt = c2.selectbox("Funktion", ak.FUNKTIONEN,
                           help="Teamleitung: Tagesverantwortung, Queue, 18-Uhr-Liste Montag · nur Telefon: immer TEL · "
                                "Aushilfe: täglich PO/SCAN + scanning")
        farbe = None
    else:
        used = {m.get("farbe") for m in cfg["mitarbeiter"]}
        default = next((c for c in tp.ZEILENFARBEN if c not in used), tp.ZEILENFARBEN[0])
        opts = farb_options(default)
        farbe = farb_hex(c2.selectbox("Zeilenfarbe im Excel", opts, index=opts.index(farb_label(default))))
        fkt = None
    st.markdown(f"**Anwesenheit** <span style='color:var(--text-mute); font-size:12px;'>{HALBTAG_HILFE}</span>",
                unsafe_allow_html=True)
    dcols = st.columns(5)
    anw = {dk: dcols[i].selectbox(dk, HALBTAG_WERTE, key=wk(f"np_anw_{dk}")) for i, dk in enumerate(DAY_KEYS)}
    c3, c4 = st.columns(2)
    regs = [m["short"] for m in cfg["mitarbeiter"] if not IS_KSC or ak.funktion(m) in ("Team", "Teamleitung")]
    vorlage = c3.selectbox("«Wer darf was» übernehmen von (optional)", ["–"] + regs,
                           format_func=lambda s: "– niemand –" if s == "–" else disp(s),
                           help="Die neue Person darf dann dieselben Aufgaben – anpassen unter «Wer darf was».")
    nach = c4.selectbox("Position im Plan", ["–"] + shorts,
                        format_func=lambda s: "ganz unten" if s == "–" else f"nach {disp(s)}")
    if st.button("➕  Hinzufügen", type="primary"):
        name = name.strip()
        if not name:
            st.error("Bitte einen Namen eingeben.")
            return
        new = _cfg_copy()
        short = E.new_short(name, new)
        tmpl = _member(new, vorlage) if vorlage != "–" else {}
        m = {"short": short, "name": name, "pct": int(pct), "aktiv": True,
             "kann": list(tmpl.get("kann", [])),
             "anwesend": {dk: v for dk, v in anw.items() if v != "ganz"},
             "frei_alle_2_wochen": None, "fix": [], "bemerkung": ""}
        if IS_KSC:
            m.update({"tl": False, "nur_tel": False, "aushilfe": False})
        else:
            m.update({"reserve": list(tmpl.get("reserve", [])), "farbe": farbe})
        lst = new["mitarbeiter"]
        lst.insert(len(lst) if nach == "–" else [x["short"] for x in lst].index(nach) + 1, m)
        if IS_KSC:
            ak.set_funktion(new, short, fkt)
        errors = try_commit(new, msg=f"{name} hinzugefügt – «Wer darf was» prüfen", remount=True)
        if errors:
            st.error(" · ".join(errors))
            return
        st.rerun()


@st.dialog("Person entfernen")
def dialog_remove_person() -> None:
    cfg = ts["cfg"]
    shorts = [m["short"] for m in cfg["mitarbeiter"]]
    who = st.selectbox("Wen?", shorts,
                       format_func=lambda s: disp(s) + ("" if _member(cfg, s).get("aktiv", True) else " (inaktiv)"))
    aktiv = _member(cfg, who).get("aktiv", True)
    st.warning(f"**{disp(who)}** wird mit allen Rollen und festen Einsätzen gelöscht.\n\n"
               "Nur vorübergehend weg (krank, Mutterschutz, lange Ferien)? Dann besser **deaktivieren** – "
               "so bleibt alles erhalten, und später genügt das Häkchen «Aktiv».")
    refs = E.rotation_refs(cfg, who)
    if refs:
        st.info("Wird beim Entfernen auch hier herausgenommen: " + ", ".join(refs))
    c1, c2, c3 = st.columns(3)
    if c1.button("Deaktivieren", width="stretch", disabled=not aktiv):
        new = _cfg_copy()
        _member(new, who)["aktiv"] = False
        try_commit(new, msg=f"{disp(who)} deaktiviert", remount=True)
        st.rerun()
    if c2.button("Endgültig entfernen", type="primary", width="stretch"):
        name = disp(who)
        new = _cfg_copy()
        E.remove_member(new, who)
        try_commit(new, msg=f"{name} entfernt", remount=True)
        st.rerun()
    if c3.button("Abbrechen", width="stretch"):
        st.rerun()


def section_people(ver: int) -> None:
    cfg = ts["cfg"]
    members = cfg["mitarbeiter"]
    c1, c2, c3 = st.columns([1.2, 1.5, 6], vertical_alignment="center")
    if c1.button("➕  Neue Person", key=wk(f"btn_add_person_{ver}"), type="primary", width="stretch"):
        dialog_add_person()
    if c2.button("🗑️  Person entfernen …", key=wk(f"btn_rm_person_{ver}"), width="stretch"):
        dialog_remove_person()
    c3.caption("Direkt in der Tabelle ändern – jede Änderung gilt sofort. **Aktiv** ausschalten, wenn jemand "
               "vorübergehend nicht eingeplant werden soll; alles andere bleibt erhalten. "
               "**Nr.** ändern verschiebt die Person im Plan." +
               ("" if IS_KSC else " **Farbe** = Zeilenfarbe im Excel."))
    extra = "Funktion" if IS_KSC else "Farbe"
    rows = [{"Nr.": i + 1, "Name": m.get("name") or m["short"], "%": int(m.get("pct", 100)),
             "Aktiv": bool(m.get("aktiv", True)),
             extra: ak.funktion(m) if IS_KSC else farb_label(m.get("farbe")),
             **{dk: m.get("anwesend", {}).get(dk, "ganz") for dk in DAY_KEYS},
             "Bemerkung": m.get("bemerkung", "")} for i, m in enumerate(members)]
    df = pd.DataFrame(rows, columns=["Nr.", "Name", "%", "Aktiv", extra] + DAY_KEYS + ["Bemerkung"])
    col_config = {
        "Nr.": st.column_config.NumberColumn("Nr.", min_value=1, max_value=max(1, len(rows)), step=1, width=48,
                                             help="Reihenfolge im Plan – Zahl ändern, um die Person zu verschieben"),
        "Name": st.column_config.TextColumn("Name", width=140, pinned=True,
                                            help="So erscheint die Person im Plan, z. B. «Sandra K.»"),
        "%": st.column_config.NumberColumn("%", min_value=0, max_value=100, step=5, width=56, help="Pensum in Prozent"),
        "Aktiv": st.column_config.CheckboxColumn("Aktiv", width=60,
                                                 help="Aus = wird vorübergehend nicht eingeplant"),
        "Bemerkung": st.column_config.TextColumn("Bemerkung", width="large"),
    }
    if IS_KSC:
        col_config["Funktion"] = st.column_config.SelectboxColumn(
            "Funktion", options=ak.FUNKTIONEN, required=True, width=120,
            help="Teamleitung: Tagesverantwortung, Queue, 18-Uhr-Liste Montag · nur Telefon: immer TEL · "
                 "Aushilfe: täglich PO/SCAN + scanning")
    else:
        col_config["Farbe"] = st.column_config.SelectboxColumn(
            "Farbe", options=farb_options(*(m.get("farbe") for m in members)), required=True, width=110,
            help="Zeilenfarbe (Name/Vorname) im Excel")
    for dk in DAY_KEYS:
        col_config[dk] = st.column_config.SelectboxColumn(dk, options=HALBTAG_WERTE, required=True, width=64,
                                                          help=f"Anwesenheit {dk}: {HALBTAG_HILFE}")
    key = wk(f"people_{ver}")
    st.data_editor(df, key=key, column_config=col_config, hide_index=True, num_rows="fixed", width="stretch",
                   height=_editor_height(len(df)), on_change=_on_people_edit, args=(key, [m["short"] for m in members]))
    if not IS_KSC:
        st.markdown('<div class="legend-wrap">' + "".join(
            f'<span class="legend-item"><span class="swatch" style="background:#{m.get("farbe")};"></span>'
            f'{html_escape(m.get("name"))}</span>' for m in members) + "</div>", unsafe_allow_html=True)


# ── 🎯 Wer darf was ───────────────────────────────────────────────────

def _on_roles_edit(key: str, shorts: list[str]) -> None:
    cfg = _cfg_copy()
    by_label = {lbl: r for r, lbl in ROLLEN_SPALTEN.items()}
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        m = _member(cfg, shorts[int(ridx)])
        kann = list(m.get("kann", []))
        for col, val in changes.items():
            if col == "Schalter":
                kann = [r for r in kann if r not in ("Schalter", "Schalter-Reserve")]
                if val == "ja":
                    kann.append("Schalter")
                elif val == "Reserve":
                    kann.append("Schalter-Reserve")
            elif col in by_label:
                role = by_label[col]
                kann = [r for r in kann if r != role] + ([role] if _b(val) else [])
        m["kann"] = kann
    commit_cb(cfg, remount=True)


def section_roles_ksc(ver: int) -> None:
    cfg = ts["cfg"]
    regs = [m for m in cfg["mitarbeiter"] if not m.get("nur_tel") and not m.get("aushilfe")]
    st.caption("Häkchen = darf diese Aufgabe übernehmen – der Generator verteilt Aufgaben nur an Personen mit Häkchen. "
               "**Schalter**: «ja» = regulär (zwei Gruppen im Wochenwechsel), «Reserve» = nur, wenn aus der Gruppe "
               "niemand frei ist. Teamleitungen (TL) bekommen nie TAGES PA oder Postöffnung. "
               "Personen mit «nur Telefon» oder «Aushilfe» fehlen hier – sie haben feste Aufgaben.")
    rows = []
    for m in regs:
        kann = set(m.get("kann", []))
        label = (m.get("name") or m["short"]) + (" · TL" if m.get("tl") else "") + \
                ("" if m.get("aktiv", True) else " (inaktiv)")
        row = {"Name": label,
               "Schalter": "ja" if "Schalter" in kann else "Reserve" if "Schalter-Reserve" in kann else "–"}
        row.update({lbl: r in kann for r, lbl in ROLLEN_SPALTEN.items()})
        rows.append(row)
    df = pd.DataFrame(rows, columns=["Name", "Schalter"] + list(ROLLEN_SPALTEN.values()))
    col_config = {
        "Name": st.column_config.TextColumn("Name", disabled=True, pinned=True, width=170),
        "Schalter": st.column_config.SelectboxColumn("Schalter", options=SCHALTER_WAHL, required=True, width=90,
                                                     help="ja = regulär · Reserve = nur wenn aus der Gruppe niemand frei ist"),
    }
    for r, lbl in ROLLEN_SPALTEN.items():
        col_config[lbl] = st.column_config.CheckboxColumn(lbl, help=ROLLEN[r], width=max(62, int(7.5 * len(lbl)) + 26))
    key = wk(f"roles_{ver}")
    st.data_editor(df, key=key, column_config=col_config, hide_index=True, num_rows="fixed", width="stretch",
                   height=_editor_height(len(df)), on_change=_on_roles_edit, args=(key, [m["short"] for m in regs]))
    act = [m for m in regs if m.get("aktiv", True)]
    counts = [("Schalter", sum("Schalter" in m.get("kann", []) for m in act))] + \
             [(lbl, sum(r in m.get("kann", []) for m in act)) for r, lbl in ROLLEN_SPALTEN.items() if r != "VBZ"]
    st.caption("Aktive Personen pro Aufgabe: " +
               " · ".join(f"{lbl} **{n}**" + (" ⚠️" if n <= 1 else "") for lbl, n in counts))


STUFE_LABEL = {"ja": "✅ ja", "Reserve": "🟨 Reserve", "–": "— nie"}
STUFE_WERT = {v: k for k, v in STUFE_LABEL.items()}


def _on_stufen_edit(key: str, shorts: list[str], codes: list[str]) -> None:
    cfg = _cfg_copy()
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        m = _member(cfg, shorts[int(ridx)])
        for col, val in changes.items():
            if col in codes:
                tp.set_stufe(m, col, STUFE_WERT.get(_s(val), "–"))
    commit_cb(cfg, remount=True)


def section_roles_ksv(ver: int) -> None:
    cfg = ts["cfg"]
    autos = tp.auto_tasks(cfg)
    fill = cfg["regeln"].get("auffuellen") or "–"
    st.caption("Pro Person und Aufgabe: **✅ ja** = wird regelmässig eingeplant · **🟨 Reserve** = nur, wenn niemand "
               "mit «ja» frei ist · **— nie**. Der Generator verteilt die Aufgaben gleichmässig auf alle mit «ja». "
               f"Das Auffüllen ({fill}) darf jede Person – dafür gibt es keine Spalte.")
    if not autos:
        st.info("Noch keine Aufgaben, die automatisch geplant werden – zuerst unter 📋 Aufgaben & Besetzung anlegen.")
        return
    codes = [t["code"] for t in autos]
    members = cfg["mitarbeiter"]
    rows = [{"Name": (m.get("name") or m["short"]) + ("" if m.get("aktiv", True) else " (inaktiv)"),
             **{c: STUFE_LABEL[tp.stufe(m, c)] for c in codes}} for m in members]
    df = pd.DataFrame(rows, columns=["Name"] + codes)
    col_config = {"Name": st.column_config.TextColumn("Name", disabled=True, pinned=True, width=160)}
    for t in autos:
        col_config[t["code"]] = st.column_config.SelectboxColumn(
            t["code"], options=list(STUFE_LABEL.values()), required=True, width=max(110, 9 * len(t["code"]) + 30),
            help=f"{t.get('bedeutung') or t['code']} · {t['dauer']}")
    key = wk(f"stufen_{ver}")
    st.data_editor(df, key=key, column_config=col_config, hide_index=True, num_rows="fixed", width="stretch",
                   height=_editor_height(len(df)), on_change=_on_stufen_edit,
                   args=(key, [m["short"] for m in members], codes))
    act = [m for m in members if m.get("aktiv", True)]
    parts = []
    for t in autos:
        ja = sum(t["code"] in m.get("kann", []) for m in act)
        res = sum(t["code"] in m.get("reserve", []) for m in act)
        warn = " ⚠️" if ja + res <= max(t["anzahl"].values()) else ""
        parts.append(f"{t['code']} **{ja}**" + (f" + {res} Reserve" if res else "") + warn)
    st.caption("Aktive Personen pro Aufgabe: " + " · ".join(parts) +
               (" — ⚠️ = kaum mehr Personen als pro Tag gebraucht werden" if any("⚠️" in x for x in parts) else ""))


# ── 📌 Feste Einsätze ────────────────────────────────────────────────

def _fix_options(cfg: dict) -> list[str]:
    if IS_KSC:
        opts = [FREI] + [t for t in TASK_OPTIONS if t and t not in ("KRANK", "FERIEN")]
    else:
        opts = [FREI, NUR_ZUSATZ] + tp.task_codes(cfg)
    for r in E.recurring_rows(cfg):
        if r["task"] not in ("frei", "") and r["task"] not in opts:
            opts.append(r["task"])
    return opts


def _einsatz_label(task: str) -> str:
    return FREI if task == "frei" else NUR_ZUSATZ if not task else task


def _einsatz_code(label: str) -> str:
    return "frei" if label == FREI else "" if label == NUR_ZUSATZ else label


def _commit_recurring(rows: list[dict], msg: str = "", remount: bool = False) -> list[str]:
    cfg = _cfg_copy()
    errors = E.apply_recurring_rows(cfg, rows)
    return errors or try_commit(cfg, msg=msg, remount=remount)


def _on_fixed_edit(key: str, base: list[dict]) -> None:
    rows = [dict(r) for r in base]
    removed = set()
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        i = int(ridx)
        r = rows[i]
        for col, val in changes.items():
            if col == "entfernen" and _b(val):
                removed.add(i)
            elif col == "Tag":
                r["tag"] = _s(val) or r["tag"]
            elif col == "Halbtag":
                r["halbtag"] = _s(val) or r["halbtag"]
            elif col == "Einsatz" and _s(val):
                r["task"] = _einsatz_code(_s(val))
            elif col == "Zusatz":
                r["zusatz"] = _s(val)
            elif col == "Wochen":
                r["kw"] = WOCHEN.get(_s(val))
    errors = _commit_recurring([r for i, r in enumerate(rows) if i not in removed],
                               msg="Fester Einsatz entfernt" if removed else "", remount=True)
    if errors:
        _fail(errors)


def section_fixed(ver: int) -> None:
    cfg = ts["cfg"]
    shorts = [m["short"] for m in cfg["mitarbeiter"]]
    zusatz = not IS_KSC
    st.caption("Was jede Woche – oder jede zweite Woche – fix ist, z. B. " +
               ("«Mo ganz ERF9», «Di VM KGS» oder «alle 2 Wochen Fr frei». " if IS_KSC else
                "«Di VM Quality» oder «alle 2 Wochen Fr frei». **Zusatz** = Text, der im Plan an die Aufgabe "
                "angehängt wird, z. B. die Arbeitszeit «12-16» (Einsatz «– (nur Zusatz)»: die Aufgabe plant der Generator). ") +
               "Der Generator plant diese Einsätze vor allem anderen. "
               f"Die angezeigte **KW {info['kw']}** ist eine **{'gerade' if info['is_even'] else 'ungerade'}** Woche.")
    opts = _fix_options(cfg)
    with st.form(wk(f"add_fix_{ver}"), border=True):
        st.markdown("**Neuer fester Einsatz**")
        widths = [2, 0.9, 1, 1.6, 1.3, 1.3, 1.3] if zusatz else [2, 0.9, 1, 1.6, 1.3, 1.3]
        c = st.columns(widths, vertical_alignment="bottom")
        who = c[0].selectbox("Person", shorts, format_func=disp)
        tag = c[1].selectbox("Tag", DAY_KEYS)
        halb = c[2].selectbox("Halbtag", ["ganz", "VM", "NM"])
        manual = [t["code"] for t in tp.task_list(cfg)
                  if t["dauer"] == "nur manuell" and t["code"] != cfg["regeln"].get("auffuellen")] if not IS_KSC else []
        default = "HO" if IS_KSC else (manual[0] if manual else NUR_ZUSATZ)
        task = c[3].selectbox("Einsatz", opts, index=opts.index(default) if default in opts else 0,
                              help="«frei» = an diesem Tag nicht da (nur mit «gerade KW» / «ungerade KW»)")
        ztext = c[4].text_input("Zusatz", placeholder="z. B. 12-16") if zusatz else ""
        wochen = c[-2].selectbox("Wochen", list(WOCHEN))
        if c[-1].form_submit_button("➕  Hinzufügen", type="primary", width="stretch"):
            if task == NUR_ZUSATZ and not ztext.strip():
                st.error("Bei «– (nur Zusatz)» bitte einen Zusatztext eingeben, z. B. «12-16».")
            else:
                new_row = {"short": who, "tag": tag, "halbtag": halb, "task": _einsatz_code(task),
                           "zusatz": ztext.strip(), "kw": WOCHEN[wochen]}
                errors = _commit_recurring(E.recurring_rows(cfg) + [new_row],
                                           msg=f"{disp(who)}: {tag} {halb} {task} {ztext}".strip() + " hinzugefügt",
                                           remount=True)
                if errors:
                    st.error(" · ".join(errors))
                else:
                    st.rerun()

    base = E.recurring_rows(cfg)
    if not base:
        st.info("Noch keine festen Einsätze.")
        return
    inaktiv = {m["short"] for m in cfg["mitarbeiter"] if not m.get("aktiv", True)}
    cols = ["Person", "Tag", "Halbtag", "Einsatz"] + (["Zusatz"] if zusatz else []) + ["Wochen", "entfernen"]
    df = pd.DataFrame([{"Person": disp(r["short"]) + (" (inaktiv)" if r["short"] in inaktiv else ""),
                        "Tag": r["tag"], "Halbtag": r["halbtag"], "Einsatz": _einsatz_label(r["task"]),
                        "Zusatz": r.get("zusatz", ""),
                        "Wochen": WOCHEN_LABEL.get(r["kw"], "jede Woche"), "entfernen": False} for r in base],
                      columns=cols)
    key = wk(f"fixed_{ver}")
    col_config = {
        "Person": st.column_config.TextColumn("Person", width="medium"),
        "Tag": st.column_config.SelectboxColumn("Tag", options=DAY_KEYS, required=True, width="small"),
        "Halbtag": st.column_config.SelectboxColumn("Halbtag", options=["ganz", "VM", "NM"], required=True, width="small"),
        "Einsatz": st.column_config.SelectboxColumn("Einsatz", options=opts, required=True, width="medium"),
        "Wochen": st.column_config.SelectboxColumn("Wochen", options=list(WOCHEN), required=True, width="small"),
        "entfernen": st.column_config.CheckboxColumn("entfernen", width="small", help="Anhaken = Zeile löschen"),
    }
    if zusatz:
        col_config["Zusatz"] = st.column_config.TextColumn("Zusatz", width="small",
                                                           help="wird im Plan angehängt, z. B. «12-16»")
    st.data_editor(df, key=key, hide_index=True, num_rows="fixed", width="stretch", height=_editor_height(len(df)),
                   disabled=["Person"], on_change=_on_fixed_edit, args=(key, base), column_config=col_config)


# ── 🔄 Rotationen (KSC) ──────────────────────────────────────────────

def _on_rot_list(key: str, field: str, hidden: list[str]) -> None:
    cfg = _cfg_copy()
    cfg["rotation"][field] = list(st.session_state[key]) + [n for n in hidden if n not in st.session_state[key]]
    commit_cb(cfg)


def _on_rot_now(key: str, off_key: str, pool: list, widx: int, mult: int) -> None:
    cfg = _cfg_copy()
    cfg["rotation"]["offset"][off_key] = (pool.index(st.session_state[key]) - widx * mult) % len(pool)
    commit_cb(cfg)


def _on_rot_value(key: str, field: str) -> None:
    cfg = _cfg_copy()
    v = st.session_state[key]
    cfg["rotation"][field] = "" if v == "–" else v
    commit_cb(cfg)


def _on_phc_fix(key: str, dk: str) -> None:
    cfg = _cfg_copy()
    fix = dict(cfg["rotation"]["phc_fix"])
    if st.session_state[key] == "–":
        fix.pop(dk, None)
    else:
        fix[dk] = st.session_state[key]
    cfg["rotation"]["phc_fix"] = {d: fix[d] for d in DAY_KEYS if d in fix}
    commit_cb(cfg)


def _on_phc_excl(key: str, shorts: list[str]) -> None:
    cfg = _cfg_copy()
    ex = {dk: list(v) for dk, v in cfg["rotation"]["phc_tag_ausschluss"].items()}
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        s = shorts[int(ridx)]
        for dk, val in changes.items():
            lst = [n for n in ex.get(dk, []) if n != s]
            ex[dk] = lst + [s] if _b(val) else lst
    cfg["rotation"]["phc_tag_ausschluss"] = {dk: ex[dk] for dk in DAY_KEYS if ex.get(dk)}
    commit_cb(cfg, remount=True)


def _rotation_block(ver: int, title: str, caption: str, field: str, off_key: str, options: list[str],
                    now_label: str, mult: int = 1) -> list[str]:
    """Liste «wer, in welcher Reihenfolge» + «wer ist in der angezeigten Woche dran»."""
    rot = ts["cfg"]["rotation"]
    widx = info["widx"]
    st.markdown(f"**{title}**")
    st.caption(caption)
    current = [n for n in rot[field] if n in options]
    hidden = [n for n in rot[field] if n not in options]
    key = wk(f"rot_{field}_{ver}")
    pool = st.multiselect(title, options, default=current, format_func=disp, key=key, label_visibility="collapsed",
                          placeholder="Personen wählen – Reihenfolge = Reihenfolge der Auswahl",
                          on_change=_on_rot_list, args=(key, field, hidden))
    if hidden:
        st.caption("Inaktiv, wird übersprungen: " + ", ".join(disp(n) for n in hidden))
    if pool:
        idx = (widx * mult + rot["offset"].get(off_key, 0)) % len(pool)
        nkey = wk(f"rot_{off_key}_now_{ver}_{'|'.join(pool)}")
        st.selectbox(now_label, pool, index=idx, format_func=disp, key=nkey,
                     on_change=_on_rot_now, args=(nkey, off_key, list(pool), widx, mult))
    return pool


def _rotation_preview(cfg: dict, n_weeks: int = 6) -> pd.DataFrame:
    """Wer ist in den nächsten Wochen laut Rotation dran (ohne Abwesenheiten)."""
    rot, off = cfg["rotation"], cfg["rotation"]["offset"]
    act = [m for m in cfg["mitarbeiter"] if m.get("aktiv", True)]
    act_set = {m["short"] for m in act}
    reg = [m for m in act if not m.get("nur_tel") and not m.get("aushilfe")]
    tv = [n for n in rot["tv_reihenfolge"] if n in act_set] or [m["short"] for m in act if m.get("tl")]
    lists = {k: [n for n in rot[k] if n in act_set] for k in ("phc_montag_tl", "phc_freitag_wechsel", "labor_mi_nm")}
    wa_pools = {lbl: [m["short"] for m in reg if lbl in m.get("kann", [])] for lbl in ("Direkt", "ONB", "BTM")}

    def at(lst, w, off_key, mult=1):
        return disp(lst[(w * mult + off.get(off_key, 0)) % len(lst)]) if lst else "–"

    rows = []
    for i in range(n_weeks):
        w = info["widx"] + i
        monday = info["monday"] + timedelta(weeks=i)
        used, wa = set(), {}
        for lbl, pool in wa_pools.items():
            if not pool:
                wa[lbl] = "–"
                continue
            idx = (w + off.get(lbl, 0)) % len(pool)
            cand = next((pool[(idx + j) % len(pool)] for j in range(len(pool))
                         if pool[(idx + j) % len(pool)] not in used), pool[idx])
            used.add(cand)
            wa[lbl] = disp(cand)
        rows.append({
            "KW": f"KW {monday.isocalendar()[1]}", "ab": monday.strftime("%d.%m."),
            "TV Montag": at(tv, w, "tv", 5), "18-Uhr Mo": at(lists["phc_montag_tl"], w, "phc_montag"),
            "18-Uhr Fr": at(lists["phc_freitag_wechsel"], w, "phc_freitag"),
            "Labor Mi": at(lists["labor_mi_nm"], w, "labor"),
            "Schalter": "Gruppe A" if (w + off.get("schalter", 0)) % 2 == 0 else "Gruppe B", **wa,
        })
    return pd.DataFrame(rows)


def section_rotation(ver: int) -> None:
    cfg = ts["cfg"]
    rot = cfg["rotation"]
    widx = info["widx"]
    kw = info["kw"]
    act = [m for m in cfg["mitarbeiter"] if m.get("aktiv", True)]
    reg = [m for m in act if not m.get("nur_tel") and not m.get("aushilfe")]
    tls = [m["short"] for m in act if m.get("tl")]
    phc_ok = [m["short"] for m in reg if "PHC" in m.get("kann", [])]
    st.caption(f"Rotationen ergeben sich aus der Kalenderwoche. Mit «diese Woche» stellst du ein, wer in der "
               f"angezeigten **KW {kw}** dran ist – die Folgewochen laufen von dort aus weiter (Vorschau unten). "
               f"Die Reihenfolge entspricht der Reihenfolge der Auswahl: zum Umsortieren Person entfernen (×) und neu wählen.")

    left, right = st.columns(2, gap="large")
    with left:
        with st.container(border=True):
            tv = _rotation_block(ver, "Tagesverantwortung (TV)",
                                 "Die Teamleitungen wechseln sich Tag für Tag ab. Leer = alle Teamleitungen.",
                                 "tv_reihenfolge", "tv", tls, f"TV am Montag, KW {kw}", mult=5)
            tv_list = tv or tls
            if tv_list:
                start = (widx * 5 + rot["offset"]["tv"]) % len(tv_list)
                st.caption("Diese Woche: " + " · ".join(f"{DAY_KEYS[d]} {disp(tv_list[(start + d) % len(tv_list)])}"
                                                         for d in range(5)))
        with st.container(border=True):
            _rotation_block(ver, "Labor Mi Nachmittag (ERF5)", "Wöchentlicher Wechsel.",
                            "labor_mi_nm", "labor", [m["short"] for m in reg], f"Labor in KW {kw}")
        with st.container(border=True):
            st.markdown("**Schalter-Gruppe**")
            sch_pool = [m["short"] for m in reg if "Schalter" in m.get("kann", [])]
            mid = (len(sch_pool) + 1) // 2
            groups = [sch_pool[:mid], sch_pool[mid:]]
            st.caption("Alle mit Schalter «ja» (unter «Wer darf was») werden in Plan-Reihenfolge in zwei Gruppen "
                       "geteilt, die sich wöchentlich abwechseln.")
            key = wk(f"rot_schalter_{ver}_{'|'.join(sch_pool)}")
            st.selectbox(f"Aktive Gruppe in KW {kw}", [0, 1], index=(widx + rot["offset"]["schalter"]) % 2, key=key,
                         format_func=lambda g: f"Gruppe {'AB'[g]}: {', '.join(disp(n) for n in groups[g]) or '–'}",
                         on_change=_on_rot_now, args=(key, "schalter", [0, 1], widx, 1))
        with st.container(border=True):
            st.markdown("**Wochenaufgaben**")
            st.caption("Je eine Person pro Woche, nie dieselbe doppelt. Wer infrage kommt, steht unter «Wer darf was».")
            wcols = st.columns(3)
            for col, label in zip(wcols, ("Direkt", "ONB", "BTM")):
                pool = [m["short"] for m in reg if label in m.get("kann", [])]
                if not pool:
                    col.caption(f"{label}: niemand berechtigt")
                    continue
                key = wk(f"rot_wa_{label}_{ver}_{'|'.join(pool)}")
                col.selectbox(f"{label} in KW {kw}", pool, index=(widx + rot["offset"].get(label, 0)) % len(pool),
                              format_func=disp, key=key, on_change=_on_rot_now, args=(key, label, pool, widx, 1))

    with right:
        with st.container(border=True):
            st.markdown("**18-Uhr-Liste (PHC)**")
            _rotation_block(ver, "Montag – Teamleitungen im Wechsel", "Wöchentlicher Wechsel.",
                            "phc_montag_tl", "phc_montag", tls, f"Montag in KW {kw}")
            _rotation_block(ver, "Freitag – im Wechsel", "Wöchentlicher Wechsel. Leer = aus dem Pool.",
                            "phc_freitag_wechsel", "phc_freitag", phc_ok, f"Freitag in KW {kw}")
            st.markdown("**Fix an einem Wochentag**")
            st.caption("Diese Person macht die 18-Uhr-Liste an diesem Tag immer (wenn sie da ist). "
                       "Montag und Freitag laufen über die Wechsel oben, sofern dort jemand eingetragen ist. "
                       "– = jemand aus dem Pool.")
            fcols = st.columns(5)
            for i, dk in enumerate(DAY_KEYS):
                cur = rot["phc_fix"].get(dk, "–")
                opts = ["–"] + phc_ok + ([cur] if cur not in phc_ok and cur != "–" else [])
                key = wk(f"rot_phcfix_{dk}_{ver}")
                fcols[i].selectbox(dk, opts, index=opts.index(cur), key=key, format_func=lambda s: s if s == "–" else disp(s),
                                   on_change=_on_phc_fix, args=(key, dk))
            st.markdown("**Nie an diesem Tag**")
            st.caption("Häkchen = diese Person kommt an diesem Wochentag nicht auf die 18-Uhr-Liste.")
            ex = rot["phc_tag_ausschluss"]
            df = pd.DataFrame([{"Name": disp(s), **{dk: s in ex.get(dk, []) for dk in DAY_KEYS}} for s in phc_ok],
                              columns=["Name"] + DAY_KEYS)
            key = wk(f"rot_phcex_{ver}")
            st.data_editor(df, key=key, hide_index=True, num_rows="fixed", width="stretch",
                           height=_editor_height(len(df)), disabled=["Name"], on_change=_on_phc_excl, args=(key, phc_ok),
                           column_config={"Name": st.column_config.TextColumn("Name", width="medium"),
                                          **{dk: st.column_config.CheckboxColumn(dk, width="small") for dk in DAY_KEYS}})

    st.markdown(f"**Vorschau – wer ist laut Rotation dran** "
                f"<span style='color:var(--text-mute); font-size:12px;'>ohne Abwesenheiten; ab KW {kw}</span>",
                unsafe_allow_html=True)
    st.dataframe(_rotation_preview(cfg), hide_index=True, width="stretch")


# ── ⚖️ Regeln (KSC) ───────────────────────────────────────────────────

def _on_rule(key: str, field: str, kind: type) -> None:
    cfg = _cfg_copy()
    cfg["regeln"][field] = kind(st.session_state[key])
    commit_cb(cfg)


def _rule_number(container, label: str, field: str, lo: int, hi: int, ver: int, help: str | None = None) -> int:
    key = wk(f"rule_{field}_{ver}")
    return container.number_input(label, lo, hi, int(ts["cfg"]["regeln"][field]), step=1, key=key, help=help,
                                  on_change=_on_rule, args=(key, field, int))


def _rule_select(container, label: str, field: str, options: list[str], ver: int, help: str | None = None) -> None:
    cur = ts["cfg"]["regeln"].get(field, options[0])
    key = wk(f"rule_{field}_{ver}")
    container.selectbox(label, options, index=options.index(cur) if cur in options else 0, key=key, help=help,
                        on_change=_on_rule, args=(key, field, str))


def section_rules(ver: int) -> None:
    cfg = ts["cfg"]
    rules = cfg["regeln"]
    st.caption("Wie viele Personen braucht es wofür? Jede Änderung gilt sofort für den nächsten Plan. "
               "Wer welche Aufgabe übernehmen darf, steht unter «Wer darf was».")
    left, right = st.columns(2, gap="large")
    with left:
        with st.container(border=True):
            st.markdown("**📞 Telefon (TEL)**")
            c1, c2 = st.columns(2)
            _rule_number(c1, "Montag Vormittag – Personen", "tel_montag_vm", 0, 10, ver)
            _rule_number(c2, "übrige Halbtage – Personen", "tel_sonst", 0, 10, ver)
            c3, c4 = st.columns(2)
            tl_min = _rule_number(c3, "Teamleitung am Telefon – mind. Halbtage/Woche", "tl_tel_min_pro_woche", 0, 10, ver,
                                  help="Wird für jede TL vorab reserviert. 0 = kein Minimum.")
            tl_max = _rule_number(c4, "… höchstens Halbtage/Woche", "tl_tel_max_pro_woche", 0, 10, ver,
                                  help="Bis zu dieser Zahl bekommen TL beim Telefon den Vorzug.")
            if tl_min > tl_max:
                st.warning("Minimum ist grösser als Maximum – das Minimum gewinnt.")
            _rule_select(st, "TAGES PA am Nachmittag zusätzlich am Telefon (TAGES PA/TEL)", "tpa_tel_nm", TPA_TEL_MODI, ver,
                         help="bei Bedarf: nur wenn die TEL-Besetzung sonst nicht reicht · immer · nie")
        with st.container(border=True):
            st.markdown("**🔎 Abklärung (ABKL)**")
            _rule_number(st, "Personen pro Halbtag", "abkl_pro_halbtag", 0, 10, ver)
        with st.container(border=True):
            st.markdown("**📬 Postöffnung (PO)**")
            _rule_number(st, "Zusätzliche PO-Personen am Vormittag (zur Aushilfe mit PO/SCAN)", "po_zusatz_vm", 0, 5, ver)
            po_opts = ["–"] + [m["short"] for m in cfg["mitarbeiter"]
                               if m.get("aktiv", True) and "PO" in m.get("kann", []) and not m.get("tl")]
            cur = cfg["rotation"].get("po_dienstag_bevorzugt") or "–"
            if cur not in po_opts:
                po_opts.append(cur)
            key = wk(f"rot_popref_{ver}")
            st.selectbox("Dienstag bevorzugt", po_opts, index=po_opts.index(cur), key=key,
                         format_func=lambda s: "– niemand –" if s == "–" else disp(s),
                         on_change=_on_rot_value, args=(key, "po_dienstag_bevorzugt"))
            key = wk(f"rule_scan_nm_{ver}")
            st.checkbox("Fehlt die Aushilfe: Nachmittags-Scan durch jemand anderen ersetzen",
                        bool(rules["scan_nm_ersatz_wenn_aushilfe_fehlt"]), key=key,
                        on_change=_on_rule, args=(key, "scan_nm_ersatz_wenn_aushilfe_fehlt", bool))
    with right:
        with st.container(border=True):
            st.markdown("**🏬 Schalter**")
            st.caption("Eine Person pro Tag, ganztags. Wer Schalter oder Reserve macht: «Wer darf was».")
            c1, c2 = st.columns(2)
            _rule_select(c1, "Reserve einsetzen …", "schalter_reserve_modus", SCHALTER_RESERVE_MODI, ver,
                         help="vor / nach der anderen Gruppe, wenn aus der aktiven Gruppe niemand frei ist · nie")
            _rule_select(c2, "Code im Plan", "schalter_code", SCHALTER_CODES, ver,
                         help="April-Liste: ERF4/SCH · seit Mai im Plan: ERF7/SCH")
        with st.container(border=True):
            st.markdown("**🧾 HUB und RX Abo**")
            c1, c2 = st.columns(2)
            _rule_number(c1, "HUB-Halbtage pro Tag", "hub_pro_tag", 0, 2, ver,
                         help="2 = eine Person am Vormittag + eine andere am Nachmittag · 1 = nur Vormittag · 0 = aus")
            _rule_number(c2, "RX Abo – Halbtage pro Woche", "rx_abo_pro_woche", 0, 10, ver,
                         help="höchstens einer pro Tag · 0 = aus")
        with st.container(border=True):
            st.markdown("**🗂️ Scanning (Stammteam)**")
            _rule_number(st, "höchstens Halbtage pro Person und Woche", "scan_max_halbtage_pro_woche", 0, 10, ver)
            key = wk(f"rule_scan_anteil_{ver}")
            p = st.slider("Wie oft wird ein freier Halbtag zu «scanning»?", 0.0, 1.0, float(rules["scan_anteil"]), 0.05,
                          key=key, on_change=_on_rule, args=(key, "scan_anteil", float))
            st.caption("≈ nie" if p == 0 else f"≈ jeder {max(1, round(1 / p))}. freie Halbtag bei Personen mit Scanning")
    with st.expander("🔒 Fest im Programm eingebaut (Änderung braucht eine Anpassung im Code)"):
        st.markdown("\n".join(f"- {r}" for r in FESTE_REGELN))
        st.markdown("**Nur manuell** (nie automatisch, aber im Plan-Editor wählbar): " + ", ".join(NUR_MANUELL) +
                    ". Vorbezüge (VBZ) sind als Rolle hinterlegt, werden aber nicht geplant.")


# ── 📋 Aufgaben & Besetzung (KSV) ─────────────────────────────────────

def _on_tasks_edit(key: str, codes: list[str], auto: bool) -> None:
    cfg = _cfg_copy()
    removed, renames, moves, errors = [], [], [], []
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        code = codes[int(ridx)]
        t = tp.task_by_code(cfg, code)
        for col, val in changes.items():
            if col == "entfernen" and _b(val):
                removed.append(code)
            elif col == "Code":
                new = _s(val)
                if not new:
                    errors.append("Code darf nicht leer sein")
                elif new != code:
                    renames.append((code, new))
            elif col == "Bedeutung":
                t["bedeutung"] = _s(val)
            elif col == "Farbe" and _s(val):
                t["farbe"] = farb_hex(val, t["farbe"])
            elif col == "Dauer" and _s(val) in tp.DAUER:
                t["dauer"] = _s(val)
                if t["dauer"] in tp.AUTO_DAUER and not any(t["anzahl"].values()):
                    t["anzahl"] = {dk: 1 for dk in DAY_KEYS}
            elif col in DAY_KEYS:
                t["anzahl"][col] = max(0, int(float(val or 0)))
            elif col == "Nr." and _s(val):
                moves.append((code, int(float(val))))
    if errors:
        _fail(errors)
        return
    if moves:
        autos = [t for t in cfg["aufgaben"] if t["dauer"] in tp.AUTO_DAUER]
        rest = [t for t in cfg["aufgaben"] if t["dauer"] not in tp.AUTO_DAUER]
        for code, pos in moves:
            t = next((x for x in autos if x["code"] == code), None)
            if t:
                autos.remove(t)
                autos.insert(max(0, min(len(autos), pos - 1)), t)
        cfg["aufgaben"] = autos + rest
    for code in removed:
        tp.remove_task(cfg, code)
    for old, new in renames:
        if old in removed:
            continue
        if new.lower() in {c.lower() for c in tp.task_codes(cfg) if c != old}:
            _fail([f"Code «{new}» gibt es schon"])
            return
        tp.rename_task(cfg, old, new)
    msg = "Aufgabe entfernt" if removed else ("Reihenfolge geändert – Nr. 1 wird zuerst besetzt" if moves else "")
    commit_cb(cfg, msg=msg, remount=True)


def _on_fill(key: str) -> None:
    cfg = _cfg_copy()
    v = st.session_state[key]
    cfg["regeln"]["auffuellen"] = "" if v == "– leer lassen –" else v
    commit_cb(cfg)


def _on_frei(key: str) -> None:
    cfg = _cfg_copy()
    cfg["regeln"]["frei_anzeigen"] = bool(st.session_state[key])
    commit_cb(cfg)


def _on_custom_color(code_key: str, color_key: str) -> None:
    cfg = _cfg_copy()
    t = tp.task_by_code(cfg, st.session_state[code_key])
    if t:
        t["farbe"] = tp.norm_hex(st.session_state[color_key], t["farbe"])
        commit_cb(cfg, msg=f"Farbe von {t['code']} geändert", remount=True)


def section_tasks(ver: int) -> None:
    cfg = ts["cfg"]
    autos = tp.auto_tasks(cfg)
    manuals = [t for t in tp.task_list(cfg) if t["dauer"] not in tp.AUTO_DAUER]
    fill = cfg["regeln"].get("auffuellen") or ""
    st.markdown(
        '<div class="explain"><b>So plant der Generator</b><ol>'
        '<li>Abwesenheiten aus Schritt 1 und <b>feste Einsätze</b> (z. B. Quality) zuerst.</li>'
        '<li>Dann die Aufgaben unten <b>von Nr. 1 abwärts</b>: pro Tag so viele Personen wie in Mo–Fr steht – '
        'gleichmässig verteilt unter allen, die laut «Wer darf was» dürfen.</li>'
        f'<li>Alle übrigen Halbtage bekommen <b>{html_escape(fill or "nichts (leer)")}</b> (Auffüllen).</li>'
        '<li>Nicht-Arbeitstage laut Anwesenheit stehen als <b>Frei</b> im Plan.</li></ol></div>',
        unsafe_allow_html=True)

    st.markdown("**Automatisch besetzen** – Anzahl Personen pro Tag (ganzer Tag) bzw. pro Halbtag")
    all_colors = [t["farbe"] for t in tp.task_list(cfg)]
    color_opts = farb_options(*all_colors)
    df = pd.DataFrame([{"Nr.": i + 1, "Code": t["code"], "Bedeutung": t.get("bedeutung", ""),
                        "Farbe": farb_label(t["farbe"]), "Dauer": t["dauer"],
                        **{dk: t["anzahl"][dk] for dk in DAY_KEYS}, "entfernen": False}
                       for i, t in enumerate(autos)],
                      columns=["Nr.", "Code", "Bedeutung", "Farbe", "Dauer"] + DAY_KEYS + ["entfernen"])
    col_config = {
        "Nr.": st.column_config.NumberColumn("Nr.", min_value=1, max_value=max(1, len(autos)), step=1, width=48,
                                             help="Wichtigkeit: Nr. 1 wird zuerst besetzt, wenn zu wenig Leute da sind"),
        "Code": st.column_config.TextColumn("Code", width=100, help="So steht es im Plan, z. B. «KSV1»"),
        "Bedeutung": st.column_config.TextColumn("Bedeutung", width="medium", help="Erklärung in der Legende"),
        "Farbe": st.column_config.SelectboxColumn("Farbe", options=color_opts, required=True, width=105),
        "Dauer": st.column_config.SelectboxColumn(
            "Dauer", options=tp.DAUER, required=True, width=115,
            help=" · ".join(f"{k}: {v}" for k, v in tp.DAUER_HILFE.items())),
        "entfernen": st.column_config.CheckboxColumn("entfernen", width="small", help="Anhaken = Aufgabe löschen"),
    }
    for dk in DAY_KEYS:
        col_config[dk] = st.column_config.NumberColumn(dk, min_value=0, max_value=10, step=1, width=52,
                                                       help=f"Personen am {dk} (ganzer Tag) bzw. pro Halbtag")
    if autos:
        key = wk(f"tasks_auto_{ver}")
        st.data_editor(df, key=key, column_config=col_config, hide_index=True, num_rows="fixed", width="stretch",
                       height=_editor_height(len(df)), on_change=_on_tasks_edit,
                       args=(key, [t["code"] for t in autos], True))
    else:
        st.info("Noch keine Aufgabe wird automatisch geplant – unten «➕ Neue Aufgabe» oder bei einem Code "
                "die Dauer von «nur manuell» auf «ganzer Tag» / «Halbtag» stellen.")

    c1, c2 = st.columns(2, vertical_alignment="bottom")
    fill_opts = ["– leer lassen –"] + tp.task_codes(cfg)
    key = wk(f"fill_{ver}")
    c1.selectbox("Freie Halbtage auffüllen mit", fill_opts,
                 index=fill_opts.index(fill) if fill in fill_opts else 0, key=key, on_change=_on_fill, args=(key,),
                 help="Was jede anwesende Person bekommt, wenn keine andere Aufgabe für sie übrig ist")
    key = wk(f"frei_{ver}")
    c2.checkbox("Nicht-Arbeitstage (laut Anwesenheit) als «Frei» eintragen",
                bool(cfg["regeln"].get("frei_anzeigen", True)), key=key, on_change=_on_frei, args=(key,),
                help="Aus = diese Halbtage bleiben im Plan leer")

    st.markdown("**Weitere Codes** – plant der Generator nie; sie kommen über feste Einsätze, Schritt 1 "
                "(«Aufgabe setzen») oder «Plan bearbeiten» in den Plan")
    if manuals:
        df2 = pd.DataFrame([{"Code": t["code"], "Bedeutung": t.get("bedeutung", ""), "Farbe": farb_label(t["farbe"]),
                             "Dauer": t["dauer"], "entfernen": False} for t in manuals],
                           columns=["Code", "Bedeutung", "Farbe", "Dauer", "entfernen"])
        key = wk(f"tasks_manual_{ver}")
        st.data_editor(df2, key=key, hide_index=True, num_rows="fixed", width="stretch",
                       height=_editor_height(len(df2)), on_change=_on_tasks_edit,
                       args=(key, [t["code"] for t in manuals], False),
                       column_config={k: v for k, v in col_config.items() if k in ("Code", "Bedeutung", "Farbe",
                                                                                  "Dauer", "entfernen")})

    with st.form(wk(f"add_task_{ver}"), border=True):
        st.markdown("**➕ Neue Aufgabe**")
        c = st.columns([1.1, 2, 1.1, 1.3, 1.1, 1.3], vertical_alignment="bottom")
        code = c[0].text_input("Code", placeholder="z. B. Redcare")
        bed = c[1].text_input("Bedeutung", placeholder="z. B. Redcare-Kundendienst")
        farbe = c[2].selectbox("Farbe", list(tp.FARBEN))
        dauer = c[3].selectbox("Dauer", tp.DAUER, help=" · ".join(f"{k}: {v}" for k, v in tp.DAUER_HILFE.items()))
        anzahl = c[4].number_input("Personen pro Tag", 0, 10, 1, help="Gilt für Mo–Fr; einzeln in der Tabelle änderbar")
        alle = st.checkbox("Alle aktiven Personen dürfen diese Aufgabe («ja»)", value=True,
                           help="Danach unter «Wer darf was» anpassen")
        if c[5].form_submit_button("Hinzufügen", type="primary", width="stretch"):
            code = code.strip()
            if not code:
                st.error("Bitte einen Code eingeben.")
            else:
                new = _cfg_copy()
                new["aufgaben"].append({"code": code, "bedeutung": bed.strip(), "farbe": tp.FARBEN[farbe],
                                        "dauer": dauer, "anzahl": {dk: int(anzahl) for dk in DAY_KEYS}})
                if alle and dauer in tp.AUTO_DAUER:
                    for m in new["mitarbeiter"]:
                        if m.get("aktiv", True):
                            tp.set_stufe(m, code, "ja")
                errors = try_commit(new, msg=f"Aufgabe {code} hinzugefügt", remount=True)
                if errors:
                    st.error(" · ".join(errors))
                else:
                    st.rerun()

    items = tp.legend_items(cfg)
    lc1, lc2 = st.columns([5, 1.4], vertical_alignment="center")
    lc1.markdown("**So sieht es im Plan aus:** " + "".join(
        f'<span class="legend-item" style="margin-right:10px;"><span class="pill" style="background:#{c}; color:{_fg_for(c)};">'
        f'{html_escape(code)}</span></span>' for code, _e, c in items), unsafe_allow_html=True)
    with lc2.popover("🎨 Eigene Farbe", width="stretch"):
        codes = tp.task_codes(cfg)
        if codes:
            ck = wk(f"cc_code_{ver}")
            sel = st.selectbox("Aufgabe", codes, key=ck)
            cur = tp.task_by_code(cfg, sel)["farbe"]
            pk = wk(f"cc_pick_{ver}_{sel}")
            st.color_picker("Farbe", f"#{cur}", key=pk, on_change=_on_custom_color, args=(ck, pk))
    with st.expander("🔒 Fest eingebaut (Änderung braucht eine Anpassung im Code)"):
        st.markdown("\n".join(f"- {r}" for r in FESTE_REGELN_KSV))


# ── 🗂️ BackUp-Team ────────────────────────────────────────────────────

def _on_backup_edit(key: str, base: list[dict]) -> None:
    rows = [dict(b) for b in base]
    removed, errors = set(), []
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        i = int(ridx)
        for col, val in changes.items():
            if col == "entfernen" and _b(val):
                removed.add(i)
            elif col == "Name":
                if _s(val):
                    rows[i]["name"] = _s(val)
                else:
                    errors.append("Name darf nicht leer sein")
            elif col == "Standard":
                rows[i]["standard"] = _s(val)
    if errors:
        _fail(errors)
        return
    cfg = _cfg_copy()
    cfg["backup_team"] = [b for i, b in enumerate(rows) if i not in removed]
    commit_cb(cfg, msg="BackUp-Zeile entfernt" if removed else "", remount=True)


def section_backup(ver: int) -> None:
    cfg = ts["cfg"]
    st.caption("Zeilen unterhalb des Teams im Plan – werden nie automatisch eingeplant. «Standard» steht in jeder "
               "Zelle (z. B. ERP beim SPP-Team); einzelne Einträge über Schritt 1 mit «📌 Aufgabe setzen».")
    with st.form(wk(f"add_backup_{ver}"), border=True):
        c = st.columns([3, 2, 1.3], vertical_alignment="bottom")
        name = c[0].text_input("Name", placeholder="z. B. Petra M.")
        standard = c[1].text_input("Standard-Eintrag (optional)", placeholder="z. B. ERP")
        if c[2].form_submit_button("➕  Hinzufügen", type="primary", width="stretch"):
            if not name.strip():
                st.error("Bitte einen Namen eingeben.")
            else:
                new = _cfg_copy()
                new["backup_team"].append({"short": E.new_short(name, new), "name": name.strip(),
                                           "standard": standard.strip()})
                errors = try_commit(new, msg=f"{name.strip()} hinzugefügt", remount=True)
                if errors:
                    st.error(" · ".join(errors))
                else:
                    st.rerun()
    base = [dict(b) for b in cfg["backup_team"]]
    if not base:
        st.info("Kein BackUp-Team – der Plan zeigt nur das Team.")
        return
    df = pd.DataFrame([{"Name": b.get("name") or b["short"], "Standard": b.get("standard", ""), "entfernen": False}
                       for b in base], columns=["Name", "Standard", "entfernen"])
    key = wk(f"backup_{ver}")
    st.data_editor(df, key=key, hide_index=True, num_rows="fixed", width="stretch", height=_editor_height(len(df)),
                   on_change=_on_backup_edit, args=(key, base),
                   column_config={"Name": st.column_config.TextColumn("Name", width="medium"),
                                  "Standard": st.column_config.TextColumn("Standard-Eintrag", width="medium"),
                                  "entfernen": st.column_config.CheckboxColumn("entfernen", width="small")})


# ── Seite ─────────────────────────────────────────────────────────────

def _keep_section() -> None:
    """Segment-Auswahl lässt sich abwählen – dann beim letzten Bereich bleiben."""
    k = wk("team_section")
    if not st.session_state.get(k):
        st.session_state[k] = st.session_state.get(wk("team_section_last"), TEAM_SECTIONS[TEAM][0])
    st.session_state[wk("team_section_last")] = st.session_state[k]


def render_team_page() -> None:
    ver = ts["cfg_version"]
    sections = TEAM_SECTIONS[TEAM]
    st.markdown(f'<div class="page-head"><h2>Team & Regeln · {TEAM}</h2><div class="sub">{n_active} von {n_total} '
                'Personen aktiv · alles, was der Generator über das Team weiss. Änderungen gelten sofort für den '
                'nächsten Plan – «Speichern» macht sie dauerhaft.</div></div>', unsafe_allow_html=True)
    render_save_bar()
    errors = st.session_state.pop(wk("team_error"), None)
    if errors:
        st.error("Nicht übernommen: " + " · ".join(errors))
    hints = E.config_hints(ts["cfg"])
    if hints:
        st.warning("Hinweis: " + " · ".join(hints))

    k = wk("team_section")
    if st.session_state.get(k) not in sections:
        st.session_state[k] = sections[0]
    section = st.segmented_control("Bereich", sections, key=k, on_change=_keep_section,
                                   label_visibility="collapsed") or sections[0]
    with st.expander("❓  Wie ändere ich …?"):
        st.markdown("| Ich möchte … | So geht's |\n|---|---|\n" +
                    "\n".join(f"| {was} | {wie} |" for was, wie in HILFE[TEAM]))
    {
        "👤 Personen": section_people,
        "🎯 Wer darf was": section_roles_ksc if IS_KSC else section_roles_ksv,
        "📌 Feste Einsätze": section_fixed,
        "🔄 Rotationen": section_rotation,
        "⚖️ Regeln": section_rules,
        "📋 Aufgaben & Besetzung": section_tasks,
        "🗂️ BackUp-Team": section_backup,
    }[section](ver)


# ══════════════════════════════════════════════════════════════════════
#  Schritt 1 – Planung
# ══════════════════════════════════════════════════════════════════════

def _set_days(days: list[int]) -> None:
    st.session_state[wk("f_days")] = days


def _add_entry_cb() -> None:
    """Eintrag aus dem Formular übernehmen und das Formular für den nächsten Eintrag leeren."""
    ss = st.session_state
    days = sorted(DAY_KEYS.index(d) for d in (ss.get(wk("f_days")) or []))
    person = employee_options().get(ss.get(wk("f_employee")) or "")
    art = ss.get(wk("f_type")) or T.gruende[0]
    task = _s(ss.get(wk("f_task"))) if art == "Custom" else ""
    if not person:
        ss[wk("form_msg")] = "Bitte eine Person wählen."
        return
    if not days:
        ss[wk("form_msg")] = "Bitte mindestens einen Tag wählen."
        return
    if art == "Custom" and not task:
        ss[wk("form_msg")] = "Bitte die Aufgabe wählen, die fest eingetragen werden soll."
        return
    slot = ss.get(wk("f_slot")) or "Ganzer Tag"
    ts["entries"].append({
        "name": person, "type": art, "days": days,
        "slots": [0, 1] if slot == "Ganzer Tag" else [0] if slot == "Vormittag" else [1],
        "note": _s(ss.get(wk("f_note"))), "task": task,
        "jahr": info["iso_year"], "kw_start": info["kw"], "weeks": max(1, int(ss.get(wk("f_weeks")) or 1)),
    })
    persist_entries()
    ss[wk("f_days")] = []
    ss[wk("f_note")] = ""
    ss[wk("f_weeks")] = 1
    st.toast(f"{disp(person)}: {GRUENDE[art][0]} eingetragen", icon="✅")


def _remove_entry_cb(idx: int) -> None:
    if 0 <= idx < len(ts["entries"]):
        e = ts["entries"].pop(idx)
        persist_entries()
        st.toast(f"{disp(e.get('name', ''))}: {GRUENDE.get(e.get('type'), (e.get('type'), ''))[0]} gelöscht", icon="🗑️")


def _purge_cb() -> None:
    ts["entries"] = [e for e in ts["entries"] if entry_span(e)[1] >= _as_date(info["monday"])]
    persist_entries()


def _save_entries_cb() -> None:
    st.session_state.save_msgs = save_everything(push_github=bool(github_conf()))


def _entry_row(idx: int, e: dict, active: bool) -> None:
    art = e.get("type", "Custom")
    label = GRUENDE.get(art, (art, ""))[0]
    if art == "Custom" and e.get("task"):
        label += f" «{e['task']}»"
    note = e.get("note") if e.get("note") and e.get("note") != e.get("task") else ""
    note_html = f'<span class="entry-note">· „{html_escape(note)}"</span>' if note else ""
    dot = {"Home Office": "HO"}.get(art, art)
    tag = '<span class="entry-tag" title="Zeitraum aus dem 📆 Jahresplan">📆</span>' if entry_range(e) else ""
    l, r = st.columns([10, 1.3], vertical_alignment="center")
    l.markdown(
        f'<div class="entry-row" {"" if active else "style=opacity:0.5;"}><span class="entry-dot {html_escape(dot)}"></span>'
        f'<span class="entry-name">{html_escape(disp(e["name"]))}</span>'
        f'<span class="entry-meta">· {html_escape(label)} · {html_escape(entry_when_label(e))}</span>{note_html}{tag}</div>',
        unsafe_allow_html=True,
    )
    r.button("✕", key=wk(f"rm_{idx}_{len(ts['entries'])}"), width="stretch", on_click=_remove_entry_cb, args=(idx,),
             help="Eintrag löschen" + (" (ganzer Zeitraum)" if entry_range(e) else ""))


def _feiertag_rows(monday, tage: dict) -> None:
    for d, txt in tage.items():
        st.markdown(f'<div class="entry-row fe-row"><span class="entry-dot Feiertag"></span>'
                    f'<span class="entry-name">🎉 {html_escape(txt)}</span><span class="entry-meta">· '
                    f'{datum_label(_as_date(monday) + timedelta(days=d))} · ganzes Team frei</span>'
                    f'<span class="entry-tag" title="Feiertag aus dem 📆 Jahresplan">📆</span></div>', unsafe_allow_html=True)


def render_planung() -> None:
    st.markdown(
        f'<div class="howto"><span class="howto-step"><b>1</b>Woche oben wählen – jetzt KW {info["kw"]}</span>'
        '<span class="howto-arrow">→</span><span class="howto-step"><b>2</b>Abwesenheiten &amp; Wünsche erfassen</span>'
        '<span class="howto-arrow">→</span><span class="howto-step"><b>3</b>oben rechts «Plan generieren»</span>'
        '<span class="howto-arrow">→</span><span class="howto-step"><b>4</b>prüfen, anpassen, Excel herunterladen</span>'
        '</div>', unsafe_allow_html=True)
    left, right = st.columns([1.2, 1], gap="medium")

    with left:
        with st.container(border=True):
            st.markdown('<div class="card-head"><div><h2 style="font-size:15.5px; margin:0;">Abwesenheit oder Sonderwunsch</h2>'
                        f'<div style="font-size:12.5px; color:var(--text-mute);">gilt ab KW {info["kw"]} · auch für das '
                        'BackUp-Team</div></div></div>', unsafe_allow_html=True)
            opts = list(employee_options())
            if st.session_state.get(wk("f_employee")) not in opts:
                st.session_state[wk("f_employee")] = opts[0] if opts else None
            st.markdown('<div class="field-label">Wer?</div>', unsafe_allow_html=True)
            st.selectbox("Person", opts, key=wk("f_employee"), label_visibility="collapsed")

            st.markdown('<div class="field-label">Was?</div>', unsafe_allow_html=True)
            if st.session_state.get(wk("f_type")) not in T.gruende:
                st.session_state[wk("f_type")] = T.gruende[0]
            art = st.segmented_control("Grund", list(T.gruende), key=wk("f_type"), label_visibility="collapsed",
                                       format_func=lambda g: GRUENDE[g][0]) or T.gruende[0]
            st.caption(GRUENDE[art][1])
            if art == "Custom":
                task_opts = [t for t in team_task_options(ts["cfg"]) if t and t not in ("KRANK", "FERIEN", "Ferien", "Frei")]
                st.selectbox("Welche Aufgabe?", task_opts, key=wk("f_task"), accept_new_options=True,
                             help="Aus der Liste wählen oder einen eigenen Code eintippen (z. B. ERP)")

            fe_hint = "".join(f" · 🎉 {DAY_KEYS[d]} {html_escape(t)}" for d, t in info["feiertage"].items())
            st.markdown(f'<div class="field-label">Wann? <span>Tage – mehrere möglich · {info["monday"]:%d.%m.} – '
                        f'{info["friday"]:%d.%m.}{fe_hint}</span></div>', unsafe_allow_html=True)
            d1, d2, d3 = st.columns([2.6, 1.2, 0.8], vertical_alignment="center")
            with d1:
                st.segmented_control("Tage", DAY_KEYS, selection_mode="multi", key=wk("f_days"),
                                     label_visibility="collapsed")
            d2.button("ganze Woche", key=wk("day_all"), width="stretch", on_click=_set_days, args=(list(DAY_KEYS),))
            d3.button("keine", key=wk("day_none"), width="stretch", on_click=_set_days, args=([],))

            c1, c2 = st.columns([2.6, 2])
            with c1:
                st.markdown('<div class="field-label">Halbtag</div>', unsafe_allow_html=True)
                if st.session_state.get(wk("f_slot")) not in ("Ganzer Tag", "Vormittag", "Nachmittag"):
                    st.session_state[wk("f_slot")] = "Ganzer Tag"
                st.segmented_control("Halbtag", ["Ganzer Tag", "Vormittag", "Nachmittag"], key=wk("f_slot"),
                                     label_visibility="collapsed")
            with c2:
                st.markdown('<div class="field-label">Wie viele Wochen? <span>ab KW ' + str(info["kw"]) + '</span></div>',
                            unsafe_allow_html=True)
                st.number_input("Wochen", min_value=1, max_value=52, step=1, key=wk("f_weeks"),
                                label_visibility="collapsed", help="z. B. 3 = Ferien über drei Wochen")

            st.markdown('<div class="field-label">Notiz <span>optional · bei «Termin» steht sie im Plan</span></div>',
                        unsafe_allow_html=True)
            st.text_input("Notiz", key=wk("f_note"), label_visibility="collapsed",
                          placeholder="z. B. Arzttermin 14:30")

            # Vorschau: was wird eingetragen?
            ss = st.session_state
            sel_days = sorted(DAY_KEYS.index(d) for d in (ss.get(wk("f_days")) or []))
            who = ss.get(wk("f_employee")) or "–"
            what = GRUENDE[art][0] + (f" «{_s(ss.get(wk('f_task')))}»" if art == "Custom" and ss.get(wk("f_task")) else "")
            weeks = max(1, int(ss.get(wk("f_weeks")) or 1))
            slot = {"Ganzer Tag": "ganzer Tag"}.get(ss.get(wk("f_slot")) or "Ganzer Tag", ss.get(wk("f_slot")))
            when = (f"{days_label(sel_days)} ({slot})" if sel_days
                    else "<i>noch kein Tag gewählt</i>")
            st.markdown(f'<div class="preview-line">→ <b>{html_escape(who)}</b> · {html_escape(what)} · {when} · '
                        f'{kw_label(info["iso_year"], info["kw"], weeks)}</div>', unsafe_allow_html=True)
            msg = st.session_state.pop(wk("form_msg"), None)
            if msg:
                st.warning(msg)
            st.button("＋  Eintrag hinzufügen", key=wk("add_entry"), type="primary", width="stretch",
                      on_click=_add_entry_cb)

    with right:
        entries = ts["entries"]
        monday = _as_date(info["monday"])
        now = [(i, e) for i, e in enumerate(entries) if entry_days(e, monday)]
        later = sorted(((i, e) for i, e in enumerate(entries) if not entry_days(e, monday)),
                       key=lambda ie: entry_span(ie[1])[0])
        with st.container(border=True):
            st.markdown(f'<div class="card-head"><div><h2 style="font-size:15.5px; margin:0;">Einträge</h2>'
                        f'<div style="font-size:12.5px; color:var(--text-mute);">{len(now)} für KW {info["kw"]} · '
                        f'{len(entries)} insgesamt</div></div></div>', unsafe_allow_html=True)
            if not entries and not info["feiertage"]:
                st.markdown('<div class="empty-state"><div class="icon">📅</div><div class="ttl">Keine Abwesenheiten erfasst.</div>'
                            '<div>Links eintragen – oder direkt «Plan generieren», wenn alle da sind. '
                            'Ferien übers Jahr: 📆 Jahresplan.</div></div>',
                            unsafe_allow_html=True)
            else:
                st.markdown(f'<div class="group-label">gilt für KW {info["kw"]}</div>', unsafe_allow_html=True)
                _feiertag_rows(monday, info["feiertage"])
                if not now and not info["feiertage"]:
                    st.caption("Keine – in dieser Woche sind alle laut Stammdaten da.")
                for i, e in now:
                    _entry_row(i, e, True)
                if later:
                    with st.expander(f"Andere Wochen ({len(later)})"):
                        for i, e in later:
                            _entry_row(i, e, False)
                b1, b2 = st.columns(2)
                b1.button("🗑  Vergangene aufräumen", key=wk("purge_old"), width="stretch", on_click=_purge_cb,
                          help="Löscht Einträge, die vor der angezeigten Woche enden")
                b2.button("💾  Einträge speichern", key=wk("save_entries"), width="stretch", disabled=not entries_dirty,
                          on_click=_save_entries_cb, help="Lokal und – falls konfiguriert – ins GitHub-Repo")
                for ok, m in st.session_state.pop("save_msgs", []):
                    (st.success if ok else st.warning)(m)
        st.markdown(f'<div class="team-hint">👥 <b>{n_active}</b> von {n_total} Personen aktiv'
                    f'{" · <b>ungespeicherte Team-Änderungen</b>" if cfg_dirty else ""}. '
                    'Personen, Teilzeit, wer darf was und Regeln: <b>⚙ Team & Regeln</b> (oben).<br>'
                    '🏖️ Ferien über mehrere Wochen und 🎉 Feiertage: <b>📆 Jahresplan</b> (oben) – '
                    'sie erscheinen automatisch in jeder betroffenen Woche.</div>',
                    unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════
#  Schritt 2 – Ergebnis
# ══════════════════════════════════════════════════════════════════════

def render_banner(kind: str, title: str, items: list[str]) -> None:
    lis = "".join(f"<li>{html_escape(i)}</li>" for i in items)
    mark = {"ok": "✓", "warn": "!", "error": "!", "info": "i"}[kind]
    st.markdown(f'<div class="banner banner-{kind}"><div class="check">{mark}</div><div><h4>{html_escape(title)}</h4>'
                f'{"<ul>" + lis + "</ul>" if items else ""}</div></div>', unsafe_allow_html=True)


def render_plan_editor(res: dict) -> None:
    sched = ts.get("sched")
    if sched is None:
        return
    slot_cols = [f"{DAY_KEYS[d]} {SLOT_KEYS[s]}" for d in range(5) for s in range(2)]
    with st.expander("✏️  Plan bearbeiten · einzelne Halbtage ändern (Team und BackUp-Team)", expanded=False):
        c1, c2 = st.columns([4, 1.2], vertical_alignment="center")
        c1.caption("Pro Halbtag eine Aufgabe wählen, leer = frei. «Übernehmen» aktualisiert Plan, Regelprüfung und "
                   "Excel. Achtung: Wochenwechsel (‹ ›) oder neu generieren erstellt den Plan neu.")
        freitext = c2.toggle("Freitext", key=wk("editor_freitext"),
                             help="Statt Auswahlliste eigenen Text eintippen, z. B. «KSV2 10.00» oder «Arzt»")
        all_rows = res["rows"] + res["backup_rows"]
        df_rows = []
        for r in all_rows:
            row = {"Name": r["full_name"], "%": f"{r['pct']}%" if r["pct"] is not None else "BackUp"}
            for i, col in enumerate(slot_cols):
                t = r["cells"][i]["task"] or ""
                row[col] = t[1:] if t.startswith("*") else t
            df_rows.append(row)
        df = pd.DataFrame(df_rows, columns=["Name", "%"] + slot_cols)
        options = list(team_task_options(ts["cfg"]))
        for r in df_rows:
            for col in slot_cols:
                if r[col] not in options:
                    options.append(r[col])
        col_config = {"Name": st.column_config.TextColumn("Name", disabled=True, width="medium"),
                      "%": st.column_config.TextColumn("%", disabled=True, width="small")}
        for col in slot_cols:
            col_config[col] = (st.column_config.TextColumn(col, width="small") if freitext else
                               st.column_config.SelectboxColumn(col, options=options, required=False, width="small"))
        edited = st.data_editor(df, column_config=col_config, hide_index=True, num_rows="fixed", width="stretch",
                                key=wk(f"plan_editor_{res['iso_year']}_{res['kw']}_{'t' if freitext else 's'}"))
        if not st.button("✓  Änderungen übernehmen", key=wk(f"apply_edits_{res['kw']}"), type="primary"):
            return
        changes = 0
        backups = set(backup_shorts())
        for ridx, r in enumerate(all_rows):
            short = r["short"]
            for cidx, col in enumerate(slot_cols):
                new_val = _s(edited.at[ridx, col])
                old_val = r["cells"][cidx]["task"] or ""
                if old_val.startswith("*"):
                    old_val = old_val[1:]
                if new_val == old_val:
                    continue
                d, s = divmod(cidx, 2)
                if short in backups:
                    sched.backup_tasks[(short, d, s)] = new_val
                elif short in sched.schedule:
                    sched.schedule[short][(d, s)] = new_val
                changes += 1
        if changes == 0:
            st.info("Keine Änderungen erkannt.")
            return
        if IS_KSC:  # Zähler neu aufbauen, damit die Regelprüfung stimmt
            sched.tel_count = {d: {s: 0 for s in range(2)} for d in range(5)}
            sched.abkl_count = {d: {s: 0 for s in range(2)} for d in range(5)}
            sched.erf8_assigned = {d: False for d in range(5)}
            for name, cells in sched.schedule.items():
                for (d, s), t in cells.items():
                    sched._count(d, s, t)
        refresh_result_from_sched(sched)
        st.toast(f"{changes} Änderung(en) übernommen.", icon="✅")
        st.rerun()


def _plan_table_html(res: dict) -> str:
    tvc = ak.tv_colors(ts["cfg"]) if IS_KSC else {}
    monday = res["monday"]
    legend = "TV · Tagesverantwortung · dunkelblau = PHC-Liste" if IS_KSC else "Farben wie im Excel"
    parts = ['<div class="plan-wrap"><table class="plan"><thead><tr>',
             f'<th class="team-head" colspan="2" rowspan="2"><div class="ttl">TEAM {TEAM}</div>'
             f'<div class="legend">{legend}</div></th>']
    fe_tage = res.get("feiertage") or {}
    for d in range(5):
        datum = (monday + timedelta(days=d)).strftime("%d.%m.")
        pill = ""
        if IS_KSC and str(res["tagesverantwortung"][d]) != ak.FEIERTAG:
            tv = str(res["tagesverantwortung"][d])
            color = tvc.get(tv)
            style = f'style="background:#{color}; color:{_fg_for(color)};"' if color else ""
            pill = (f'<span class="tv-pill{" open" if "OFFEN" in tv else ""}"><span class="tv-tag" {style}>TV</span>'
                    f'<span class="tv-name">{html_escape(tv)}</span></span>')
        if d in fe_tage:
            pill += f'<span class="fe-pill">🎉 {html_escape(fe_tage[d])}</span>'
        parts.append(f'<th class="day-head" colspan="2"><span class="day-name">{DAYS[d].upper()}</span>'
                     f'<span class="day-date">{datum}</span>{pill}</th>')
    parts.append("</tr><tr>")
    for _ in range(5):
        parts.append('<th class="slot-head">VM</th><th class="slot-head">NM</th>')
    parts.append("</tr></thead><tbody>")

    def row_html(r):
        swatch = f'<span class="swatch" style="background:#{r["farbe"]};"></span>' if r.get("farbe") else ""
        out = [f'<tr><td class="name-cell">{swatch}{html_escape(r["full_name"])}</td>',
               f'<td class="pct-cell">{str(r["pct"]) + "%" if r["pct"] is not None else ""}</td>']
        for c in r["cells"]:
            task = c["task"] or ""
            note = f'<span class="pill-note">{html_escape(c["note"])}</span>' if c.get("note") else ""
            if not task and not c.get("note"):
                out.append('<td class="task-cell"></td>')
            elif c["is_phc"]:
                out.append(f'<td class="task-cell"><span class="pill pill-phc">{html_escape(task)}</span>{note}</td>')
            elif task.startswith("*") or (not task and c.get("note")):
                out.append(f'<td class="task-cell"><span class="pill pill-termin">{html_escape(task[1:] or c["note"])}</span></td>')
            else:
                bg, fg = pill_for_task(task)
                out.append(f'<td class="task-cell"><span class="pill" style="background:{bg}; color:{fg};">'
                           f'{html_escape(task)}</span>{note}</td>')
        out.append("</tr>")
        return "".join(out)

    for r in res["rows"]:
        parts.append(row_html(r))
    if res["backup_rows"]:
        parts.append('<tr><td class="sep-cell" colspan="12">BackUp Team · nur manuell</td></tr>')
        for r in res["backup_rows"]:
            parts.append(row_html(r))
    parts.append("</tbody></table></div>")
    return "".join(parts)


def _ksc_details_html(res: dict) -> str:
    parts = ['<div class="section-title">Tagesrollen</div><table class="roles"><thead><tr><th></th>' +
             "".join(f"<th>{DAY_KEYS[d]}</th>" for d in range(5)) + "</tr></thead><tbody>"]
    for lbl, key in (("Tagesverantwortung", "tagesverantwortung"), ("PHC-Liste (18 Uhr)", "phc_liste"),
                     ("TAGES PA", "tagespa"), ("Postöffnung (PO/SCAN)", "po_liste"), ("Zusatz-PO", "po_zusatz"),
                     ("Schalter", "schalter"), ("ERF9", "erf9")):
        cells = "".join(f'<td class="{"open" if "OFFEN" in str(v) else ""}">{html_escape(v)}</td>' for v in res[key])
        parts.append(f'<tr><td class="lbl">{lbl}</td>{cells}</tr>')
    parts.append("</tbody></table>")
    wa = res["wochenaufgaben"]
    parts.append('<div class="wa-strip">' + "".join(
        f'<div class="wa-card"><div class="wa-icon">{ico}</div><div class="wa-meta"><div class="lbl">{lbl}</div>'
        f'<div class="val">{html_escape(wa[k] or "–")}</div></div></div>'
        for ico, lbl, k in (("📦", "Direkt", "direkt"), ("🆕", "ONB", "onb"), ("💊", "BTM", "btm"))) + "</div>")
    return "".join(parts)


def _ksv_details_html(res: dict) -> str:
    cfg = ts["cfg"]
    parts = ['<div class="section-title">Wer hat was – pro Tag</div><table class="roles"><thead><tr><th></th>' +
             "".join(f"<th>{DAY_KEYS[d]}</th>" for d in range(5)) + "</tr></thead><tbody>"]
    for code, cells in res["overview"]:
        bg, fg = pill_for_task(code)
        lbl = f'<span class="pill" style="background:{bg}; color:{fg}; min-width:0;">{html_escape(code)}</span>'
        tds = "".join(f'<td class="{"open" if "OFFEN" in c or "⚠" in c else ""}">{html_escape(c)}</td>' for c in cells)
        parts.append(f'<tr><td class="lbl">{lbl}</td>{tds}</tr>')
    parts.append("</tbody></table>")
    summary = res.get("summary") or []
    if summary:
        codes = [k for k in summary[0] if k not in ("name", "da")]
        parts.append('<div class="section-title">Verteilung pro Person – Halbtage in dieser Woche</div>'
                     '<table class="roles"><thead><tr><th style="text-align:left;">Person</th><th>anwesend</th>' +
                     "".join(f"<th>{html_escape(c)}</th>" for c in codes) + "</tr></thead><tbody>")
        for row in summary:
            tds = "".join(f"<td>{row[c] or ''}</td>" for c in codes)
            parts.append(f'<tr><td class="lbl">{html_escape(row["name"])}</td><td>{row["da"]}</td>{tds}</tr>')
        parts.append("</tbody></table>")
        fill = cfg["regeln"].get("auffuellen")
        parts.append('<div style="font-size:12px; color:var(--text-mute); margin-top:6px;">Ganztags-Aufgaben zählen '
                     f'2 Halbtage. So siehst du auf einen Blick, ob die Arbeit fair verteilt ist'
                     f'{" – " + html_escape(fill) + " = Auffüllen" if fill else ""}.</div>')
    return "".join(parts)


def _coverage_cards(res: dict) -> None:
    cov = res["coverage"]
    for i in range(0, len(cov), 5):
        cols = st.columns(5)
        for col, c in zip(cols, cov[i:i + 5]):
            if c.get("feiertag"):
                items = []
            elif IS_KSC:
                items = [("TEL", c["tel"], c["tel_target"], c["tel_ok"]), ("ABKL", c["abkl"], c["abkl_target"], c["abkl_ok"])]
            else:
                items = [(x["code"], x["have"], x["need"], x["ok"]) for x in c["items"]]
            empty = (f'<div style="color:#64748b;">🎉 {html_escape(c["feiertag"])}</div>' if c.get("feiertag")
                     else '<div style="color:#94a3b8;">nichts zu besetzen</div>')
            lines = "".join(
                f'<div style="display:flex; justify-content:space-between;"><span>{html_escape(code)}</span>'
                f'<span style="color:{"#16a34a" if ok else "#dc2626"}; font-weight:700;">{have}/{need}</span></div>'
                for code, have, need, ok in items) or empty
            col.markdown(
                f'<div style="background:#f8fafc; border:1px solid #e3e8f0; border-radius:8px; padding:10px 12px; font-size:12px;">'
                f'<div style="display:flex; justify-content:space-between; color:#475569; margin-bottom:4px;">'
                f'<strong style="color:#0b1020;">{c["day"]}</strong><span>{c["slot"]}</span></div>{lines}</div>',
                unsafe_allow_html=True,
            )


def _legend_items() -> list[tuple[str, str, str, str]]:
    """(Code, Erklärung, Hintergrund, Schrift) für die Legende in der App."""
    items = []
    if IS_KSC:
        for row in ak.legend_rows(ts["cfg"]):
            for code, expl in row:
                if not code:
                    continue
                if code == "PHC-Liste":
                    items.append((code, expl, f"#{ak.PHC_COLOR}", "#FFFFFF"))
                else:
                    items.append((code, expl, *pill_for_task(code)))
    else:
        for code, expl, color in tp.legend_items(ts["cfg"]):
            items.append((code, expl, f"#{color}", _fg_for(color)))
    return items


def render_ergebnis() -> None:
    res = ts["result"]
    if not res:
        st.info("Noch kein Plan generiert. Klicke oben rechts auf **Plan generieren**.")
        return

    hard = [i for i in res["issues"] if i.startswith("❌")]
    soft = [i for i in res["issues"] if i.startswith("⚠️")]
    infos = [i for i in res["issues"] if not i.startswith("❌") and not i.startswith("⚠️")]
    if res.get("config_problems"):
        render_banner("warn", "Konfiguration prüfen (Team & Regeln)", res["config_problems"])
    if hard:
        render_banner("error", f"{len(hard)} Regel{'n' if len(hard) != 1 else ''} nicht eingehalten – "
                               "bitte im Plan anpassen («Plan bearbeiten» unten)", hard)
    else:
        render_banner("ok", "Alle Pflicht-Regeln eingehalten.", [])
    if soft:
        render_banner("warn", "Hinweise", soft)
    if res.get("feiertage"):
        render_banner("info", "Feiertag – das ganze Team hat frei, an diesen Halbtagen wird nichts geplant",
                      [f"{DAYS[d]} {(res['monday'] + timedelta(days=d)).strftime('%d.%m.')}: {t}"
                       for d, t in res["feiertage"].items()])
    if res["warnings"]:
        render_banner("warn", "Auswirkungen der erfassten Abwesenheiten", res["warnings"])
    if infos:
        render_banner("info", "Vertretungen und Rotation" if IS_KSC else "Zur Info", infos)

    with st.container(border=True):
        h1, h2, h3 = st.columns([4, 1.3, 1.6], vertical_alignment="center")
        h1.markdown(f'<h2 style="font-size:15.5px; margin:0;">Wochenplan {TEAM} · KW {res["kw"]}</h2>'
                    f'<div style="font-size:12.5px; color:var(--text-mute);">{res["monday_str"]} – {res["friday_str"]}</div>',
                    unsafe_allow_html=True)
        if h2.button("↻  Neu generieren", key=wk("regen"), width="stretch",
                     help="Plan mit dem aktuellen Stand (Einträge, Team & Regeln) neu erstellen – manuelle Änderungen gehen verloren"):
            generate_now()
        h3.download_button("⬇  Excel herunterladen", data=res["excel_bytes"], file_name=res["filename"],
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           width="stretch", key=wk("dl_bottom"),
                           help="KSC: Vorlage mit Legende" if IS_KSC else
                           "Spalten A–L = ein Wochenblock wie im KSV-Jahresplan – direkt dort einfügbar")
        st.markdown(_plan_table_html(res), unsafe_allow_html=True)
        st.markdown(_ksc_details_html(res) if IS_KSC else _ksv_details_html(res), unsafe_allow_html=True)

    render_plan_editor(res)

    with st.expander("Besetzung pro Halbtag · " + ("TEL und ABKL" if IS_KSC else "geplante Aufgaben (Ist/Soll)")):
        _coverage_cards(res)

    with st.expander("Legende · Farben wie im Excel"):
        st.markdown('<div class="legend-wrap">' + "".join(
            f'<span class="legend-item"><span class="pill" style="background:{bg}; color:{fg};">{html_escape(code)}</span>'
            f'{html_escape(expl)}</span>' for code, expl, bg, fg in _legend_items()) + "</div>", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════
#  📆 Jahresplan – Ferien und Feiertage fürs ganze Jahr
# ══════════════════════════════════════════════════════════════════════
#  Grobe Übersicht pro Kalenderwoche. Einträge landen in derselben Liste wie
#  in Schritt 1 (als Zeitraum von–bis) und gelten automatisch für jede
#  betroffene Woche; Feiertage gelten fürs ganze Team.

# Farbe pro Art: (ein Teil der Arbeitstage weg, die ganze Woche weg)
JP_FARBEN = {"Ferien": ("#bae6fd", "#0284c7"), "Frei": ("#a5f3fc", "#0891b2"), "Krank": ("#fecaca", "#dc2626"),
             "andere": ("#ede9fe", "#7c3aed")}
JP_ARTEN_LABEL = {"Ferien": "Ferien", "Frei": "Frei", "Krank": "krank", "andere": "Termin / HO / Aufgabe"}
MONATE = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
FREI_WAHL = {"–": None, "ganzer Tag": "ganz", "Vormittag": "VM", "Nachmittag": "NM"}
FREI_WAHL_LABEL = {v: k for k, v in FREI_WAHL.items()}


def jahr_wochen(jahr: int) -> list[date]:
    """Montage aller Kalenderwochen (ISO) des Jahres – 52 oder 53."""
    d, out = date.fromisocalendar(jahr, 1, 1), []
    while d.isocalendar()[0] == jahr:
        out.append(d)
        d += timedelta(weeks=1)
    return out


def halbe_tage(n: int) -> str:
    """Halbtage → Tage: 4 → '2', 5 → '2½', 1 → '½'."""
    tage, rest = divmod(n, 2)
    return f"{tage or ''}{'½' if rest else ''}" or "0"


def jahres_daten(jahr: int) -> list[dict]:
    """Pro KW: Feiertage, Abwesenheiten pro Person und ob der Plan aufgeht.
    Rechnet für jede kommende Woche den echten Plan durch (1–2 ms pro Woche) – nur bei Änderungen neu."""
    key = hashlib.md5((E.config_to_json(ts["cfg"]) + absences_json() + f"{jahr}|{date.today()}")
                      .encode("utf-8")).hexdigest()
    cache = ts.get("jp_cache")
    if cache and cache[0] == key:
        return cache[1]
    heute = date.today()
    out = []
    for monday in jahr_wochen(jahr):
        vorbei = monday + timedelta(days=4) < heute
        if vorbei:
            inp, hard, knapp = week_inputs(monday), [], []
        else:
            sched, warnings, inp = build_week(monday)
            hard = [i for i in E.validate_schedule(sched) if i.startswith("❌")]
            knapp = [w for w in warnings if w.startswith(("KRITISCH", "Knapp"))]
        moeglich = {n: sum(1 for d in range(5) for s in range(2)
                           if e.is_available(d, s) and (d, s) not in inp["feiertage"])
                    for n, e in inp["employees"].items()}
        out.append({"monday": monday, "kw": monday.isocalendar()[1], "vorbei": vorbei,
                    "feiertage": ft.pro_tag(inp["feiertage"]), "halbtage": inp["halbtage"],
                    "details": inp["details"], "moeglich": moeglich, "hard": hard, "knapp": knapp})
    ts["jp_cache"] = (key, out)
    return out


def _jp_zelle(w: dict, short: str) -> tuple[str, str, str] | None:
    """(Text, Hintergrund, Schrift) für Person × KW – None, wenn nichts eingetragen ist."""
    ht = w["halbtage"].get(short) or {}
    n = sum(ht.values())
    if not n:
        return None
    art = max(ht, key=ht.get)
    hell, kraeftig = JP_FARBEN.get(art, JP_FARBEN["andere"])
    if n >= w["moeglich"].get(short, 0):       # an allen Arbeitstagen der Woche weg
        return halbe_tage(n), kraeftig, "#ffffff"
    return halbe_tage(n), hell, "#0f172a"


def _jp_grid_html(daten: list[dict]) -> str:
    sel = _as_date(info["monday"])
    monate: list[list[int]] = []
    month_start = set()
    for i, w in enumerate(daten):
        m = (w["monday"] + timedelta(days=3)).month     # die KW gehört zum Monat ihres Donnerstags
        if not monate or monate[-1][0] != m:
            monate.append([m, 0])
            month_start.add(i)
        monate[-1][1] += 1

    def cls(i: int, w: dict) -> str:
        c = [x for x, on in (("past", w["vorbei"]), ("sel", w["monday"] == sel), ("ms", i in month_start)) if on]
        return f' class="{" ".join(c)}"' if c else ""

    heute = date.today()
    rtl = daten[0]["monday"].year <= heute.year and heute.isocalendar()[1] > 26 and not daten[-1]["vorbei"] \
        and daten[0]["monday"] <= heute
    h = [f'<div class="jp-wrap{" rtl" if rtl else ""}"><table class="jp"><thead><tr><th class="jp-name"></th>']
    h += [f'<th class="jp-month ms" colspan="{n}">{MONATE[m - 1]}</th>' for m, n in monate]
    h.append('</tr><tr><th class="jp-name">KW</th>')
    for i, w in enumerate(daten):
        tip = f"KW {w['kw']}: {range_label(w['monday'], w['monday'] + timedelta(days=4))}"
        h.append(f'<th{cls(i, w)} title="{html_escape(tip)}">{w["kw"]}</th>')
    h.append('</tr></thead><tbody><tr class="jp-fe"><td class="jp-name">🎉 Feiertage</td>')
    for i, w in enumerate(daten):
        if w["feiertage"]:
            tip = " · ".join(f"{datum_label(w['monday'] + timedelta(days=d))} {t}" for d, t in w["feiertage"].items())
            h.append(f'<td{cls(i, w)} title="{html_escape(tip)}">🎉</td>')
        else:
            h.append(f"<td{cls(i, w)}></td>")
    h.append("</tr>")
    for m in active_members():
        short, name = m["short"], m.get("name") or m["short"]
        h.append(f'<tr><td class="jp-name">{html_escape(name)}</td>')
        for i, w in enumerate(daten):
            z = _jp_zelle(w, short)
            if not z:
                h.append(f"<td{cls(i, w)}></td>")
                continue
            tip = f"{name} · KW {w['kw']}: " + " · ".join(w["details"].get(short, [])) + f" – {z[0]} Arbeitstag(e)"
            h.append(f'<td{cls(i, w)} title="{html_escape(tip)}"><span class="jp-cell" '
                     f'style="background:{z[1]}; color:{z[2]};">{z[0]}</span></td>')
        h.append("</tr>")
    h.append('<tr class="jp-status"><td class="jp-name" title="Für jede kommende Woche wird der Plan mit allen '
             'Einträgen und Feiertagen durchgerechnet">Plan geht auf?</td>')
    for i, w in enumerate(daten):
        if w["vorbei"]:
            h.append(f"<td{cls(i, w)}></td>")
            continue
        if w["hard"]:
            color = "var(--err)"
            tip = (f"KW {w['kw']}: geht nicht auf, {len(w['hard'])} Regel(n) verletzt – im Wochenplan von Hand anpassen: "
                   + " · ".join(x[2:] for x in w["hard"][:4]) + (" …" if len(w["hard"]) > 4 else ""))
        elif w["knapp"]:
            color, tip = "var(--warn)", f"KW {w['kw']}: knapp – " + " · ".join(w["knapp"][:3])
        else:
            color, tip = "var(--ok)", f"KW {w['kw']}: genug Personen da"
        h.append(f'<td{cls(i, w)} title="{html_escape(tip)}"><span class="jp-dot" style="background:{color};"></span></td>')
    h.append("</tr></tbody></table></div>")
    return "".join(h)


def _jp_legend_html() -> str:
    arten = [a for a in ("Ferien", "Frei", "Krank") if a in T.gruende] + ["andere"]
    items = [f'<span><span class="jp-cell" style="background:{JP_FARBEN[a][1]}; color:#fff;">5</span> '
             f'<span class="jp-cell" style="background:{JP_FARBEN[a][0]};">2</span> {JP_ARTEN_LABEL[a]}</span>'
             for a in arten]
    items.append("<span>Zahl = Arbeitstage weg in der KW · kräftig = die ganze Woche · Maus über die Zelle = Details</span>")
    items.append('<span><span class="jp-dot" style="background:var(--ok);"></span> geht auf &nbsp;'
                 '<span class="jp-dot" style="background:var(--warn);"></span> knapp &nbsp;'
                 '<span class="jp-dot" style="background:var(--err);"></span> geht nicht auf (Regeln verletzt)</span>')
    return '<div class="jp-legend">' + "".join(items) + "</div>"


def entry_halbtage(e: dict) -> int:
    """Arbeits-Halbtage, die ein Eintrag betrifft – laut Anwesenheit der Person (Teilzeit),
    ohne Wochenenden und Feiertage. So zählen Ferientage richtig."""
    von, bis = entry_span(e)
    monday, n = von - timedelta(days=von.weekday()), 0
    while monday <= bis:
        emp = E.create_employees(monday.isocalendar()[1], ts["cfg"]).get(e.get("name"))
        fe = ft.woche(ts["feiertage"], monday)
        n += sum(1 for d in entry_days(e, monday) for s in e["slots"]
                 if (d, s) not in fe and (emp is None or emp.is_available(d, s)))
        monday += timedelta(weeks=1)
    return n


def jahresplan_excel(jahr: int, daten: list[dict]) -> bytes:
    """Jahresübersicht zum Ausdrucken: Personen × KW, dazu Einträge und Feiertage als Liste."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def font(size=10, **kw):
        return Font(name="Calibri", size=size, **kw)

    def fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_color.lstrip("#").upper())

    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    wb = Workbook()
    ws = wb.active
    ws.title = f"Jahresplan {jahr}"
    ws.cell(1, 1, f"Jahresplan {jahr} · Ferien & Feiertage · Team {TEAM}").font = font(14, bold=True)
    ws.cell(2, 1, "Zahl = Arbeitstage abwesend in der Kalenderwoche · kräftige Farbe = die ganze Woche · "
                  f"FT = Feiertag · Stand {date.today():%d.%m.%Y}").font = font(9, italic=True, color="64748B")
    r_mon, r_kw, r_fe, r0 = 4, 5, 6, 7
    ws.cell(r_kw, 1, "KW").font = font(bold=True)
    ws.cell(r_fe, 1, "Feiertage").font = font(bold=True)
    last = None
    for i, w in enumerate(daten):
        col = 2 + i
        m = (w["monday"] + timedelta(days=3)).month
        if m != last:
            ws.cell(r_mon, col, MONATE[m - 1]).font = font(9, bold=True)
            last = m
        c = ws.cell(r_kw, col, w["kw"])
        c.font, c.alignment, c.border, c.fill = font(9, bold=True), center, border, fill("F1F5F9")
        c = ws.cell(r_fe, col)
        c.border = border
        if w["feiertage"]:
            c.value, c.alignment, c.fill = "FT", center, fill("E9D5FF")
            c.font = font(8, bold=True, color="6B21A8")
    members = active_members()
    for j, m in enumerate(members):
        row = r0 + j
        c = ws.cell(row, 1, m.get("name") or m["short"])
        c.font, c.border = font(), border
        for i, w in enumerate(daten):
            c = ws.cell(row, 2 + i)
            c.border, c.alignment = border, center
            z = _jp_zelle(w, m["short"])
            if z:
                c.value, c.fill = z[0], fill(z[1])
                c.font = font(9, bold=True, color=z[2].lstrip("#").upper())
    row = r0 + len(members) + 1
    ws.cell(row, 1, "Legende").font = font(bold=True)
    for k, a in enumerate([a for a in ("Ferien", "Frei", "Krank") if a in T.gruende] + ["andere"]):
        for off, (color, txt) in enumerate(((JP_FARBEN[a][1], "ganze Woche"), (JP_FARBEN[a][0], "einzelne Tage"))):
            c = ws.cell(row + 1 + k, 2 + off * 6)
            c.fill, c.border = fill(color), border
            ws.cell(row + 1 + k, 3 + off * 6, f"{JP_ARTEN_LABEL[a]} – {txt}").font = font(9)
    ws.column_dimensions["A"].width = 20
    for i in range(len(daten)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 3.7
    ws.freeze_panes = ws.cell(r0, 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    start, end = daten[0]["monday"], daten[-1]["monday"] + timedelta(days=4)
    ws2 = wb.create_sheet("Einträge")
    head = ["Name", "Was", "Von", "Bis", "Halbtag", "Arbeitstage", "Wochentage", "Notiz"]
    for col, txt in enumerate(head, 1):
        c = ws2.cell(1, col, txt)
        c.font, c.fill, c.border = font(bold=True), fill("E2E8F0"), border
    items = sorted((e for e in ts["entries"] if entry_span(e)[1] >= start and entry_span(e)[0] <= end),
                   key=lambda e: (entry_span(e)[0], disp(e.get("name", ""))))
    for r, e in enumerate(items, 2):
        von, bis = entry_span(e)
        art = e.get("type", "Custom")
        was = GRUENDE.get(art, (art, ""))[0].split(" ", 1)[-1] + (f" «{e['task']}»" if art == "Custom" and e.get("task") else "")
        vals = [disp(e.get("name", "")), was, von, bis, slot_label(e["slots"]), halbe_tage(entry_halbtage(e)),
                "" if entry_range(e) else days_label(e.get("days", [])), e.get("note", "")]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(r, col, v)
            c.font, c.border = font(), border
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    for col, width in zip("ABCDEFGH", (20, 22, 12, 12, 12, 12, 14, 34)):
        ws2.column_dimensions[col].width = width
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Feiertage")
    for col, txt in enumerate(["Datum", "Tag", "Feiertag", "frei"], 1):
        c = ws3.cell(1, col, txt)
        c.font, c.fill, c.border = font(bold=True), fill("E2E8F0"), border
    for r, (d, (name, slots)) in enumerate(ft.tage(ts["feiertage"], start, end).items(), 2):
        for col, v in enumerate([d, WOCHENTAGE[d.weekday()], name, slot_label(list(slots))], 1):
            c = ws3.cell(r, col, v)
            c.font, c.border = font(), border
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    for col, width in zip("ABCD", (12, 6, 28, 12)):
        ws3.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Formular «Ferien oder Abwesenheit eintragen» ─────────────────────

def _jp_slots() -> list[int]:
    slot = st.session_state.get(wk("j_slot")) or "Ganzer Tag"
    return [0, 1] if slot == "Ganzer Tag" else [0] if slot == "Vormittag" else [1]


def _jp_von_cb() -> None:
    """Neues «Von» nach «Bis»? Dann «Bis» auf den Freitag derselben Woche setzen."""
    ss = st.session_state
    von, bis = ss.get(wk("j_von")), ss.get(wk("j_bis"))
    if von and (not bis or bis < von):
        ss[wk("j_bis")] = von + timedelta(days=max(0, 4 - von.weekday()))


def _jp_add_cb() -> None:
    ss = st.session_state
    person = {person_label(s): s for s in active_shorts()}.get(ss.get(wk("j_employee")) or "")
    art = ss.get(wk("j_type")) or "Ferien"
    von, bis = ss.get(wk("j_von")), ss.get(wk("j_bis"))
    task = _s(ss.get(wk("j_task"))) if art == "Custom" else ""
    msg = None
    if not person:
        msg = "Bitte eine Person wählen."
    elif not von or not bis:
        msg = "Bitte «Von» und «Bis» wählen."
    elif bis < von:
        msg = "«Bis» liegt vor «Von»."
    elif (bis - von).days > 366:
        msg = "Höchstens ein Jahr am Stück – bitte aufteilen."
    elif art == "Custom" and not task:
        msg = "Bitte die Aufgabe wählen, die fest eingetragen werden soll."
    elif not entry_halbtage({"name": person, "von": von.isoformat(), "bis": bis.isoformat(), "slots": _jp_slots()}):
        msg = f"{disp(person)} arbeitet in diesem Zeitraum nie (Wochenende, Feiertage oder freie Tage laut Pensum)."
    if msg:
        ss[wk("j_msg")] = msg
        return
    m0, m1 = von - timedelta(days=von.weekday()), bis - timedelta(days=bis.weekday())
    jahr, kw, _ = m0.isocalendar()
    ts["entries"].append({
        "name": person, "type": art, "von": von.isoformat(), "bis": bis.isoformat(), "slots": _jp_slots(),
        "note": _s(ss.get(wk("j_note"))), "task": task,
        # für ältere App-Versionen: betroffene Wochen, aber keine Wochentage
        "days": [], "jahr": jahr, "kw_start": kw, "weeks": (m1 - m0).days // 7 + 1,
    })
    persist_entries()
    ss[wk("j_note")] = ""
    st.toast(f"{disp(person)}: {GRUENDE[art][0]} {range_label(von, bis)} eingetragen", icon="✅")


def _jp_form() -> None:
    ss = st.session_state
    with st.container(border=True):
        st.markdown('<div class="card-head"><div><h2 style="font-size:15.5px; margin:0;">Ferien oder Abwesenheit '
                    'eintragen</h2><div style="font-size:12.5px; color:var(--text-mute);">von – bis, auch über '
                    'mehrere Wochen · Wochenenden und Feiertage zählen nicht</div></div></div>', unsafe_allow_html=True)
        opts = [person_label(s) for s in active_shorts()]
        if ss.get(wk("j_employee")) not in opts:
            ss[wk("j_employee")] = opts[0] if opts else None
        st.markdown('<div class="field-label">Wer?</div>', unsafe_allow_html=True)
        st.selectbox("Person", opts, key=wk("j_employee"), label_visibility="collapsed")

        st.markdown('<div class="field-label">Was?</div>', unsafe_allow_html=True)
        if ss.get(wk("j_type")) not in T.gruende:
            ss[wk("j_type")] = "Ferien"
        art = st.segmented_control("Art", list(T.gruende), key=wk("j_type"), label_visibility="collapsed",
                                   format_func=lambda g: GRUENDE[g][0]) or "Ferien"
        if art == "Custom":
            task_opts = [t for t in team_task_options(ts["cfg"])
                         if t and t not in ("KRANK", "FERIEN", "Ferien", "Frei", ak.FEIERTAG)]
            st.selectbox("Welche Aufgabe?", task_opts, key=wk("j_task"), accept_new_options=True,
                         help="Aus der Liste wählen oder einen eigenen Code eintippen (z. B. Schulung)")

        heute = date.today()
        ss.setdefault(wk("j_von"), _as_date(info["monday"]))
        ss.setdefault(wk("j_bis"), _as_date(info["monday"]) + timedelta(days=4))
        lim = {"min_value": date(heute.year - 1, 1, 1), "max_value": date(heute.year + 2, 12, 31),
               "format": "DD.MM.YYYY"}
        c1, c2 = st.columns(2)
        with c1:
            st.markdown('<div class="field-label">Von</div>', unsafe_allow_html=True)
            st.date_input("Von", key=wk("j_von"), on_change=_jp_von_cb, label_visibility="collapsed", **lim)
        with c2:
            st.markdown('<div class="field-label">Bis <span>(inklusive)</span></div>', unsafe_allow_html=True)
            st.date_input("Bis", key=wk("j_bis"), label_visibility="collapsed", **lim)
        st.markdown('<div class="field-label">Halbtag <span>gilt für jeden Tag im Zeitraum</span></div>',
                    unsafe_allow_html=True)
        if ss.get(wk("j_slot")) not in ("Ganzer Tag", "Vormittag", "Nachmittag"):
            ss[wk("j_slot")] = "Ganzer Tag"
        st.segmented_control("Halbtag", ["Ganzer Tag", "Vormittag", "Nachmittag"], key=wk("j_slot"),
                             label_visibility="collapsed")
        st.markdown('<div class="field-label">Notiz <span>optional</span></div>', unsafe_allow_html=True)
        st.text_input("Notiz", key=wk("j_note"), label_visibility="collapsed", placeholder="z. B. Sommerferien")

        von, bis = ss.get(wk("j_von")), ss.get(wk("j_bis"))
        what = GRUENDE[art][0] + (f" «{_s(ss.get(wk('j_task')))}»" if art == "Custom" and ss.get(wk("j_task")) else "")
        if von and bis and bis >= von:
            person = {person_label(x): x for x in active_shorts()}.get(ss.get(wk("j_employee")) or "", "")
            n = entry_halbtage({"name": person, "von": von.isoformat(), "bis": bis.isoformat(), "slots": _jp_slots()})
            fe_namen = [name for name, _ in ft.tage(ts["feiertage"], von, bis).values()]
            when = html_escape(f"{range_label(von, bis)} · {halbe_tage(n)} Arbeitstag{'' if n == 2 else 'e'} · "
                               f"{kw_range_label(von, bis)}" + (f" · ohne {', '.join(fe_namen)}" if fe_namen else ""))
        else:
            when = "<i>«Bis» liegt vor «Von»</i>"
        st.markdown(f'<div class="preview-line">→ <b>{html_escape(ss.get(wk("j_employee")) or "–")}</b> · '
                    f'{html_escape(what)} · {when}</div>', unsafe_allow_html=True)
        msg = ss.pop(wk("j_msg"), None)
        if msg:
            st.warning(msg)
        st.button("＋  Eintragen", key=wk("j_add"), type="primary", width="stretch", on_click=_jp_add_cb)


def _jp_entries(jahr: int, daten: list[dict]) -> None:
    start, end = daten[0]["monday"], daten[-1]["monday"] + timedelta(days=4)
    heute = date.today()
    items = sorted(((i, e) for i, e in enumerate(ts["entries"])
                    if entry_span(e)[1] >= start and entry_span(e)[0] <= end),
                   key=lambda ie: (entry_span(ie[1])[0], disp(ie[1].get("name", ""))))
    with st.container(border=True):
        st.markdown(f'<div class="card-head"><div><h2 style="font-size:15.5px; margin:0;">Einträge {jahr}</h2>'
                    f'<div style="font-size:12.5px; color:var(--text-mute);">{len(items)} · nach Datum · aus dem '
                    'Jahresplan und aus Schritt 1 · vergangene blass</div></div></div>', unsafe_allow_html=True)
        if not items:
            st.markdown('<div class="empty-state"><div class="icon">🏖️</div><div class="ttl">Noch nichts eingetragen.</div>'
                        '<div>Links Ferien oder freie Tage eintragen – sie erscheinen oben in der Übersicht und '
                        'automatisch im Wochenplan der betroffenen Wochen.</div></div>', unsafe_allow_html=True)
            return
        box = st.container(height=470, border=False) if len(items) > 9 else st.container(border=False)
        with box:
            for i, e in items:
                _entry_row(i, e, entry_span(e)[1] >= heute)


# ── Feiertage ─────────────────────────────────────────────────────────

def _set_feiertage(fe: dict, msg: str = "Feiertage übernommen – gelten sofort für jeden Wochenplan.") -> None:
    ts["feiertage"] = ft.normalize(fe)
    ts["fe_version"] += 1
    persist_entries()
    st.toast(msg, icon="🎉")


def _on_fe_standard(key: str, names: list[str]) -> None:
    fe = ft.copy_of(ts["feiertage"])
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        if "frei" in changes:
            name, wert = names[int(ridx)], FREI_WAHL.get(_s(changes["frei"]))
            if wert:
                fe["standard"][name] = wert
            else:
                fe["standard"].pop(name, None)
    _set_feiertage(fe)


def _on_fe_eigene(key: str, base: list[dict]) -> None:
    rows = [dict(e) for e in base]
    removed = set()
    for ridx, changes in st.session_state[key].get("edited_rows", {}).items():
        i = int(ridx)
        for col, val in changes.items():
            if col == "entfernen" and _b(val):
                removed.add(i)
            elif col == "Bezeichnung":
                rows[i]["name"] = _s(val) or rows[i]["name"]
            elif col == "frei":
                rows[i]["frei"] = FREI_WAHL.get(_s(val)) or "ganz"
    fe = ft.copy_of(ts["feiertage"])
    fe["eigene"] = [r for i, r in enumerate(rows) if i not in removed]
    _set_feiertage(fe, "Freier Tag entfernt" if removed else "Feiertage übernommen")


def _jp_feiertage(jahr: int) -> None:
    fe = ts["feiertage"]
    ver = ts["fe_version"]
    with st.container(border=True):
        st.markdown(f'<div class="card-head"><div><h2 style="font-size:15.5px; margin:0;">🎉 Feiertage {jahr}</h2>'
                    '<div style="font-size:12.5px; color:var(--text-mute);">Gelten fürs ganze Team: alle haben frei, '
                    'es wird nichts geplant und keine Besetzung geprüft. Halbe Tage sind möglich (z. B. Heiligabend '
                    'nur Nachmittag). Die Daten rechnet die App jedes Jahr selbst aus.</div></div></div>',
                    unsafe_allow_html=True)
        left, right = st.columns([1.1, 1], gap="medium")
        with left:
            st.markdown('<div class="field-label">Feiertage <span>«–» = wird gearbeitet</span></div>',
                        unsafe_allow_html=True)
            kat = ft.katalog(jahr)
            names = [n for n, _ in kat]
            df = pd.DataFrame([{"Feiertag": n, "Datum": datum_label(d, True) + (" · Wochenende" if d.weekday() > 4 else ""),
                                "frei": FREI_WAHL_LABEL[fe["standard"].get(n)]} for n, d in kat],
                              columns=["Feiertag", "Datum", "frei"])
            key = wk(f"fe_std_{jahr}_{ver}")
            st.data_editor(df, key=key, hide_index=True, num_rows="fixed", width="stretch",
                           height=_editor_height(len(df)), on_change=_on_fe_standard, args=(key, names),
                           column_config={
                               "Feiertag": st.column_config.TextColumn("Feiertag", disabled=True),
                               "Datum": st.column_config.TextColumn(f"{jahr}", disabled=True),
                               "frei": st.column_config.SelectboxColumn("frei", options=list(FREI_WAHL), required=True,
                                                                        help="ganzer Tag, nur Vormittag, nur Nachmittag "
                                                                             "oder «–» (normaler Arbeitstag)")})
        with right:
            st.markdown('<div class="field-label">Eigene freie Tage <span>z. B. Brückentag, Betriebsausflug</span></div>',
                        unsafe_allow_html=True)
            with st.form(wk(f"fe_add_{ver}"), border=True):
                c = st.columns([1.25, 1.6, 1.2], vertical_alignment="bottom")
                d = c[0].date_input("Datum", value=None, format="DD.MM.YYYY",
                                    min_value=date(jahr - 1, 1, 1), max_value=date(jahr + 2, 12, 31))
                name = c[1].text_input("Bezeichnung", placeholder="z. B. Brückentag")
                frei = c[2].selectbox("frei", ["ganzer Tag", "Vormittag", "Nachmittag"])
                if st.form_submit_button("➕  Freien Tag hinzufügen", type="primary", width="stretch"):
                    if not d:
                        st.error("Bitte ein Datum wählen.")
                    elif d.weekday() > 4:
                        st.error("Das ist ein Wochenende – da wird ohnehin nicht geplant.")
                    else:
                        new = ft.copy_of(fe)
                        new["eigene"].append({"datum": d.isoformat(), "name": name.strip() or "freier Tag",
                                              "frei": FREI_WAHL[frei]})
                        _set_feiertage(new, f"{datum_label(d, True)} als freier Tag eingetragen")
                        st.rerun()
            base = list(fe["eigene"])
            if base:
                df = pd.DataFrame([{"Datum": datum_label(date.fromisoformat(e["datum"]), True), "Bezeichnung": e["name"],
                                    "frei": FREI_WAHL_LABEL[e["frei"]], "entfernen": False} for e in base],
                                  columns=["Datum", "Bezeichnung", "frei", "entfernen"])
                key = wk(f"fe_own_{ver}")
                st.data_editor(df, key=key, hide_index=True, num_rows="fixed", width="stretch",
                               height=_editor_height(len(df)), on_change=_on_fe_eigene, args=(key, base),
                               column_config={
                                   "Datum": st.column_config.TextColumn("Datum", disabled=True),
                                   "Bezeichnung": st.column_config.TextColumn("Bezeichnung"),
                                   "frei": st.column_config.SelectboxColumn(
                                       "frei", options=["ganzer Tag", "Vormittag", "Nachmittag"], required=True),
                                   "entfernen": st.column_config.CheckboxColumn("entfernen", width="small")})
            else:
                st.caption("Noch keine eigenen freien Tage.")
            kommende = [(d, n, s) for d, (n, s) in ft.tage(fe, date.today(), date.today() + timedelta(days=200)).items()]
            if kommende:
                st.markdown('<div class="field-label" style="margin-top:14px;">Die nächsten freien Tage</div>',
                            unsafe_allow_html=True)
                st.markdown(" · ".join(f"**{datum_label(d, True)}** {html_escape(ft.tag_text(n, s))}"
                                       for d, n, s in kommende[:6]))


def render_jahresplan() -> None:
    heute_jahr = date.today().isocalendar()[0]
    jahre = [heute_jahr, heute_jahr + 1]
    if st.session_state.get("jp_jahr") not in jahre:
        st.session_state.jp_jahr = info["iso_year"] if info["iso_year"] in jahre else heute_jahr
    st.markdown(f'<div class="page-head"><h2>📆 Jahresplan · Ferien & Feiertage · {TEAM}</h2><div class="sub">Ferien, '
                'freie Tage und Feiertage fürs ganze Jahr. Der Wochenplan (Schritt 1–2) übernimmt sie automatisch – '
                'für die nächste Woche lässt sich dort weiterhin alles einzeln anpassen.</div></div>',
                unsafe_allow_html=True)
    render_save_bar()
    c1, c2, c3 = st.columns([1.5, 5.6, 2.1], vertical_alignment="center")
    with c1:
        st.segmented_control("Jahr", jahre, key="jp_jahr", label_visibility="collapsed")
    jahr = st.session_state.jp_jahr or heute_jahr
    daten = jahres_daten(jahr)
    rot = [w["kw"] for w in daten if w["hard"]]
    gelb = [w["kw"] for w in daten if w["knapp"] and not w["hard"]]
    n_fe = sum(len(w["feiertage"]) for w in daten)
    teile = [f"🎉 {n_fe} Feiertag{'e' if n_fe != 1 else ''}"]
    if rot:
        teile.append(f"🔴 geht nicht auf: KW {', '.join(map(str, rot[:8]))}{' …' if len(rot) > 8 else ''}")
    if gelb:
        teile.append(f"🟠 knapp: KW {', '.join(map(str, gelb[:8]))}{' …' if len(gelb) > 8 else ''}")
    if not rot and not gelb:
        teile.append("🟢 alle kommenden Wochen gehen auf")
    c2.markdown(f'<div style="font-size:13px; color:var(--text-dim);">{" · ".join(teile)}</div>', unsafe_allow_html=True)
    c3.download_button("⬇  Jahresplan als Excel", data=jahresplan_excel(jahr, daten),
                       file_name=f"Jahresplan_{TEAM}_{jahr}.xlsx", key=wk("jp_download"), width="stretch",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       help="Übersicht pro KW zum Ausdrucken, dazu alle Einträge und Feiertage als Liste")
    with st.container(border=True):
        st.markdown(_jp_grid_html(daten) + _jp_legend_html(), unsafe_allow_html=True)
    left, right = st.columns([1, 1.15], gap="medium")
    with left:
        _jp_form()
    with right:
        _jp_entries(jahr, daten)
    _jp_feiertage(jahr)


# ══════════════════════════════════════════════════════════════════════
#  Seite wählen
# ══════════════════════════════════════════════════════════════════════

if st.session_state.step == 3:
    render_team_page()
elif st.session_state.step == 4:
    render_jahresplan()
elif st.session_state.step == 1 or ts["result"] is None:
    render_planung()
else:
    render_ergebnis()
