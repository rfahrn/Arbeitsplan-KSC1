#!/usr/bin/env python3
"""
Regel-Tests für das KSV-Team (teamplaner.py + team_ksv.json). Läuft ohne pytest:

    python tests/test_ksv.py              # 12 Wochen ab nächster Woche
    python tests/test_ksv.py 2026-10-19 20

… und auch mit pytest (`pytest tests/`).
"""
from __future__ import annotations

import sys
from collections import Counter
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import teamplaner as tp  # noqa: E402

START = datetime(2026, 10, 19)   # KW 43


def wochen(start: datetime, n: int):
    for i in range(n):
        monday = start + timedelta(weeks=i)
        yield monday.isocalendar()[1], monday


def pruefe_woche(sched: "tp.Schedule") -> list[str]:
    """Harte Regeln (❌ aus validate_schedule) plus strukturelle Prüfungen."""
    cfg = sched.cfg
    fehler = [i for i in tp.validate_schedule(sched) if i.startswith("❌")]
    for n, e in sched.employees.items():
        for d in range(5):
            for s in range(2):
                task = sched.get_task(n, d, s)
                # niemand arbeitet, wenn laut Stammdaten nicht da
                if task and not tp.is_absent_code(task) and not e.is_available(d, s):
                    fehler.append(f"{n} {tp.DAY_KEYS[d]} {tp.SLOT_KEYS[s]}: {task}, aber nicht anwesend")
                # jeder anwesende Halbtag hat eine Aufgabe (Auffüllen)
                if e.is_available(d, s) and not task and (n, d, s) not in sched.notes:
                    fehler.append(f"{n} {tp.DAY_KEYS[d]} {tp.SLOT_KEYS[s]}: leer")
        # feste Einsätze stehen im Plan
        for d, s, task, _z in e.fix:
            if task and sched.is_present(n, d, s) and not sched.get_task(n, d, s).startswith(task):
                fehler.append(f"Fixeinsatz {n} {tp.DAY_KEYS[d]} {tp.SLOT_KEYS[s]}: {sched.get_task(n, d, s)!r} statt {task}")
    # Ganztags-Aufgaben: dieselbe Person VM + NM (bei voller Besetzung)
    for t in tp.auto_tasks(cfg):
        if t["dauer"] != "ganzer Tag":
            continue
        for d in range(5):
            vm, nm = set(sched.holders(t["code"], d, 0)), set(sched.holders(t["code"], d, 1))
            if vm != nm:
                fehler.append(f"{t['code']} {tp.DAY_KEYS[d]}: VM {sorted(vm)} / NM {sorted(nm)}")
    return fehler


def test_konfiguration_gueltig():
    cfg = tp.load_config()
    assert not tp.validate_config(cfg), tp.validate_config(cfg)
    assert tp.config_to_json(tp.normalize_config(cfg)) == tp.config_to_json(cfg)


def test_regeln_mehrere_wochen():
    cfg = tp.load_config()
    for kw, monday in wochen(START, 12):
        sched = tp.build_schedule(kw, monday, cfg=cfg)
        fehler = pruefe_woche(sched)
        assert not fehler, f"KW {kw}: {fehler}"
        assert not [i for i in tp.validate_schedule(sched) if not i.startswith("ℹ️")], kw


def test_regeln_aus_dem_plan_2026():
    """Muster aus dem Excel-Plan 2026: Kim nie KSV1, Pia Quality Di/Do VM, Montag 2× contact,
    Medine Mi NM mit Arbeitszeit im Wochenwechsel, Nicht-Arbeitstage = Frei."""
    cfg = tp.load_config()
    for kw, monday in wochen(START, 6):
        s = tp.build_schedule(kw, monday, cfg=cfg)
        assert not any(s.get_task("Kim", d, sl).startswith("KSV1") for d in range(5) for sl in range(2))
        assert s.get_task("Pia", 1, 0) == "Quality" and s.get_task("Pia", 3, 0) == "Quality"
        assert sorted(s.get_task(n, 0, 0) for n in s.holders("contact", 0, 0)) == ["contact 1", "contact 2"]
        assert all(len(s.holders("contact", d, 0)) == 1 for d in range(1, 5))
        zeit = "12-16" if kw % 2 else "13-17"
        assert s.get_task("Medine", 2, 1).endswith(zeit), (kw, s.get_task("Medine", 2, 1))
        assert s.get_task("Medine", 3, 0) == "Frei" and s.get_task("Kim", 4, 1) == "Frei"
        for d in range(5):
            assert len(s.holders("KD Komm", d, 0)) == 1 and len(s.holders("KD Komm", d, 1)) == 1


def test_faire_verteilung():
    """Jede Person mit «ja» hat pro Woche genau einen contact-Tag (6 contact-Tage, 6 Personen)."""
    cfg = tp.load_config()
    for kw, monday in wochen(START, 10):
        s = tp.build_schedule(kw, monday, cfg=cfg)
        tage = Counter(n for d in range(5) for n in s.holders("contact", d, 0))
        ja = [m["short"] for m in cfg["mitarbeiter"] if "contact" in m["kann"]]
        assert all(tage[n] == 1 for n in ja), (kw, tage)
        assert "Melanie" not in tage   # nur Reserve


def test_determinismus():
    cfg = tp.load_config()
    a = tp.build_schedule(43, START, cfg=cfg)
    b = tp.build_schedule(43, START, cfg=cfg)
    assert a.schedule == b.schedule


def test_abwesenheiten_und_reserve():
    """Ferien/Krank/Termin blockieren Halbtage; fehlen viele, springt die Reserve ein (mit Hinweis)."""
    cfg = tp.load_config()
    ov = {(n, d, s): "Ferien" for n in ("Joy", "Maria", "Jana") for d in range(5) for s in range(2)}
    ov[("Pia", 2, 1)] = "*Arzt 14:30"
    s = tp.build_schedule(43, START, ov, cfg=cfg)
    assert all(s.get_task("Joy", d, sl) == "Ferien" for d in range(5) for sl in range(2))
    assert s.get_task("Pia", 2, 1) == "" and s.notes[("Pia", 2, 1)] == "Arzt 14:30"
    issues = tp.validate_schedule(s)
    assert any("Reserve eingesetzt" in i for i in issues), issues
    # Melanie hilft als Reserve bei contact aus
    assert any(s.get_task("Melanie", d, 0).startswith("contact") for d in range(5))
    # Warnung vor dem Planen
    emps = tp.create_employees(43, cfg)
    warn = tp.check_rule_impact([k for k, v in ov.items() if v == "Ferien"], emps, cfg)
    assert any(w.startswith("KRITISCH") for w in warn), warn


def test_ganztags_aufteilen_wenn_niemand_ganztags_frei():
    cfg = tp.load_config()
    for m in cfg["mitarbeiter"]:   # Redcare darf nur, wer halbtags frei ist
        tp.set_stufe(m, "Redcare", "ja" if m["short"] in ("Kim",) else "–")
    for m in cfg["mitarbeiter"]:
        if m["short"] == "Medine":
            tp.set_stufe(m, "Redcare", "ja")
    s = tp.build_schedule(43, START, cfg=cfg)
    # Mittwoch: Kim nur VM, Medine nur NM → aufgeteilt (beide sind am anderen Halbtag nicht da → kein Hinweis)
    assert s.holders("Redcare", 2, 0) == ["Kim"] and s.holders("Redcare", 2, 1) == ["Medine"]
    assert dict(tp.day_overview(s))["Redcare"][2] == "VM: Kim K. · NM: Medine A."
    assert not any("aufgeteilt" in i for i in tp.validate_schedule(s))
    # manuell aufgeteilt, obwohl die Person den ganzen Tag da ist → Hinweis
    kim_do = s.get_task("Kim", 3, 0)
    s.schedule["Kim"][(3, 1)] = "KSV2"
    assert kim_do.startswith("Redcare")
    assert any("Kim K. Donnerstag: Redcare nur am VM" in i for i in tp.validate_schedule(s))


def test_neue_aufgabe_ohne_code():
    """Neue Aufgabe in der Konfiguration (wie Redcare im August) wird sofort geplant."""
    cfg = tp.load_config()
    cfg["aufgaben"].append({"code": "Retouren", "bedeutung": "Retouren", "farbe": "FF66CC", "dauer": "nur VM",
                            "anzahl": {"Mo": 1, "Di": 0, "Mi": 1, "Do": 0, "Fr": 0}})
    for m in cfg["mitarbeiter"]:
        tp.set_stufe(m, "Retouren", "ja")
    cfg = tp.normalize_config(cfg)
    assert not tp.config_errors(cfg)
    s = tp.build_schedule(43, START, cfg=cfg)
    assert len(s.holders("Retouren", 0, 0)) == 1 and len(s.holders("Retouren", 2, 0)) == 1
    assert not s.holders("Retouren", 1, 0) and not s.holders("Retouren", 0, 1)
    assert not pruefe_woche(s)
    assert tp.get_task_color("Retouren", cfg) == "FF66CC"


def test_konfiguration_bearbeiten():
    cfg = tp.load_config()
    base = tp.normalize_config(cfg)
    # Umbenennen zieht Wer-darf-was, feste Einsätze und Auffüllen mit
    tp.rename_task(cfg, "KSV2", "KSV 2")
    tp.rename_task(cfg, "Quality", "QS")
    assert cfg["regeln"]["auffuellen"] == "KSV 2"
    pia = next(m for m in cfg["mitarbeiter"] if m["short"] == "Pia")
    assert {fx["task"] for fx in pia["fix"]} == {"QS"}
    assert not tp.config_errors(cfg)
    assert "Neue Aufgabe: QS" in tp.describe_changes(base, cfg)
    # Entfernen räumt Wer-darf-was auf
    tp.remove_task(cfg, "Redcare")
    assert not any("Redcare" in m["kann"] + m["reserve"] for m in cfg["mitarbeiter"])
    assert not tp.config_errors(cfg)
    assert not pruefe_woche(tp.build_schedule(44, START + timedelta(weeks=1), cfg=cfg))
    # Fehler: doppelter Code, reservierter Code, unbekannte Auffüll-Aufgabe
    bad = tp.normalize_config(base)
    bad["aufgaben"].append({"code": "ksv1", "dauer": "nur manuell"})
    assert any(e.startswith("Aufgabe doppelt") for e in tp.config_errors(bad))
    bad = tp.normalize_config(base)
    bad["aufgaben"].append({"code": "Frei", "dauer": "nur manuell"})
    assert any("reserviert" in e for e in tp.config_errors(bad))
    bad = tp.normalize_config(base)
    bad["regeln"]["auffuellen"] = "Gibtsnicht"
    assert tp.config_errors(bad)


def test_feste_einsaetze_als_liste():
    base = tp.load_config()
    cfg = tp.normalize_config(base)
    rows = tp.recurring_rows(cfg)
    assert {"short": "Medine", "tag": "Mi", "halbtag": "NM", "task": "", "zusatz": "12-16", "kw": "ungerade"} in rows
    assert not tp.apply_recurring_rows(cfg, rows)
    assert tp.config_to_json(cfg) == tp.config_to_json(base)
    rows.append({"short": "Kim", "tag": "Fr", "halbtag": "ganz", "task": "frei", "zusatz": "", "kw": "gerade"})
    assert not tp.apply_recurring_rows(cfg, rows)
    assert tp.describe_changes(base, cfg) == ["Kim K.: alle 2 Wochen frei → Fr ganz (gerade KW)"]
    s = tp.build_schedule(44, START + timedelta(weeks=1), cfg=cfg)   # KW 44 = gerade
    assert s.get_task("Kim", 4, 0) == "Frei"
    assert tp.apply_recurring_rows(cfg, rows + [{"short": "Kim", "tag": "Mo", "halbtag": "ganz", "task": "frei",
                                                 "zusatz": "", "kw": None}])


def test_excel_wie_jahresplan():
    """Spalten A–L = ein Wochenblock wie im KSV-Jahresplan (Woche | Name | Vorname | Mo VM … Fr NM)."""
    from openpyxl import load_workbook
    cfg = tp.load_config()
    s = tp.build_schedule(43, START, cfg=cfg)
    buf = BytesIO()
    tp.write_excel(s, buf)
    ws = load_workbook(BytesIO(buf.getvalue())).active
    assert ws["A1"].value == "Woche 43 | 19.10.2026 – 23.10.2026"
    assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == ("Name", "Vorname", "Montag\n19.10.2026")
    assert "C2:D2" in {str(r) for r in ws.merged_cells.ranges}
    assert (ws["C3"].value, ws["D3"].value) == ("Vormittag", "Nachmittag")
    assert (ws["A4"].value, ws["B4"].value) == ("A.", "Medine")
    assert ws["A4"].fill.fgColor.rgb.endswith("FFC7CE")
    farben = {ws.cell(row=r, column=c).value: ws.cell(row=r, column=c).fill.fgColor.rgb[-6:]
              for r in range(4, 11) for c in range(3, 13)}
    assert farben["KSV1"] == "92D050" and farben["KSV2"] == "00B050" and farben["Frei"] == "00B0F0"
    assert any(ws.cell(row=r, column=1).value == "Legende" for r in range(11, 20))


if __name__ == "__main__":
    start = datetime.strptime(sys.argv[1], "%Y-%m-%d") if len(sys.argv) > 1 else tp.get_next_monday()
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 12
    cfg = tp.load_config()
    print("team_ksv.json:", "ok" if not tp.validate_config(cfg) else tp.validate_config(cfg))
    total = 0
    for kw, monday in wochen(start, n):
        sched = tp.build_schedule(kw, monday, cfg=cfg)
        fehler = pruefe_woche(sched)
        total += len(fehler)
        print(f"{'✅' if not fehler else '❌'} KW {kw:2d} ({monday:%d.%m.%Y}): {len(fehler)} Fehler")
        for f in fehler:
            print("     ", f)
    for name in [k for k in list(globals()) if k.startswith("test_")]:
        try:
            globals()[name]()
            print(f"✅ {name}")
        except AssertionError as exc:
            total += 1
            print(f"❌ {name}: {exc}")
    print("\nERGEBNIS:", "alle Regeln eingehalten 🎉" if total == 0 else f"{total} Fehler")
    sys.exit(0 if total == 0 else 1)
